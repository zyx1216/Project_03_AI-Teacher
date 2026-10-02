# -*- coding: utf-8 -*-
"""导出模板服务（v3.4.0）。

学校 Excel/Word/PDF 模板文件存 data/templates/exports/，元数据落
export_templates 表。提供上传、列表、设默认、删除、预览；
格式不合法直接抛 ValueError，不静默降级。
"""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import config
from models.models import ExportTemplate

EXPORT_TEMPLATE_DIR = config.DATA_DIR / "templates" / "exports"
TYPE_SUFFIXES = {"excel": (".xlsx", ".xlsm"), "word": (".docx",),
                 "pdf": (".pdf",)}
TYPE_LABELS = {"excel": "Excel", "word": "Word", "pdf": "PDF"}
PREVIEW_ROWS = 10
PREVIEW_PARAGRAPHS = 15


DEFAULT_SUFFIX = {"excel": ".xlsx", "word": ".docx", "pdf": ".pdf"}


def _suffix_for(file_name: str | None, file_type: str) -> str:
    """后缀以上传文件名称为准；没给名称时用该类型默认后缀。"""
    allowed = TYPE_SUFFIXES.get(str(file_type), ())
    raw_suffix = Path(str(file_name or "")).suffix.lower()
    if not raw_suffix:
        return DEFAULT_SUFFIX.get(str(file_type), allowed[0] if allowed else "")
    if raw_suffix not in allowed:
        raise ValueError(
            f"{TYPE_LABELS.get(file_type, file_type)} 模板只支持 "
            + "、".join(allowed))
    return raw_suffix


def _validate(file_type: str, file_bytes: bytes) -> None:
    """按类型校验文件可被对应库打开，不合法直接报错。"""
    if not file_bytes:
        raise ValueError("模板文件为空。")
    try:
        if file_type == "excel":
            from openpyxl import load_workbook
            load_workbook(io.BytesIO(file_bytes), read_only=True)
        elif file_type == "word":
            from docx import Document
            Document(io.BytesIO(file_bytes))
        elif file_type == "pdf":
            import fitz
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                if doc.page_count <= 0:
                    raise ValueError("PDF 模板没有页面。")
        else:
            raise ValueError("不支持的模板类型。")
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 —— 统一转成友好错误
        raise ValueError(f"模板文件无法解析：{exc}") from exc


def save_export_template(session, name: str, file_type: str,
                         file_bytes: bytes,
                         file_name: str | None = None) -> ExportTemplate:
    """保存学校模板：校验通过后落盘并写库。

    file_name 传上传文件的原始文件名（用于判定后缀），不传则用默认后缀。
    """
    file_type = str(file_type)
    if file_type not in TYPE_SUFFIXES:
        raise ValueError("不支持的模板类型。")
    clean_name = str(name or "").strip()
    if not clean_name:
        raise ValueError("模板名称不能为空。")
    _validate(file_type, file_bytes)
    suffix = _suffix_for(file_name, file_type)
    EXPORT_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    file_name = f"{file_type}_{stamp}_{abs(hash(clean_name)) % 10000}{suffix}"
    path = EXPORT_TEMPLATE_DIR / file_name
    path.write_bytes(file_bytes)
    row = ExportTemplate(name=clean_name[:200], type=file_type,
                         file_path=str(path), is_default=False)
    session.add(row)
    session.flush()
    return row


def list_export_templates(session, file_type: str | None = None
                          ) -> list[ExportTemplate]:
    """列出模板；默认模板排前面。"""
    query = session.query(ExportTemplate)
    if file_type:
        query = query.filter(ExportTemplate.type == str(file_type))
    return query.order_by(ExportTemplate.is_default.desc(),
                          ExportTemplate.id.desc()).all()


def get_default(session, file_type: str) -> ExportTemplate | None:
    """取某类型的默认模板。"""
    return (session.query(ExportTemplate)
            .filter(ExportTemplate.type == str(file_type),
                    ExportTemplate.is_default.is_(True))
            .order_by(ExportTemplate.id.desc()).first())


def set_default(session, template_id: int) -> ExportTemplate:
    """设为该类型唯一默认模板。"""
    row = session.get(ExportTemplate, int(template_id))
    if row is None:
        raise ValueError("模板不存在。")
    for other in (session.query(ExportTemplate)
                  .filter(ExportTemplate.type == row.type,
                          ExportTemplate.is_default.is_(True)).all()):
        other.is_default = False
    row.is_default = True
    session.flush()
    return row


def delete_export_template(session, template_id: int) -> bool:
    """删除模板记录与磁盘文件。"""
    row = session.get(ExportTemplate, int(template_id))
    if row is None:
        return False
    path = Path(row.file_path) if row.file_path else None
    session.delete(row)
    session.flush()
    if path is not None:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            pass  # 文件删不掉不影响记录清理
    return True


def template_path(row: ExportTemplate) -> Path:
    return Path(row.file_path)


def preview(session, template_id: int) -> dict:
    """预览模板：Excel 读首表前 10 行，Word 读前 15 段，PDF 读页数。"""
    row = session.get(ExportTemplate, int(template_id))
    if row is None:
        raise ValueError("模板不存在。")
    path = template_path(row)
    if not path.exists():
        raise ValueError("模板文件已丢失，请重新上传。")
    data = path.read_bytes()
    if row.type == "excel":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = []
        for index, row_values in enumerate(ws.iter_rows(values_only=True)):
            if index >= PREVIEW_ROWS:
                break
            rows.append(["" if v is None else str(v) for v in row_values])
        sheets = list(wb.sheetnames)
        wb.close()
        return {"kind": "excel", "sheets": sheets, "rows": rows}
    if row.type == "word":
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables = len(doc.tables)
        return {"kind": "word", "paragraphs": paragraphs[:PREVIEW_PARAGRAPHS],
                "table_count": tables}
    import fitz
    with fitz.open(stream=data, filetype="pdf") as doc:
        return {"kind": "pdf", "pages": doc.page_count}
