# -*- coding: utf-8 -*-
"""
PPT 生成层（离线、确定性，不调用任何模型）。

直接从已保存的教案字段排版，避免在"导出 PPT"这一步再引入一个可能失败的网络点。
固定结构：封面 → 学习目标 → 导入 → 新授（内容多自动分页）→ 例题与练习
→ 课堂小结 → 作业布置。

排版约束（开发指令要求）：
- 白底、标题深蓝 #1F4E79；
- 每页要点不超过 5 条；
- 每条不超过 20 字，超出截断（教案正文可在 Word 里看完整版）；
- 输出标准 .pptx，可在 WPS / PowerPoint 里二次编辑。
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt, Emu

# 主题色：深蓝标题
DEEP_BLUE = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x33, 0x33, 0x33)

MAX_BULLETS = 5      # 每页最多要点数
MAX_CHARS = 20       # 每条最多字数

# 环节 -> PPT 页标题（教案 5 环节映射到固定的课件结构）
STAGE_SLIDE_TITLES = {
    "导入": "课堂导入",
    "新授": "新知讲授",
    "巩固练习": "例题与练习",
    "课堂小结": "课堂小结",
    "作业布置": "作业布置",
}

# 宽屏 16:9
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)


def _split_bullets(content: str) -> list[str]:
    """把一段教学内容拆成短要点：先按换行/分号/句号切，再截断到 20 字。"""
    import re
    if not content or not content.strip():
        return []
    # 按换行、中文分号句号、问号感叹号切，保留短句
    parts = re.split(r"[\n；;。！？!?]+", content)
    bullets = []
    for part in parts:
        text = part.strip().lstrip("-·•●0123456789.、 ")
        if not text:
            continue
        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS] + "…"
        bullets.append(text)
    return bullets


def _paginate(bullets: list[str]) -> list[list[str]]:
    """把要点按每页最多 MAX_BULLETS 条分页。"""
    return [bullets[i:i + MAX_BULLETS] for i in range(0, len(bullets), MAX_BULLETS)]


def _new_presentation() -> Presentation:
    """建白底 16:9 演示文稿。"""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def _blank_slide(prs: Presentation):
    """加一张空白版式的幻灯片（不依赖模板自带占位符）。"""
    layout = prs.slide_layouts[6]  # 6 = 空白
    return prs.slides.add_slide(layout)


def _add_title(slide, text: str):
    """在幻灯片顶部加深蓝色标题，并加一条同色细线。"""
    box = slide.shapes.add_textbox(Emu(600000), Emu(350000),
                                   Emu(11000000), Emu(800000))
    frame = box.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = DEEP_BLUE

    # 标题下的装饰线
    from pptx.enum.shapes import MSO_SHAPE
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Emu(600000), Emu(1200000),
                                  Emu(11000000), Emu(45000))
    line.fill.solid()
    line.fill.fore_color.rgb = DEEP_BLUE
    line.line.fill.background()
    return slide


def _add_bullets(slide, bullets: list[str], top: int = 1500000):
    """在正文区写要点列表。"""
    box = slide.shapes.add_textbox(Emu(800000), Emu(top),
                                   Emu(10600000), Emu(4800000))
    frame = box.text_frame
    frame.word_wrap = True
    for i, text in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = f"• {text}"
        run.font.size = Pt(22)
        run.font.color.rgb = BLACK
        p.space_after = Pt(10)


def _add_cover(prs: Presentation, title: str, subtitle: str):
    """封面：大标题居中 + 副标题（年级/章节）。"""
    slide = _blank_slide(prs)

    title_box = slide.shapes.add_textbox(Emu(1000000), Emu(2400000),
                                         Emu(10200000), Emu(1400000))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = DEEP_BLUE

    if subtitle:
        sub_box = slide.shapes.add_textbox(Emu(1000000), Emu(4000000),
                                           Emu(10200000), Emu(800000))
        sf = sub_box.text_frame
        sp = sf.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        srun = sp.add_run()
        srun.text = subtitle
        srun.font.size = Pt(22)
        srun.font.color.rgb = BLACK


def _add_objectives(prs: Presentation, objectives: dict):
    """学习目标页：三维目标各成一条要点（已截到 20 字）。"""
    labels = [
        ("知识与技能", objectives.get("knowledge", "")),
        ("过程与方法", objectives.get("process", "")),
        ("情感态度", objectives.get("emotion", "")),
    ]
    bullets = []
    for label, text in labels:
        text = (text or "").strip()
        if text:
            text = f"{label}：{text}"
            bullets.append(text[:MAX_CHARS] + ("…" if len(text) > MAX_CHARS else ""))
    slide = _blank_slide(prs)
    _add_title(slide, "学习目标")
    _add_bullets(slide, bullets or ["（教案中未填写教学目标）"])


def _add_stage_slides(prs: Presentation, stage: str, content: str):
    """把一个教学环节排到一页或多页（超过 5 条自动分页，标题带序号）。"""
    base_title = STAGE_SLIDE_TITLES.get(stage, stage)
    pages = _paginate(_split_bullets(content))
    if not pages:
        # 空环节不输出整页空白，保持课件干净
        return
    total = len(pages)
    for idx, bullets in enumerate(pages):
        title = base_title if total == 1 else f"{base_title}（{idx + 1}）"
        slide = _blank_slide(prs)
        _add_title(slide, title)
        _add_bullets(slide, bullets)


def build_slides(prs: Presentation, title: str, subtitle: str, plan: dict) -> int:
    """按固定结构往演示文稿里填内容，返回页数。"""
    _add_cover(prs, title, subtitle)
    _add_objectives(prs, plan.get("objectives", {}))
    for step in plan.get("process", []):
        _add_stage_slides(prs, step.get("stage", ""), step.get("content", ""))
    return len(prs.slides._sldIdLst)


def generate_ppt(lesson, plan: dict) -> bytes:
    """
    从已保存教案生成 PPT，返回 .pptx 字节。
    lesson: LessonPlan ORM 对象（取标题/年级/章节）；plan: 教案内容字典。
    """
    import io
    prs = _new_presentation()
    subtitle = "　".join(x for x in [lesson.grade, lesson.chapter] if x)
    build_slides(prs, lesson.title, subtitle, plan)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()