# -*- coding: utf-8 -*-
"""
学情工作台页面。

顶部 5 个标签页：学生管理 / 成绩管理 / 考试分析 / 趋势分析 / 学生画像。
本文件只负责界面与交互；统计口径在 utils/stats.py，数据库读写在
utils/student_service.py、utils/exam_service.py，Excel 识别在 utils/excel_handler.py。
"""

from datetime import date, datetime

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from utils.db import SessionLocal
from utils import excel_handler as eh
from utils import exam_service as es
from utils import stats as stats_mod
from utils import student_service as ss
from utils import llm_client


# ---------------------------------------------------------------------------
# 通用小工具
# ---------------------------------------------------------------------------

def _read_excel_sheet(uploaded, sheet_name):
    """读取上传文件对象的指定工作表，回到开头以便重复读取。"""
    uploaded.seek(0)
    return eh.load_dataframe(uploaded, sheet_name=sheet_name)


def _download_xlsx(df: pd.DataFrame, filename: str, label: str = "下载 Excel"):
    """把 DataFrame 做成 Excel 下载按钮。"""
    from io import BytesIO
    buffer = BytesIO()
    df.to_excel(buffer, index=False, engine="openpyxl")
    st.download_button(label, buffer.getvalue(), file_name=filename,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _ai_button(prompt_kind: str, build_data_text, button_label: str, help_text: str):
    """
    通用 AI 总结按钮：未配置时提示去设置页；调用失败给友好信息。
    prompt_kind: 'exam' 或 'student'，决定用哪个提示词模板。
    build_data_text: 无参函数，返回注入给模型的数据文本。
    """
    if not llm_client.is_configured():
        st.info("还没配置 AI：请到「⚙️ 设置 → LLM 配置」填写 API Key 后再来。")
        return
    if st.button(button_label, help=help_text):
        import time
        from pathlib import Path
        import config
        prompt_file = {
            "exam": "exam_analysis_prompt.txt",
            "student": "student_profile_prompt.txt",
        }[prompt_kind]
        system_prompt = (Path(config.BASE_DIR) / "prompts" / prompt_file).read_text(encoding="utf-8")
        # 提示词模板里 {data} 是占位符，数据作为用户消息单独传入更安全（不做字符串拼接注入系统提示）
        system_prompt = system_prompt.replace("{data}", "").strip()
        with st.spinner("AI 正在分析，请稍候……"):
            try:
                reply = llm_client.chat(system_prompt, build_data_text(), temperature=0.5)
                st.markdown(reply)
            except llm_client.LLMConfigError as exc:
                st.warning(str(exc))
            except llm_client.LLMCallError as exc:
                st.error(str(exc))


# ---------------------------------------------------------------------------
# 标签页 1：学生管理
# ---------------------------------------------------------------------------

def tab_students():
    """学生名单：Excel 导入、手动添加、筛选搜索、编辑删除。"""
    st.subheader("学生管理")

    with st.expander("📥 从 Excel 导入学生名单", expanded=False):
        _student_import_panel()

    with st.expander("➕ 手动添加学生", expanded=False):
        _student_add_panel()

    st.divider()
    _student_list_panel()


def _student_import_panel():
    """上传名单 → 识别列 → 预览（标问题行）→ 确认入库。"""
    uploaded = st.file_uploader("选择学生名单 Excel（需含“姓名”列）",
                                type=["xlsx", "xls"], key="student_upload")
    if uploaded is None:
        st.caption("支持的列：姓名（必填）、学号、班级、性别、备注，表头叫法不同也能自动识别。")
        return

    try:
        sheets = eh.list_sheets(uploaded)
    except Exception as exc:
        st.error(f"读不了这个文件：{exc}")
        return
    sheet = st.selectbox("选择工作表", sheets, key="student_sheet") if len(sheets) > 1 else sheets[0]

    try:
        df = _read_excel_sheet(uploaded, sheet)
        detected = eh.detect_student_columns(df)
    except Exception as exc:
        st.error(f"解析表格出错：{exc}")
        return

    if detected["problems"]:
        st.error("　".join(detected["problems"]))
        return

    records = eh.extract_students(df, detected["mapping"])
    if not records:
        st.warning("表里没有有效学生行。")
        return

    preview = pd.DataFrame([{
        "Excel行号": r["row_no"], "姓名": r["name"], "学号": r["student_no"],
        "班级": r["class_name"], "性别": r["gender"], "备注": r["remark"],
        "问题": "、".join(r["problems"]),
    } for r in records])

    problem_rows = preview[preview["问题"] != ""]
    if len(problem_rows):
        st.warning(f"有 {len(problem_rows)} 行存在问题（导入时会被跳过），已在下表红色标出。")

    st.dataframe(
        preview.style.apply(
            lambda row: ["background-color:#ffe0e0"] * len(row) if row["问题"] else [""] * len(row),
            axis=1),
        width="stretch", hide_index=True)

    valid = len(records) - len(problem_rows)
    if st.button("✅ 确认导入", type="primary", key="confirm_students"):
        with SessionLocal() as session:
            result = ss.bulk_import_students(session, records)
            session.commit()
        st.success(f"导入完成：新增 {result['created']} 人，更新 {result['updated']} 人，"
                   f"跳过 {result['skipped']} 行。")
        st.cache_data.clear()
        st.rerun()


def _student_add_panel():
    """表单：手动添加单个学生。"""
    with st.form("add_student_form", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        name = c1.text_input("姓名 *")
        student_no = c2.text_input("学号")
        class_name = c3.text_input("班级")
        c4, c5 = st.columns(2)
        gender = c4.selectbox("性别", ["", "男", "女"])
        remark = c5.text_input("备注")
        submitted = st.form_submit_button("添加")
    if submitted:
        if not name.strip():
            st.error("姓名必填。")
            return
        with SessionLocal() as session:
            exists = ss.find_student(session, name.strip(), class_name.strip() or None)
            if exists:
                st.warning("这个姓名+班级已经存在，没有重复添加。")
                return
            from models.models import Student
            session.add(Student(
                name=name.strip(), student_no=student_no.strip() or None,
                class_name=class_name.strip() or None,
                gender=gender or None, remark=remark.strip() or None))
            session.commit()
        st.success(f"已添加学生：{name.strip()}")
        st.rerun()


def _student_list_panel():
    """学生列表：班级筛选、关键字搜索、编辑、删除。"""
    with SessionLocal() as session:
        classes = ss.list_classes(session)

    c1, c2 = st.columns(2)
    class_filter = c1.selectbox("按班级筛选", ["全部"] + classes)
    keyword = c2.text_input("搜索姓名或学号")

    with SessionLocal() as session:
        students = ss.list_students(
            session,
            class_name=None if class_filter == "全部" else class_filter,
            keyword=keyword.strip() or None)

        options = [f"{s.id}｜{s.name}｜{s.class_name or '未分班'}" for s in students]
        st.caption(f"共 {len(students)} 名学生")
        st.dataframe(pd.DataFrame([{
            "ID": s.id, "姓名": s.name, "学号": s.student_no or "",
            "班级": s.class_name or "", "性别": s.gender or "",
            "标签": s.tags or "", "备注": s.remark or "",
        } for s in students]), width="stretch", hide_index=True)

        if options:
            st.divider()
            st.markdown("**编辑 / 删除学生**")
            chosen = st.selectbox("选择学生", options, key="student_pick")
            sid = int(chosen.split("｜")[0])
            target = session.get(__import__("models.models", fromlist=["Student"]).Student, sid)

            with st.form("edit_student_form"):
                e1, e2, e3 = st.columns(3)
                new_name = e1.text_input("姓名", value=target.name)
                new_no = e2.text_input("学号", value=target.student_no or "")
                new_class = e3.text_input("班级", value=target.class_name or "")
                e4, e5 = st.columns(2)
                new_gender = e4.selectbox(
                    "性别", ["", "男", "女"],
                    index=["", "男", "女"].index(target.gender) if target.gender in ("男", "女") else 0)
                new_remark = e5.text_input("备注", value=target.remark or "")
                col_save, col_del = st.columns(2)
                save_clicked = col_save.form_submit_button("💾 保存修改", type="primary")
                del_clicked = col_del.form_submit_button("🗑️ 删除该学生")

            if save_clicked:
                if not new_name.strip():
                    st.error("姓名不能为空。")
                else:
                    ss.update_student(session, sid, name=new_name.strip(),
                                      student_no=new_no.strip() or None,
                                      class_name=new_class.strip() or None,
                                      gender=new_gender or None,
                                      remark=new_remark.strip() or None)
                    session.commit()
                    st.success("已保存修改。")
                    st.rerun()
            if del_clicked:
                st.session_state["confirm_delete_student"] = sid

        session.rollback()  # selectbox 阶段未提交的对象在此结束会话

    # 删除二次确认（危险操作）
    if st.session_state.get("confirm_delete_student"):
        sid = st.session_state["confirm_delete_student"]
        st.error("删除学生会一并删除该生所有成绩，且不能撤销。确定删除吗？")
        cc1, cc2 = st.columns(2)
        if cc1.button("✅ 确认删除", type="primary"):
            with SessionLocal() as session:
                ss.delete_student(session, sid)
                session.commit()
            st.session_state.pop("confirm_delete_student", None)
            st.success("已删除。")
            st.rerun()
        if cc2.button("取消"):
            st.session_state.pop("confirm_delete_student", None)
            st.rerun()

# ---------------------------------------------------------------------------
# 标签页 2：成绩管理
# ---------------------------------------------------------------------------

def tab_scores():
    """考试信息管理 + 成绩导入/手动录入/导出。"""
    st.subheader("成绩管理")

    with st.expander("🗂️ 考试管理（新建 / 修改满分 / 删除）", expanded=False):
        _exam_manage_panel()

    with st.expander("📥 从 Excel 导入成绩", expanded=False):
        _score_import_panel()

    with st.expander("⌨️ 手动录入 / 修正成绩", expanded=False):
        _score_manual_panel()

    st.divider()
    _score_export_panel()


def _exam_manage_panel():
    """新建考试、修改各科满分、删除考试。"""
    with SessionLocal() as session:
        exams = es.list_exams(session)
        exam_options = {f"{e.name}（{e.exam_date}）": e.id for e in exams}

    with st.form("create_exam_form", clear_on_submit=True):
        st.markdown("**新建考试**")
        c1, c2, c3, c4 = st.columns(4)
        name = c1.text_input("考试名称 *", placeholder="如：八年级上册期中")
        exam_date = c2.date_input("考试日期", value=date.today())
        grade = c3.text_input("年级", value="八年级")
        term = c4.text_input("学期", placeholder="如：2026 上")
        full_text = st.text_input(
            "各科满分（可留空，默认 100）",
            placeholder="格式：数学=120,语文=120,英语=120")
        if st.form_submit_button("新建考试", type="primary"):
            if not name.strip():
                st.error("考试名称必填。")
            else:
                full_scores = _parse_full_scores(full_text)
                if full_scores is None:
                    st.error("满分格式不对，请按“数学=120,语文=120”填写。")
                else:
                    with SessionLocal() as session:
                        es.create_exam(session, name, exam_date=exam_date,
                                       grade=grade.strip() or None,
                                       term=term.strip() or None,
                                       full_scores=full_scores)
                        session.commit()
                    st.success(f"已新建考试：{name.strip()}")
                    st.rerun()

    st.divider()
    if not exam_options:
        st.caption("还没有考试，先在上面新建一场。")
        return

    st.markdown("**已有考试 / 修改满分 / 删除**")
    chosen_label = st.selectbox("选择考试", list(exam_options.keys()),
                                key="exam_manage_pick")
    eid = exam_options[chosen_label]
    with SessionLocal() as session:
        exam = es.get_exam(session, eid)
        subjects = es.exam_subjects(session, eid)
        full_scores = es.full_scores_of(exam, subjects)
        st.caption(f"已录科目：{('、'.join(subjects)) or '（暂无成绩）'}")

        edit_text = ", ".join(f"{s}={full_scores[s]:g}" for s in subjects) if subjects else ""
        new_full = st.text_input("修改各科满分", value=edit_text,
                                 key="exam_full_edit")
        cc1, cc2 = st.columns(2)
        if cc1.button("💾 保存满分"):
            parsed = _parse_full_scores(new_full)
            if parsed is None:
                st.error("满分格式不对。")
            else:
                es.update_exam(session, eid, full_scores=parsed)
                session.commit()
                st.success("满分已更新。")
                st.rerun()
        if cc2.button("🗑️ 删除该考试（含全部成绩）"):
            st.session_state["confirm_delete_exam"] = eid

    if st.session_state.get("confirm_delete_exam"):
        st.error("删除考试会连带删除这场考试的全部成绩，确定吗？")
        d1, d2 = st.columns(2)
        if d1.button("✅ 确认删除", type="primary", key="exam_del_ok"):
            with SessionLocal() as session:
                es.delete_exam(session, st.session_state["confirm_delete_exam"])
                session.commit()
            st.session_state.pop("confirm_delete_exam", None)
            st.success("已删除。")
            st.rerun()
        if d2.button("取消", key="exam_del_cancel"):
            st.session_state.pop("confirm_delete_exam", None)
            st.rerun()


def _parse_full_scores(text: str):
    """解析“数学=120,语文=120”成字典；空串返回空字典，格式错返回 None。"""
    text = (text or "").strip()
    if not text:
        return {}
    result = {}
    for part in text.replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "=" not in part:
            return None
        subj, val = part.split("=", 1)
        subj = subj.strip()
        try:
            num = float(val.strip())
        except ValueError:
            return None
        if not subj or num <= 0:
            return None
        result[subj] = num
    return result


def _score_import_panel():
    """选考试 → 上传成绩 Excel → 识别科目列 → 预览标红 → 确认入库。"""
    with SessionLocal() as session:
        exams = es.list_exams(session)
    if not exams:
        st.info("请先在上方“考试管理”里新建一场考试，再导入成绩。")
        return
    options = {f"{e.name}（{e.exam_date}）": e.id for e in exams}
    chosen = st.selectbox("导入到哪场考试", list(options.keys()), key="score_import_exam")
    exam_id = options[chosen]

    uploaded = st.file_uploader("选择成绩 Excel（需含“姓名”列，其余数字列按科目识别）",
                                type=["xlsx", "xls"], key="score_upload")
    if uploaded is None:
        st.caption("数字列名就是科目，如“数学”或“数学成绩”；空格按缺考处理，不覆盖已有分数。")
        return

    try:
        sheets = eh.list_sheets(uploaded)
        sheet = st.selectbox("选择工作表", sheets, key="score_sheet") if len(sheets) > 1 else sheets[0]
        df = _read_excel_sheet(uploaded, sheet)
        detected = eh.detect_score_columns(df)
    except Exception as exc:
        st.error(f"解析出错：{exc}")
        return

    if detected["problems"]:
        st.error("　".join(detected["problems"]))
        return

    st.success(f"识别到姓名列「{detected['name_col']}」，科目："
               + "、".join(s for _, s in detected["subjects"]))
    records = eh.extract_scores(df, detected)
    if not records:
        st.warning("没有有效成绩行。")
        return

    preview = pd.DataFrame([{
        "Excel行号": r["row_no"], "姓名": r["name"], "班级": r["class_name"] or "",
        **{subj: r["scores"][subj] for _, subj in detected["subjects"]},
        "问题": "、".join(r["problems"]),
    } for r in records])

    problem_rows = preview[preview["问题"] != ""]
    if len(problem_rows):
        st.warning(f"有 {len(problem_rows)} 行提示问题（缺姓名会跳过；缺考科目不计分），已标红。")
    st.dataframe(
        preview.style.apply(
            lambda row: ["background-color:#ffe0e0"] * len(row) if row["问题"] else [""] * len(row),
            axis=1),
        width="stretch", hide_index=True)

    if st.button("✅ 确认导入成绩", type="primary", key="confirm_scores"):
        with SessionLocal() as session:
            result = es.import_scores(session, exam_id, records)
            session.commit()
        st.success(
            f"导入完成：自动新建学生 {result['created_students']} 人，"
            f"新写入成绩 {result['scores_written']} 条，"
            f"更新已有成绩 {result['scores_updated']} 条，"
            f"跳过 {result['skipped_rows']} 行。")
        st.rerun()


def _score_manual_panel():
    """选考试和班级，用可编辑表格手动录入/修正各科分数。"""
    with SessionLocal() as session:
        exams = es.list_exams(session)
        classes = ss.list_classes(session)
    if not exams:
        st.info("还没有考试。")
        return
    exam_options = {f"{e.name}（{e.exam_date}）": e.id for e in exams}
    c1, c2 = st.columns(2)
    chosen = c1.selectbox("考试", list(exam_options.keys()), key="manual_exam")
    class_filter = c2.selectbox("班级", ["全部"] + classes, key="manual_class")
    exam_id = exam_options[chosen]

    if st.button("载入录入表", key="manual_load"):
        st.session_state["manual_loaded"] = (exam_id, class_filter)

    loaded = st.session_state.get("manual_loaded")
    if not loaded or loaded[0] != exam_id:
        st.caption("选好考试和班级后点“载入录入表”。")
        return

    with SessionLocal() as session:
        subjects = es.exam_subjects(session, exam_id)
        rows = es.exam_score_rows(session, exam_id)
        if class_filter != "全部":
            rows = [r for r in rows if r["class_name"] == class_filter]
        if not subjects:
            st.warning("这场考试还没有任何科目成绩，请先导入一次成绩（哪怕只有一列）。")
            return
        editable = pd.DataFrame([{
            "学生ID": r["student_id"], "姓名": r["name"],
            **{s: r[s] for s in subjects},
        } for r in rows])

    st.caption("直接改表格里的分数，改完点下方“保存”。空格表示缺考，不会清除已有分数。")
    edited = st.data_editor(editable, num_rows="fixed", width="stretch",
                            hide_index=True, key="manual_editor")

    if st.button("💾 保存修改", type="primary", key="manual_save"):
        records = []
        for _, row in edited.iterrows():
            scores = {}
            for s in subjects:
                v = row[s]
                if pd.isna(v) or v == "":
                    scores[s] = None
                else:
                    try:
                        scores[s] = float(v)
                    except (TypeError, ValueError):
                        scores[s] = None
            records.append({"name": row["姓名"],
                            "class_name": class_filter if class_filter != "全部" else None,
                            "scores": scores})
        with SessionLocal() as session:
            result = es.import_scores(session, exam_id, records)
            session.commit()
        st.success(f"已保存：新写入 {result['scores_written']} 条，更新 {result['scores_updated']} 条。")
        st.rerun()


def _score_export_panel():
    """把某场考试的宽表（含总分、班名）导出成 Excel。"""
    with SessionLocal() as session:
        exams = es.list_exams(session)
    if not exams:
        st.caption("暂无考试数据可导出。")
        return
    options = {f"{e.name}（{e.exam_date}）": e.id for e in exams}
    c1, c2 = st.columns([2, 1])
    chosen = c1.selectbox("选择要导出的考试", list(options.keys()), key="export_exam")
    if c2.button("生成导出表", type="primary"):
        with SessionLocal() as session:
            analysis = es.analyze_exam(session, options[chosen])
        df = pd.DataFrame([{
            "班级": r["class_name"] or "", "姓名": r["name"],
            **{s: r[s] for s in analysis["subjects"]},
            "总分": r["total"], "班级总分排名": r["total_rank"],
        } for r in analysis["rows"]])
        st.dataframe(df, width="stretch", hide_index=True)
        _download_xlsx(df, f"成绩_{chosen.split('（')[0]}.xlsx", label="⬇️ 导出 Excel")
# ---------------------------------------------------------------------------
# 标签页 3：单次考试分析
# ---------------------------------------------------------------------------

def tab_exam_analysis():
    """选考试和班级，展示指标、图表、排名（含进退步）、AI 总结。"""
    st.subheader("单次考试分析")

    with SessionLocal() as session:
        exams = es.list_exams(session)
        classes = ss.list_classes(session)
    if not exams:
        st.info("还没有考试和成绩，先到「成绩管理」导入。")
        return

    exam_options = {f"{e.name}（{e.exam_date}）": e.id for e in exams}
    c1, c2, c3 = st.columns([2, 1, 1])
    chosen = c1.selectbox("选择考试", list(exam_options.keys()), key="analysis_exam")
    class_filter = c2.selectbox("班级", ["全部"] + classes, key="analysis_class")
    c3.markdown("　")  # 占位对齐
    exam_id = exam_options[chosen]

    with SessionLocal() as session:
        data = es.analyze_exam(
            session, exam_id,
            class_name=None if class_filter == "全部" else class_filter)
    if not data or not data["rows"]:
        st.warning("这场考试还没有成绩数据。")
        return

    _render_metric_cards(data)
    st.divider()
    _render_band_and_subject_charts(data)
    st.divider()
    _render_rank_table(data)
    st.divider()
    _render_ai_exam_summary(data, class_filter)


def _fmt_pct(rate):
    """0.8276 -> '82.8%'；None -> '—'。"""
    return "—" if rate is None else f"{rate * 100:.1f}%"


def _render_metric_cards(data):
    """总分维度的关键指标卡片。"""
    t = data["total_stats"]
    cols = st.columns(6)
    cols[0].metric("参考人数", t["count"])
    cols[1].metric("平均分", "—" if t["mean"] is None else f"{t['mean']:g}")
    cols[2].metric("中位数", "—" if t["median"] is None else f"{t['median']:g}")
    cols[3].metric("标准差", "—" if t["std"] is None else f"{t['std']:g}")
    cols[4].metric("及格率", _fmt_pct(t["pass_rate"]))
    cols[5].metric("优秀率", _fmt_pct(t["excellent_rate"]))
    c1, c2 = st.columns(2)
    c1.metric("最高分", "—" if t["max"] is None else f"{t['max']:g}")
    c2.metric("最低分", "—" if t["min"] is None else f"{t['min']:g}")
    st.caption(f"及格线=满分60%，优秀线=满分85%；本次总分满分 {data['total_full']:g}。")


def _render_band_and_subject_charts(data):
    """各科分数段柱状图 + 各科均分对比柱状图。"""
    subjects = data["subjects"]
    subject_stats = data["subject_stats"]

    st.markdown("**分数段分布（按科目）**")
    subj = st.selectbox("选择科目看分数段", subjects, key="band_subject")
    bands = subject_stats[subj]["bands"]
    fig = go.Figure(go.Bar(
        x=[b["label"] for b in bands], y=[b["count"] for b in bands],
        text=[b["count"] for b in bands], marker_color="#4c78a8"))
    fig.update_layout(height=320, margin=dict(l=10, r=10, t=30, b=10),
                      yaxis_title="人数", showlegend=False)
    st.plotly_chart(fig, width="stretch")

    st.markdown("**各科平均分对比**")
    means = [subject_stats[s]["mean"] for s in subjects]
    fig2 = go.Figure(go.Bar(x=subjects, y=means, text=means, marker_color="#59a14f"))
    fig2.update_layout(height=320, margin=dict(l=10, r=10, t=30, b=10),
                       yaxis_title="平均分", showlegend=False)
    st.plotly_chart(fig2, width="stretch")


def _delta_html(delta_info):
    """把进退步信息渲染成带颜色的小文本；进步绿、退步红。"""
    if not delta_info:
        return "—"
    score_delta = delta_info.get("score_delta")
    rank_delta = delta_info.get("rank_delta")
    parts = []
    if score_delta is not None:
        sign = "+" if score_delta > 0 else ""
        color = "#1a9850" if score_delta > 0 else ("#d73027" if score_delta < 0 else "#666")
        parts.append(f"<span style='color:{color}'>总分{sign}{score_delta:g}</span>")
    if rank_delta is not None and rank_delta != 0:
        sign = "↑" if rank_delta > 0 else "↓"
        color = "#1a9850" if rank_delta > 0 else "#d73027"
        parts.append(f"<span style='color:{color}'>名次{sign}{abs(rank_delta)}</span>")
    elif rank_delta == 0:
        parts.append("<span style='color:#666'>名次持平</span>")
    return "　".join(parts) if parts else "—"


def _render_rank_table(data):
    """排名表：班级筛选、搜索、进退步红绿标记。"""
    st.markdown("**班级排名表（含与上次考试的进退步）**")
    keyword = st.text_input("搜索姓名", key="rank_search")
    subjects = data["subjects"]
    deltas = data["deltas"]

    rows = data["rows"]
    if keyword.strip():
        rows = [r for r in rows if keyword.strip() in (r["name"] or "")]

    table = []
    for r in rows:
        row = {
            "班级": r["class_name"] or "", "姓名": r["name"],
            **{s: ("" if r[s] is None else f"{r[s]:g}") for s in subjects},
            "总分": "" if r["total"] is None else f"{r['total']:g}",
            "班名": r["total_rank"] if r["total_rank"] is not None else "—",
            "进退步": _delta_html(deltas.get(r["student_id"])),
        }
        table.append(row)
    df = pd.DataFrame(table)
    st.markdown(df.to_html(escape=False, index=False), unsafe_allow_html=True)
    st.caption("绿色=进步，红色=退步；只比较日期相邻的两次考试。")


def _build_exam_data_text(data, class_filter) -> str:
    """把分析结果压缩成给 AI 的纯文本数据（只含统计，不含提示词指令）。"""
    lines = [f"考试：{data['exam'].name}；范围：{class_filter}；"
             f"总分满分：{data['total_full']:g}；参考人数：{data['total_stats']['count']}"]
    t = data["total_stats"]
    lines.append(f"总分：均分{t['mean']}、中位数{t['median']}、标准差{t['std']}、"
                 f"最高{t['max']}、最低{t['min']}、"
                 f"及格率{_fmt_pct(t['pass_rate'])}、优秀率{_fmt_pct(t['excellent_rate'])}")
    for s in data["subjects"]:
        ss_stat = data["subject_stats"][s]
        band_text = "，".join(f"{b['label']}{b['count']}人" for b in ss_stat["bands"])
        lines.append(f"{s}（满分{ss_stat['full_score']:g}）：均分{ss_stat['mean']}，"
                     f"及格率{_fmt_pct(ss_stat['pass_rate'])}，"
                     f"优秀率{_fmt_pct(ss_stat['excellent_rate'])}；分数段：{band_text}")
    return "\n".join(lines)


def _render_ai_exam_summary(data, class_filter):
    """AI 考试分析文字总结 + Word 导出。"""
    st.markdown("**AI 考试分析总结**")
    data_text = _build_exam_data_text(data, class_filter)
    with st.expander("查看发给 AI 的统计数据", expanded=False):
        st.code(data_text)
    _ai_button("exam", lambda: data_text, "🤖 生成考试分析",
               help_text="依据上面的统计数据生成不超过 300 字的分析")
    _export_word_report(data, class_filter)


def _export_word_report(data, class_filter):
    """把统计结果和排名表导出成可打印的 Word 报告（不含 AI 文本，AI 文本可另行复制）。"""
    if not st.button("⬇️ 导出分析报告（Word）"):
        return
    try:
        from docx import Document
        from io import BytesIO
        import config as cfg

        doc = Document()
        doc.add_heading(f"{data['exam'].name} 成绩分析", level=1)
        doc.add_paragraph(f"范围：{class_filter}　总分满分：{data['total_full']:g}")
        t = data["total_stats"]
        doc.add_paragraph(
            f"参考 {t['count']} 人，均分 {t['mean']}，中位数 {t['median']}，"
            f"标准差 {t['std']}，最高 {t['max']}，最低 {t['min']}，"
            f"及格率 {_fmt_pct(t['pass_rate'])}，优秀率 {_fmt_pct(t['excellent_rate'])}。")

        doc.add_heading("各科情况", level=2)
        for s in data["subjects"]:
            stt = data["subject_stats"][s]
            doc.add_paragraph(
                f"{s}：均分 {stt['mean']}，及格率 {_fmt_pct(stt['pass_rate'])}，"
                f"优秀率 {_fmt_pct(stt['excellent_rate'])}。", style="List Bullet")

        doc.add_heading("班级排名", level=2)
        subjects = data["subjects"]
        table = doc.add_table(rows=1, cols=4 + len(subjects))
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        headers = ["班级", "姓名"] + subjects + ["总分", "班名"]
        for i, h in enumerate(headers):
            hdr[i].text = h
        for r in data["rows"]:
            cells = table.add_row().cells
            cells[0].text = r["class_name"] or ""
            cells[1].text = r["name"]
            for j, s in enumerate(subjects):
                cells[2 + j].text = "" if r[s] is None else f"{r[s]:g}"
            cells[2 + len(subjects)].text = "" if r["total"] is None else f"{r['total']:g}"
            cells[3 + len(subjects)].text = ("" if r["total_rank"] is None
                                             else str(r["total_rank"]))

        buffer = BytesIO()
        doc.save(buffer)
        cfg.ensure_dirs()
        st.download_button(
            "💾 保存 Word 文件", buffer.getvalue(),
            file_name=f"{data['exam'].name}_成绩分析.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        st.success("报告已生成，点上方按钮保存。")
    except Exception as exc:
        st.error(f"导出失败：{exc}")
# ---------------------------------------------------------------------------
# 标签页 4：趋势分析
# ---------------------------------------------------------------------------

def tab_trends():
    """学生个人趋势 + 班级整体趋势，可选最近 3/5 次/全部考试。"""
    st.subheader("趋势分析")

    with SessionLocal() as session:
        exams = es.list_exams(session)
        classes = ss.list_classes(session)
    if len(exams) < 1:
        st.info("还没有考试数据。")
        return

    c1, c2 = st.columns(2)
    scope = c1.selectbox("显示范围", ["最近 3 次", "最近 5 次", "全部"], index=2)
    class_filter = c2.selectbox("班级", ["全部"] + classes, key="trend_class")
    limit = {"最近 3 次": 3, "最近 5 次": 5, "全部": None}[scope]

    with st.expander("📈 班级整体趋势", expanded=True):
        _class_trend_chart(class_filter, limit)

    st.divider()
    with st.expander("🧑‍🎓 学生个人趋势", expanded=True):
        _student_trend_panel(class_filter)


def _apply_limit(names, limit):
    """按最后 limit 场截断（None 表示全部）。"""
    if limit and len(names) > limit:
        return names[-limit:]
    return names


def _class_trend_chart(class_filter, limit):
    """历次考试各科均分与总分均分折线。"""
    with SessionLocal() as session:
        trend = es.class_trend(session,
                               None if class_filter == "全部" else class_filter)
    exam_names = _apply_limit(trend["exams"], limit)
    cut = len(exam_names)
    start = len(trend["exams"]) - cut

    fig = go.Figure()
    for subject in trend["subjects"]:
        fig.add_trace(go.Scatter(
            x=exam_names, y=trend["series"][subject][start:],
            mode="lines+markers", name=subject))
    fig.add_trace(go.Scatter(
        x=exam_names, y=trend["series"]["total"][start:],
        mode="lines+markers", name="总分",
        line=dict(width=4, dash="dot")))
    fig.update_layout(height=420, margin=dict(l=10, r=10, t=30, b=10),
                      yaxis_title="平均分", hovermode="x unified")
    if exam_names:
        st.plotly_chart(fig, width="stretch")
    else:
        st.caption("所选范围内没有数据。")


def _student_trend_panel(class_filter):
    """选学生 → 历次总分折线 + 单科切换折线。"""
    with SessionLocal() as session:
        students = ss.list_students(
            session,
            class_name=None if class_filter == "全部" else class_filter)
    if not students:
        st.caption("没有学生。")
        return
    options = {f"{s.name}｜{s.class_name or '未分班'}": s.id for s in students}
    chosen = st.selectbox("选择学生", list(options.keys()), key="trend_student")
    sid = options[chosen]

    with SessionLocal() as session:
        history = es.student_scores_over_time(session, sid)
    if not history:
        st.caption("该生还没有成绩。")
        return

    exam_names = [h["exam_name"] for h in history]
    subjects = [k for k in history[0].keys()
                if k not in ("exam_id", "exam_name", "exam_date", "total")]

    # 总分折线
    fig = go.Figure(go.Scatter(
        x=exam_names, y=[h["total"] for h in history],
        mode="lines+markers+text",
        text=[h["total"] for h in history], textposition="top center",
        line=dict(width=3), name="总分"))
    fig.update_layout(height=360, title="历次总分",
                      margin=dict(l=10, r=10, t=40, b=10), yaxis_title="总分")
    st.plotly_chart(fig, width="stretch")

    # 单科折线
    subject = st.selectbox("切换单科趋势", subjects, key="trend_single_subject")
    fig2 = go.Figure(go.Scatter(
        x=exam_names, y=[h.get(subject) for h in history],
        mode="lines+markers+text",
        text=[h.get(subject) for h in history], textposition="top center",
        line=dict(width=3, color="#e15759"), name=subject))
    fig2.update_layout(height=320, title=f"{subject}历次成绩",
                       margin=dict(l=10, r=10, t=40, b=10), yaxis_title="分数")
    st.plotly_chart(fig2, width="stretch")


# ---------------------------------------------------------------------------
# 标签页 5：学生画像
# ---------------------------------------------------------------------------

def tab_profile():
    """学生信息卡、各科雷达图、自动标签、趋势、AI 学情分析。"""
    st.subheader("学生画像")
    with SessionLocal() as session:
        students = ss.list_students(session)
    if not students:
        st.info("还没有学生，先到「学生管理」导入名单。")
        return

    options = {f"{s.name}｜{s.class_name or '未分班'}": s.id for s in students}
    chosen = st.selectbox("选择学生", list(options.keys()), key="profile_student")
    sid = options[chosen]

    with SessionLocal() as session:
        student = session.get(__import__("models.models", fromlist=["Student"]).Student, sid)
        history = es.student_scores_over_time(session, sid)

    if not history:
        st.info(f"{student.name} 还没有成绩数据。")
        return

    # 每科按历次"得分/满分"算平均得分率（雷达图用），并收集趋势得分率
    subject_rates_all = {}
    subject_latest_rate = {}
    # 逐次考试取各科满分，把原始分换算成得分率
    with SessionLocal() as session:
        for h in history:
            exam = es.get_exam(session, h["exam_id"])
            subs = es.exam_subjects(session, h["exam_id"])
            full = es.full_scores_of(exam, subs)
            for s in subs:
                if h.get(s) is not None and full.get(s):
                    subject_rates_all.setdefault(s, []).append(h[s] / full[s])

    subject_avg_rate = {s: sum(v) / len(v) for s, v in subject_rates_all.items()}
    subjects = sorted(subject_avg_rate.keys())

    # 信息卡 + 自动标签
    c1, c2 = st.columns([1, 2])
    with c1:
        st.markdown(f"#### {student.name}")
        st.write(f"**班级：** {student.class_name or '未分班'}")
        st.write(f"**学号：** {student.student_no or '—'}")
        st.write(f"**性别：** {student.gender or '—'}")
        exams_count = len(history)
        st.write(f"**参考次数：** {exams_count}")
        latest = history[-1]
        st.write(f"**最近一次：** {latest['exam_name']}，总分 {latest['total']:g}")

    # 自动标签
    overall_rate = (sum(subject_avg_rate.values()) / len(subject_avg_rate)
                    if subject_avg_rate else None)
    level = stats_mod.level_tag(overall_rate)
    # 趋势标签用总分得分率序列（需要总分满分，按各科满分之和）
    total_rates = []
    with SessionLocal() as session:
        for h in history:
            exam = es.get_exam(session, h["exam_id"])
            subs = es.exam_subjects(session, h["exam_id"])
            full_total = sum(es.full_scores_of(exam, subs).values())
            if h["total"] is not None and full_total:
                total_rates.append(h["total"] / full_total)
    trend = stats_mod.trend_label(total_rates)
    strengths = stats_mod.subject_strength(subject_avg_rate)
    strong = [s for s, t in strengths.items() if t == "强项"]
    weak = [s for s, t in strengths.items() if t == "弱项"]

    with c2:
        tag_cols = st.columns(3)
        tag_cols[0].metric("成绩层次", level)
        tag_cols[1].metric("趋势", trend)
        tag_cols[2].metric("参考", f"{len(total_rates)} 次")
        st.write("**强项科目：** " + ("、".join(strong) if strong else "暂无明显强项"))
        st.write("**薄弱科目：** " + ("、".join(weak) if weak else "暂无明显弱项"))
        if student.tags:
            st.write("**老师标签：** " + student.tags)

    st.divider()
    # 雷达图：按得分率画，满分不同也可比
    st.markdown("**各科实力雷达图（平均得分率 %）**")
    fig = go.Figure(go.Scatterpolar(
        r=[round(subject_avg_rate[s] * 100, 1) for s in subjects] + [round(subject_avg_rate[subjects[0]] * 100, 1)],
        theta=subjects + [subjects[0]],
        fill="toself", name="平均得分率"))
    fig.update_layout(polar=dict(radialaxis=dict(range=[0, 100])),
                      height=420, margin=dict(l=30, r=30, t=30, b=10))
    st.plotly_chart(fig, width="stretch")

    # 历次成绩表
    st.markdown("**历次成绩**")
    table = pd.DataFrame([{
        "考试": h["exam_name"],
        **{s: ("" if h.get(s) is None else f"{h[s]:g}") for s in subjects},
        "总分": "" if h["total"] is None else f"{h['total']:g}",
    } for h in history])
    st.dataframe(table, width="stretch", hide_index=True)

    st.divider()
    st.markdown("**AI 学情分析**")
    data_text = _build_student_data_text(
        student, history, subjects, subject_avg_rate, level, trend, strong, weak)
    with st.expander("查看发给 AI 的学生数据", expanded=False):
        st.code(data_text)
    _ai_button("student", lambda: data_text, "🤖 生成学情分析",
               help_text="结合历次成绩和标签生成 150 字内的个性化分析")


def _build_student_data_text(student, history, subjects, subject_avg_rate,
                             level, trend, strong, weak) -> str:
    """组装给 AI 的学生数据文本。"""
    lines = [f"学生：{student.name}；班级：{student.class_name or '未分班'}；"
             f"参考 {len(history)} 次。",
             f"系统标签：成绩层次「{level}」，趋势「{trend}」，"
             f"强项：{('、'.join(strong)) or '无'}，弱项：{('、'.join(weak)) or '无'}。",
             "各科平均得分率："
             + "，".join(f"{s}{subject_avg_rate[s]*100:.1f}%" for s in subjects) + "。",
             "历次成绩："]
    for h in history:
        parts = [f"{s}{h.get(s):g}" for s in subjects if h.get(s) is not None]
        lines.append(f"- {h['exam_name']}：" + "，".join(parts)
                     + f"，总分{h['total']:g}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 页面入口
# ---------------------------------------------------------------------------

def show():
    """学情工作台：顶部 5 个标签页路由。"""
    st.title("📊 学情工作台")
    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        ["学生管理", "成绩管理", "考试分析", "趋势分析", "学生画像"])
    with tab1:
        tab_students()
    with tab2:
        tab_scores()
    with tab3:
        tab_exam_analysis()
    with tab4:
        tab_trends()
    with tab5:
        tab_profile()