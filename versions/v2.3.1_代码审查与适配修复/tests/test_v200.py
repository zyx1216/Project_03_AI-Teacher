# -*- coding: utf-8 -*-
"""v2.0.0 功能深化批次1测试。

覆盖：题目质量审核、相似度查重、家校沟通报告、分层作业、教案版本 diff、PPT 控制。
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from models.models import (
    Base, Exam, Homework, HomeworkQuestion, Question, Score, Student)
from utils import agent_advisor
from utils import lesson_service
from utils import parent_report_service
from utils import ppt_generator
from utils import question_quality_service
from utils import question_service
from utils import question_similarity_service


# ---------------------------------------------------------------------------
# 基础 fixture
# ---------------------------------------------------------------------------


@pytest.fixture()
def session(tmp_path):
    db_path = tmp_path / "v200.db"
    engine = create_engine(
        f"sqlite:///{db_path.as_posix()}", future=True)
    Base.metadata.create_all(engine)
    sess = sessionmaker(bind=engine, future=True)()
    yield sess
    sess.close()
    engine.dispose()


def _question(content="题目", answer="答案", qtype="solution",
               difficulty=2, subject="数学", grade="初二",
               kps=None, status="approved"):
    return {
        "content": content, "question_type": qtype,
        "difficulty": difficulty,
        "knowledge_points": question_service.normalize_knowledge_points(
            kps or ["知识点"]),
        "answer": answer, "analysis": "解析", "error_points": "",
        "verify": None,
        "_subject": subject, "_grade": grade, "_status": status,
    }


def _create_q(session, **kwargs):
    subject = kwargs.pop("_subject", "数学")
    grade = kwargs.pop("_grade", "初二")
    status = kwargs.pop("_status", "approved")
    return question_service.create_question(
        session, kwargs, source="manual", status=status,
        subject=subject, grade=grade)


def _seed_exam(session, class_name="一班", scores=None):
    students = []
    exam = Exam(
        name="月考", exam_date=date(2026, 9, 1), grade="八年级",
        full_scores='{"数学": 100}')
    session.add(exam)
    session.flush()
    scores = scores or [100, 80, 70, 60, 50]
    for idx, value in enumerate(scores, start=1):
        student = Student(name=f"学生{idx}", class_name=class_name)
        session.add(student)
        session.flush()
        session.add(Score(
            exam_id=exam.id, student_id=student.id,
            subject="数学", score=float(value)))
        students.append(student)
    session.flush()
    return exam, students


def _seed_question_pool(session):
    """造 1-3 难度各 5 道已审核题，够三份分层作业各取 5 道。"""
    for difficulty in (1, 2, 3):
        for qtype in ("choice", "choice", "fill", "fill", "solution"):
            _create_q(session, **_question(
                content=f"{qtype}-{difficulty}", answer="标准答案",
                qtype=qtype, difficulty=difficulty,
                kps=["数学基础"]))
    session.flush()


# ---------------------------------------------------------------------------
# 题目质量审核
# ---------------------------------------------------------------------------


def test_quality_rule_audit_finds_hard_errors():
    q = SimpleNamespace(
        content="短", answer="", question_type="solution",
        difficulty=5, knowledge_points="[]", analysis="")
    result = question_quality_service.audit_question(q)
    assert result["passed"] is False
    assert result["score"] < 60
    assert "答案为空" in result["issues"]
    assert "未标注知识点" in result["issues"]


def test_quality_llm_audit_and_repair():
    q = SimpleNamespace(
        content="1+1等于几？", answer="2", question_type="solution",
        difficulty=1, knowledge_points='["加减法"]', analysis="")

    def fake_chat(system, user):
        if "修复" in system:
            return (
                '{"content":"1+1等于几？","question_type":"solution",'
                '"difficulty":1,"knowledge_points":["加减法"],'
                '"answer":"2","analysis":"1加1等于2"}')
        return (
            '{"passed":true,"score":92,"issues":[],'
            '"suggestions":["补充解析"]}')

    result = question_quality_service.audit_question(q, fake_chat)
    assert result["passed"] is True
    assert result["score"] == 92
    repaired = question_quality_service.repair_question(
        q, result, fake_chat)
    assert repaired is not None
    assert repaired["analysis"] == "1加1等于2"


def test_quality_batch_keeps_going_when_one_item_bad():
    good = SimpleNamespace(
        content="这是一道完整的题目", answer="标准答案",
        difficulty=2, knowledge_points='["知识点"]', analysis="解析")
    bad = object()
    result = question_quality_service.batch_audit([good, bad])
    assert len(result) == 2
    assert result[0]["score"] >= 80
    assert result[1]["score"] == 0


# ---------------------------------------------------------------------------
# 题目相似度查重
# ---------------------------------------------------------------------------


def test_similar_find_duplicate_and_ignore_different(session):
    one = _create_q(session, **_question(
        content="计算 1+1 的结果", kps=["加减法"]))
    two = _create_q(session, **_question(
        content="计算 1+1 的结果？", kps=["加减法"]))
    different = _create_q(session, **_question(
        content="阅读短文回答作者情感", answer="略",
        qtype="solution", kps=["阅读理解"]))

    matches = question_similarity_service.find_similar(
        session, one, threshold=0.75)
    assert matches
    assert matches[0]["question_id"] == two.id
    assert matches[0]["score"] >= 0.75

    low_matches = question_similarity_service.find_similar(
        session, one, pool=[different], threshold=0.75)
    assert low_matches == []


def test_duplicate_on_save_and_empty_pool(session):
    old = _create_q(session, **_question(
        content="计算 2+2", kps=["加法"]))
    item = {
        "content": "计算 2+2", "question_type": "solution",
        "difficulty": 2, "knowledge_points": ["加法"],
        "answer": "4", "subject": "数学"}
    match = question_similarity_service.duplicate_on_save(session, item)
    assert match is not None
    assert match["question_id"] == old.id
    assert question_similarity_service.find_similar(
        session, item, pool=[], threshold=0.8) == []


# ---------------------------------------------------------------------------
# 家校沟通报告
# ---------------------------------------------------------------------------


def test_parent_report_aggregates_real_data(session):
    exam, students = _seed_exam(session)
    report = parent_report_service.generate_parent_report(
        session, students[0].id, exam_id=exam.id)
    assert report["facts"]["exams"]
    assert "学生1" in report["facts"]["student_name"]
    assert report["strengths"]
    pdf = parent_report_service.export_parent_report_pdf(report)
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 1000


def test_parent_report_data_notice_without_scores(session):
    student = Student(name="无成绩学生", class_name="一班")
    session.add(student)
    session.flush()
    report = parent_report_service.generate_parent_report(
        session, student.id)
    assert report["data_notice"] == (
        "需要更多成绩数据，本报告基于有限信息生成。")


# ---------------------------------------------------------------------------
# 分层作业
# ---------------------------------------------------------------------------


def test_generate_tiered_homework(session):
    _seed_exam(session)
    _seed_question_pool(session)
    result = agent_advisor.generate_tiered_homework(
        session, "一班", "数学", questions_per_layer=5)
    assert set(result["tiers"]) == {"A", "B", "C"}
    assert len(result["tiers"]["A"]) == 1
    assert len(result["tiers"]["B"]) == 3
    assert len(result["tiers"]["C"]) == 1
    assert len(result["homework_ids"]) == 3

    expected_diff = {"A": 3, "B": 2, "C": 1}
    for tier, homework_id in result["homework_ids"].items():
        pairs = (session.query(HomeworkQuestion, Question)
                 .join(Question, HomeworkQuestion.question_id == Question.id)
                 .filter(HomeworkQuestion.homework_id == homework_id).all())
        assert len(pairs) == 5
        assert {q.difficulty for _link, q in pairs} == {expected_diff[tier]}


def test_generate_tiered_homework_without_scores(session):
    with pytest.raises(ValueError, match="分层"):
        agent_advisor.generate_tiered_homework(
            session, "空班", "数学")


# ---------------------------------------------------------------------------
# 教案版本 diff
# ---------------------------------------------------------------------------


def _plan_data():
    plan = lesson_service.empty_plan()
    plan["subject"] = "数学"
    plan["objectives"]["knowledge"] = "理解概念"
    plan["key_points"] = "概念理解"
    plan["process"] = [
        {"stage": "导入", "minutes": 5, "content": "复习旧知"},
        {"stage": "新授", "minutes": 20, "content": "讲解概念"},
        {"stage": "巩固练习", "minutes": 12, "content": "完成练习"},
        {"stage": "课堂小结", "minutes": 5, "content": "总结本课"},
        {"stage": "作业布置", "minutes": 3, "content": "课后习题"},
    ]
    return plan


def test_compare_plan_versions_added_step(session):
    first = _plan_data()
    lesson = lesson_service.save_plan(
        session, "测试课", first, grade="初二",
        chapter="第一章", subject="数学")
    second = _plan_data()
    second["objectives"]["knowledge"] = "理解并应用概念"
    second["process"][1]["content"] = "讲解概念\n增加例题演示"
    lesson_service.save_plan(
        session, "测试课", second, grade="初二",
        chapter="第一章", subject="数学", plan_id=lesson.id)
    versions = lesson_service.list_plan_versions(session, lesson.id)
    assert len(versions) >= 2
    text = lesson_service.compare_plan_versions(
        session, versions[1].id, versions[0].id)
    assert "【修改】" in text
    assert "教学目标·知识" in text


def test_compare_same_version_no_diff(session):
    lesson = lesson_service.save_plan(
        session, "测试课", _plan_data(), grade="初二",
        subject="数学")
    version = lesson_service.list_plan_versions(session, lesson.id)[0]
    text = lesson_service.compare_plan_versions(
        session, version.id, version.id)
    assert text == "两个版本内容无差异。"


# ---------------------------------------------------------------------------
# PPT 内容控制
# ---------------------------------------------------------------------------


def _ppt_lesson():
    return SimpleNamespace(title="测试课", grade="初二", chapter="第一章")


def _ppt_plan(*, include_example=True, stages=None):
    plan = lesson_service.empty_plan()
    plan["objectives"] = {
        "knowledge": "理解知识", "process": "合作探究",
        "emotion": "主动思考"}
    stages = stages or ["导入", "新授", "巩固练习", "课堂小结", "作业布置"]
    contents = {
        "导入": "复习旧知",
        "新授": "讲解概念\n例题演示" if include_example else "讲解概念",
        "巩固练习": "课堂练习",
        "课堂小结": "总结知识",
        "作业布置": "完成习题",
    }
    plan["process"] = [
        {"stage": stage, "minutes": 5, "content": contents[stage]}
        for stage in stages]
    return plan


def _slide_texts(data: bytes):
    import io
    prs = Presentation(io.BytesIO(data))
    return [" ".join(shape.text for shape in slide.shapes
                      if hasattr(shape, "text"))
            for slide in prs.slides]


def test_ppt_controls_examples_practice_and_pages():
    full_texts = _slide_texts(ppt_generator.generate_ppt(
        _ppt_lesson(), _ppt_plan()))
    assert any("测试课" in text for text in full_texts)
    assert any("完成习题" in text for text in full_texts)

    no_practice = _slide_texts(ppt_generator.generate_ppt(
        _ppt_lesson(), _ppt_plan(), include_practice=False))
    assert all("例题与练习" not in text for text in no_practice)

    no_example = _slide_texts(ppt_generator.generate_ppt(
        _ppt_lesson(), _ppt_plan(include_example=False),
        include_examples=False))
    assert all("例题演示" not in text for text in no_example)

    limited = _slide_texts(ppt_generator.generate_ppt(
        _ppt_lesson(),
        _ppt_plan(stages=["新授"]), target_pages=3))
    assert len(limited) == 3


# ---------------------------------------------------------------------------
# AppTest
# ---------------------------------------------------------------------------


def _app_engine(tmp_path):
    db_file = tmp_path / "app_v200.db"
    engine = create_engine(f"sqlite:///{db_file.as_posix()}", future=True)
    Base.metadata.create_all(engine)
    return db_file, engine


def test_v200_app_question_quality_and_ppt(tmp_path):
    from streamlit.testing.v1 import AppTest
    from tests.test_app_smoke import _isolated_app_code

    db_file, engine = _app_engine(tmp_path)
    sess = sessionmaker(bind=engine)()
    question_service.create_question(
        sess,
        {"content": "AppTest质量题：完整题干", "question_type": "solution",
         "difficulty": 2, "knowledge_points": '["App知识点"]',
         "answer": "标准答案", "analysis": "标准解析",
         "error_points": "", "verify": None},
        source="manual", status="pending", subject="数学", grade="初二")
    lesson_service.save_plan(
        sess, "App测试课", _ppt_plan(stages=["导入", "新授", "作业布置"]),
        grade="初二", chapter="第一章", subject="数学")
    sess.commit()
    sess.close()

    (tmp_path / "feature.json").write_text(
        '{"materials_subject": "数学", "question_gen_subject": "数学", '
        '"question_bank_subject": "数学", "wrong_book_subject": "数学", '
        '"lesson_plan_subject": "数学"}', encoding="utf-8")
    at = AppTest.from_string(
        _isolated_app_code(db_file, tmp_path / "feature.json"),
        default_timeout=30)
    at.run()

    at.session_state["app_top_page"] = "📚 备课"
    at.session_state["lesson_plan_tab"] = "题库管理"
    at.run()
    next(b for b in at.button if b.key == "quality_audit_1").click().run()
    assert not at.exception, [str(x) for x in at.exception]
    assert any(x.key == "quality_repair_1" for x in at.button)

    at.session_state["lesson_plan_tab"] = "PPT 生成"
    at.run()
    next(b for b in at.button if "生成 PPT" in b.label).click().run()
    assert not at.exception, [str(x) for x in at.exception]
    engine.dispose()


def test_v200_app_parent_report_and_tiered_homework(tmp_path):
    from streamlit.testing.v1 import AppTest
    from tests.test_app_smoke import _isolated_app_code

    db_file, engine = _app_engine(tmp_path)
    sess = sessionmaker(bind=engine)()
    _seed_exam(sess)
    _seed_question_pool(sess)
    sess.commit()
    sess.close()

    (tmp_path / "feature.json").write_text(
        '{"materials_subject": "数学", "question_gen_subject": "数学", '
        '"question_bank_subject": "数学", "wrong_book_subject": "数学", '
        '"lesson_plan_subject": "数学"}', encoding="utf-8")
    at = AppTest.from_string(
        _isolated_app_code(db_file, tmp_path / "feature.json"),
        default_timeout=30)
    at.run()
    at.session_state["app_top_page"] = "📊 学情"
    at.run()

    radio = next(x for x in at.sidebar.radio if x.key == "analysis_tab")
    radio.set_value("学生画像").run()
    next(b for b in at.button if b.label == "📨 生成家校报告").click().run()
    assert not at.exception, [str(x) for x in at.exception]
    assert any(x.key == "parent_pdf_1" for x in at.download_button)

    at.session_state["analysis_tab"] = "趋势分析"
    at.run()
    next(b for b in at.button
         if b.key == "generate_tiered_homework").click().run()
    assert not at.exception, [str(x) for x in at.exception]
    check_session = sessionmaker(bind=engine)()
    assert check_session.query(Homework).count() == 3
    check_session.close()
    engine.dispose()



