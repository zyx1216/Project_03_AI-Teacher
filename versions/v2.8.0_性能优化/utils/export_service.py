# -*- coding: utf-8 -*-
"""多格式导出服务（v1.9.0）。

- lesson_plan_pdf：用 PyMuPDF Story 直出教案 PDF（不依赖 Word 转 PDF）；
- homework_latex：作业导出为 ctexart LaTeX 源码；
- analysis_report_ppt：成绩分析导出为 PPT（python-pptx）。

原有 Word / Excel / ZIP 导出保留在各自服务里，这里不重复。
"""

from __future__ import annotations

import html
import io
from pathlib import Path


# ---------------------------------------------------------------------------
# 教案 PDF（PyMuPDF Story）
# ---------------------------------------------------------------------------

def _pdf_archive():
    """微软雅黑存在时嵌入；否则用 PyMuPDF 内置中文字体。"""
    import pymupdf
    font = Path(r"C:\Windows\Fonts\msyh.ttc")
    if not font.exists():
        return None, "body { font-family: china-s, sans-serif; }"
    archive = pymupdf.Archive()
    archive.add(font.read_bytes(), "yh.ttc")
    return archive, (
        "@font-face { font-family: yh; src: url(yh.ttc); }"
        "body { font-family: yh, china-s, sans-serif; }")


def _esc(value) -> str:
    return html.escape(str(value or ""), quote=False).replace("\n", "<br>")


def lesson_plan_html(lesson, plan: dict) -> str:
    """教案 PDF 的 HTML；分页交给 Story。"""
    objectives = plan.get("objectives", {})
    parts = [
        "<html><head><meta charset='utf-8'><style>",
        """
        @page { margin: 48px 46px; }
        body { font-size: 10.5pt; line-height: 1.6; color: #222; }
        h1 { color: #1f4e79; font-size: 22pt; margin-bottom: 4px; }
        h2 { color: #1f4e79; font-size: 14pt; margin: 18px 0 8px;
             border-bottom: 1px solid #9eadcc; padding-bottom: 3px;
             page-break-after: avoid; }
        .meta { color: #555; margin-bottom: 16px; }
        .block { page-break-inside: avoid; margin: 6px 0; }
        .stage { page-break-inside: avoid; border: 1px solid #c8d4e6;
                 border-radius: 6px; padding: 8px 10px; margin: 9px 0; }
        .stage-t { color: #1f4e79; font-weight: bold; }
        """,
        "</style></head><body>",
        f"<h1>{_esc(lesson.title)}</h1>",
        f"<div class='meta'>学科：{_esc(plan.get('subject'))}　"
        f"年级：{_esc(lesson.grade or '未指定')}　"
        f"章节：{_esc(lesson.chapter or '未指定')}</div>",
        "<h2>教学目标</h2>",
        f"<div class='block'>知识与技能：{_esc(objectives.get('knowledge'))}</div>",
        f"<div class='block'>过程与方法：{_esc(objectives.get('process'))}</div>",
        f"<div class='block'>情感态度：{_esc(objectives.get('emotion'))}</div>",
        f"<h2>教学重点</h2>{_esc(plan.get('key_points'))}",
        f"<h2>教学难点</h2>{_esc(plan.get('difficult_points'))}",
        "<h2>教学过程</h2>",
    ]
    for step in plan.get("process", []):
        parts.append(
            "<div class='stage'>"
            f"<div class='stage-t'>{_esc(step.get('stage'))}"
            f"（{_esc(step.get('minutes'))}分钟）</div>"
            f"<div>{_esc(step.get('content'))}</div></div>")
    assignment = next(
        (s.get("content") for s in plan.get("process", [])
         if s.get("stage") == "作业布置"), "")
    parts.append(f"<h2>作业布置</h2>{_esc(assignment or '见课堂安排')}")
    if plan.get("board_design"):
        parts.append(f"<h2>板书设计</h2>{_esc(plan.get('board_design'))}")
    parts.append("</body></html>")
    return "".join(parts)


def lesson_plan_pdf(lesson, plan: dict) -> bytes:
    """教案导出为 PDF（PyMuPDF Story 直出，支持中文）。"""
    import pymupdf
    archive, font_css = _pdf_archive()
    page_html = lesson_plan_html(lesson, plan).replace(
        "</style>", font_css + "</style>", 1)
    story = pymupdf.Story(html=page_html, archive=archive)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    mediabox = pymupdf.paper_rect("a4")
    while True:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(mediabox + (46, 48, -46, -48))
        story.draw(dev)
        writer.end_page()
        if not more:
            break
    writer.close()
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 作业 LaTeX
# ---------------------------------------------------------------------------

def _latex_text(value) -> str:
    """普通文本最小转义；保留题干里已有的 $...$ 公式。"""
    text = str(value or "")
    # 按公式段切分，只转义非公式部分。
    parts = []
    cursor = 0
    in_math = False
    i = 0
    while i < len(text):
        if text[i] == "$":
            seg = text[cursor:i]
            parts.append(_latex_escape(seg) if not in_math else seg)
            parts.append("$")
            in_math = not in_math
            cursor = i + 1
        i += 1
    tail = text[cursor:]
    parts.append(_latex_escape(tail) if not in_math else tail)
    return "".join(parts).replace("\n", "\\\\\n")


def _latex_escape(seg: str) -> str:
    for ch in ("\\", "&", "%", "#", "_", "{", "}"):
        seg = seg.replace(ch, "\\" + ch)
    return seg


def homework_latex(homework, pairs) -> str:
    """作业导出为 ctexart LaTeX 源码。pairs 为 [(关联行, 题目)]。"""
    lines = [
        "\\documentclass{ctexart}",
        "\\usepackage[a4paper,margin=2.2cm]{geometry}",
        "\\usepackage{amsmath,amssymb}",
        "\\begin{document}",
        f"\\section*{{{_latex_text(homework.name)}}}",
    ]
    current_type = None
    type_names = {"choice": "一、选择题", "fill": "二、填空题",
                  "judge": "三、判断题", "solution": "四、解答题"}
    for index, (_link, q) in enumerate(pairs, start=1):
        if q.question_type != current_type:
            current_type = q.question_type
            lines.append(f"\\subsection*{{{type_names.get(current_type, '题目')}}}")
        lines.append(f"{index}. {_latex_text(q.content)}")
        if q.answer:
            lines.append("")
    lines.append("\\end{document}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 成绩分析 PPT
# ---------------------------------------------------------------------------

def analysis_report_ppt(session, source_id: int, source_type: str) -> bytes:
    """成绩分析导出 PPT。

    source_type="exam" 走 exam_service.analyze_exam；
    source_type="homework" 走 homework_score_service.analyze_homework。
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title_slide(title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle

    def add_bullets(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0] if bullets else ""
        for text in bullets[1:]:
            frame.add_paragraph().text = text
        return slide

    def add_table(title, headers, rows):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        rows_n = len(rows) + 1
        cols_n = len(headers)
        table_shape = slide.shapes.add_table(
            rows_n, cols_n, Inches(0.6), Inches(1.4),
            Inches(12.1), Inches(min(5.2, 0.45 * rows_n)))
        table = table_shape.table
        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = str(h)
            cell.text_frame.paragraphs[0].font.bold = True
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                table.cell(i, j).text = str(value)

    if source_type == "exam":
        from utils import exam_service
        data = exam_service.analyze_exam(session, source_id)
        exam = data["exam"]
        add_title_slide(
            f"考试分析：{exam.name}",
            f"考试日期：{exam.exam_date}")
        stats = data["total_stats"]
        add_bullets("总体指标", [
            f"参考人数：{stats.get('count')}",
            f"平均分：{stats.get('mean')}",
            f"最高分：{stats.get('max')}",
            f"最低分：{stats.get('min')}",
            f"及格率：{stats.get('pass_rate')}%",
            f"优秀率：{stats.get('excellent_rate')}%",
        ])
        subject_rows = []
        for subject in data["subjects"]:
            s = data["subject_stats"][subject]
            subject_rows.append([
                subject, s.get("count"), s.get("mean"),
                s.get("pass_rate"), s.get("excellent_rate")])
        add_table("各科统计",
                  ["科目", "人数", "平均分", "及格率", "优秀率"],
                  subject_rows)
        add_bullets("结论与建议", [
            "结合各科及格率、优秀率确定后续教学重点。",
            "关注进退步明显的学生，分层布置巩固练习。"])
    else:
        from utils import homework_score_service as hscore
        data = hscore.analyze_homework(session, source_id)
        name = data.get("name", f"作业{source_id}")
        add_title_slide(f"作业分析：{name}", "学业测评")
        stats = data.get("stats", {})
        add_bullets("总体指标", [
            f"提交人数：{stats.get('submitted_count')}",
            f"平均分：{stats.get('mean')}",
            f"最高分：{stats.get('max')}",
            f"最低分：{stats.get('min')}",
        ])
        add_bullets("结论与建议", [
            "依据错题知识点安排专项复习。",
            "对未提交学生及时跟进。"])

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
