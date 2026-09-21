# -*- coding: utf-8 -*-
"""
成绩文档解析层（纯函数，不依赖 Streamlit、不碰数据库）。

把 Word（python-docx 真表格）和 PDF（PyMuPDF 线框表格 find_tables）里的
成绩表提取成 DataFrame，再交给 excel_handler 复用现有的姓名/科目识别规则。

约定：
- Word 取第一个“表头含姓名列且至少有一个数字列”的表格；
- PDF 只认规范线框表格，多页同一张表纵向合并；识别不出返回 None，由页面友好提示；
- 单元格空白一律转 None，分数保持原始字符串/数字，由后续 _to_score 归一。
"""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd

from utils import excel_handler as eh

# 表头里“姓名”的各种写法（规范化后比较），复用成绩表的姓名别名
_NAME_TOKENS = {eh.normalize_text(a) for a in eh.SCORE_META_ALIASES["name"]}


def _clean(value) -> str:
    """单元格转去空白字符串；None/NaN 转空串。"""
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value).strip()


def _header_has_name(header) -> bool:
    """表头行里是否出现姓名列。"""
    return any(eh.normalize_text(c) in _NAME_TOKENS for c in header)


def _row_has_number(values) -> bool:
    """一行里是否至少有一个能当成绩的数字单元格。"""
    return any(eh._to_score(v) is not None for v in values)


def _grid_to_dataframe(rows: list[list[str]]) -> pd.DataFrame | None:
    """二维字符串网格 -> DataFrame；不满足成绩表特征时返回 None。"""
    rows = [list(r) for r in rows if r is not None]
    if len(rows) < 2:
        return None
    header = [_clean(c) for c in rows[0]]
    if not _header_has_name(header):
        return None

    ncols = len(header)
    # 空表头给占位名，重复表头加序号，避免 DataFrame 列冲突
    seen = {}
    columns = []
    for i, name in enumerate(header):
        name = name or f"未命名列{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        columns.append(name)

    data = []
    for raw in rows[1:]:
        raw = [_clean(c) for c in raw]
        if not any(raw):  # 整行空白跳过
            continue
        if len(raw) < ncols:
            raw = raw + [""] * (ncols - len(raw))
        elif len(raw) > ncols:
            raw = raw[:ncols]
        data.append([(v if v != "" else None) for v in raw])

    if not data:
        return None
    if not any(_row_has_number([v for v in row if v is not None]) for row in data):
        return None
    return pd.DataFrame(data, columns=columns, dtype=object)


def _is_score_grid(rows: list[list[str]]) -> bool:
    """网格是否像一张成绩表（含姓名表头 + 至少一个数字单元格）。"""
    df = _grid_to_dataframe(rows)
    return df is not None


def docx_score_dataframe(file) -> pd.DataFrame | None:
    """读取 Word 里第一个成绩表格；没有合格表格返回 None。"""
    from docx import Document

    if isinstance(file, (bytes, bytearray)):
        doc = Document(io.BytesIO(file))
    else:
        if hasattr(file, "seek"):
            file.seek(0)
        doc = Document(file)

    for table in doc.tables:
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        df = _grid_to_dataframe(rows)
        if df is not None:
            return df
    return None


def _merge_grids(frames: list[list[list[str]]]) -> pd.DataFrame | None:
    """合并多张同构表：以第一张表头为准，后续若首行重复表头则去掉。"""
    if not frames:
        return None
    base_rows = [list(r) for r in frames[0]]
    header = [_clean(c) for c in base_rows[0]]
    merged = list(base_rows)
    for frame in frames[1:]:
        rows = [list(r) for r in frame]
        if rows and [_clean(c) for c in rows[0]] == header:
            rows = rows[1:]
        merged.extend(rows)
    return _grid_to_dataframe(merged)


def pdf_score_dataframe(file) -> pd.DataFrame | None:
    """读取 PDF 各页的线框成绩表并纵向合并；没有可识别表格返回 None。"""
    import pymupdf

    if isinstance(file, (bytes, bytearray)):
        doc = pymupdf.open(stream=io.BytesIO(file), filetype="pdf")
    elif isinstance(file, (str, Path)):
        # 路径必须按 filename 传；Path 不能塞进 stream 参数。
        doc = pymupdf.open(file)
    else:
        # Streamlit UploadedFile 是类文件对象，PyMuPDF 必须显式走 stream，
        # 直接 pymupdf.open(file_like) 在 1.28 会按文件名处理导致识别失败。
        if hasattr(file, "seek"):
            file.seek(0)
        doc = pymupdf.open(stream=file, filetype="pdf")

    try:
        frames = []
        for page in doc:
            found = page.find_tables()
            for table in found.tables:
                grid = table.extract()
                if grid and _is_score_grid(grid):
                    frames.append(grid)
        return _merge_grids(frames)
    finally:
        doc.close()
