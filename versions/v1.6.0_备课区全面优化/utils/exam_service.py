# -*- coding: utf-8 -*-
"""
考试与成绩数据服务层。

集中处理：
- 考试的增删改查、成绩批量导入（自动建学生、重复导入更新）；
- 单次考试分析所需的数据组装（各科指标、排名、进退步）；
- 趋势分析和学生画像所需的历次成绩数据。

所有函数都接收外部传入的 session，不自己开关连接。
"""

from __future__ import annotations

import json
from collections import defaultdict
from typing import Iterable, Optional

from models.models import Exam, Score, Student
from utils import academic_time, stats as stats_mod
from utils.student_service import get_or_create_student


def list_exams(session, term: str | None = None) -> list[Exam]:
    """考试按日期升序返回；传 term 时只看手填的 term 字段。"""
    q = session.query(Exam)
    if term:
        q = q.filter(Exam.term == term)
    return q.order_by(Exam.exam_date, Exam.id).all()


def list_exam_terms(session) -> list[str]:
    """返回已使用过的非空学期，去重排序。"""
    rows = (session.query(Exam.term)
            .filter(Exam.term.isnot(None), Exam.term != "")
            .distinct().all())
    return sorted({row[0] for row in rows})


def get_exam(session, exam_id: int) -> Optional[Exam]:
    return session.get(Exam, exam_id)


def create_exam(session, name: str, exam_date=None, grade: str | None = None,
                term: str | None = None, full_scores: dict | None = None,
                remark: str | None = None,
                allow_same_name_date: bool = False) -> Exam:
    """新建考试；同名同日期默认视为重复，阻止创建。"""
    name = (name or "").strip()
    if not name:
        raise ValueError("考试名称不能为空")
    if exam_date is not None and not allow_same_name_date:
        duplicate = (session.query(Exam)
                     .filter(Exam.name == name, Exam.exam_date == exam_date)
                     .first())
        if duplicate is not None:
            raise ValueError("同名、同日期的考试已存在，请直接选择已有考试或修改日期。")
    exam = Exam(
        name=name, exam_date=exam_date, grade=grade, term=term,
        full_scores=json.dumps(full_scores or {}, ensure_ascii=False),
        remark=remark,
    )
    session.add(exam)
    session.flush()
    return exam


def update_exam(session, exam_id: int, **fields) -> None:
    allowed = {"name", "exam_date", "grade", "term", "remark"}
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise ValueError(f"考试不存在：id={exam_id}")
    for key, value in fields.items():
        if key in allowed:
            setattr(exam, key, value)
    if "full_scores" in fields and isinstance(fields["full_scores"], dict):
        exam.full_scores = json.dumps(fields["full_scores"], ensure_ascii=False)


def delete_exam(session, exam_id: int) -> None:
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise ValueError(f"考试不存在：id={exam_id}")
    session.delete(exam)


def rows_to_full_scores(rows: list[dict]) -> dict[str, float]:
    """逐行满分配置转 {科目: 满分}；空行忽略，重复/非正数报错。"""
    result: dict[str, float] = {}
    for row in rows or []:
        subject = str(row.get("subject") or row.get("学科") or "").strip()
        if not subject:
            continue
        try:
            score = float(row.get("score", row.get("满分")))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"“{subject}”的满分必须是数字。") from exc
        if score <= 0:
            raise ValueError(f"“{subject}”的满分必须大于 0。")
        if subject in result:
            raise ValueError(f"学科重复：{subject}")
        result[subject] = score
    return result


def full_scores_of(exam: Exam, subjects: list[str],
                   default_full: float = 100.0) -> dict:
    """读取考试各科满分；缺失科目用默认满分 100 兜底。"""
    try:
        saved = json.loads(exam.full_scores) if exam.full_scores else {}
    except (TypeError, ValueError):
        saved = {}
    return {s: float(saved.get(s, default_full)) for s in subjects}


def _get_or_create_score(session, exam_id: int, student_id: int,
                         subject: str) -> Score:
    score = (session.query(Score)
             .filter(Score.exam_id == exam_id, Score.student_id == student_id,
                     Score.subject == subject).first())
    if score is None:
        score = Score(exam_id=exam_id, student_id=student_id, subject=subject)
        session.add(score)
    return score


def import_scores(session, exam_id: int, records: list[dict]) -> dict:
    """
    批量导入成绩；系统里没有的学生按 姓名+班级 自动新建。
    分数为 None 的缺考不覆盖已有分数；重复导入更新。
    """
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise ValueError(f"考试不存在：id={exam_id}")

    created_students = scores_written = scores_updated = skipped_rows = 0
    for rec in records or []:
        name = (rec.get("name") or "").strip() if rec.get("name") else ""
        if not name:
            skipped_rows += 1
            continue
        class_name = (rec.get("class_name") or "").strip() or None
        student, created = get_or_create_student(
            session, name, class_name,
            student_no=(rec.get("student_no") or None))
        if created:
            created_students += 1
        for subject, raw_score in (rec.get("scores") or {}).items():
            subject = str(subject).strip()
            if not subject:
                continue
            score_value = None if raw_score is None else float(raw_score)
            score = _get_or_create_score(session, exam_id, student.id, subject)
            if score_value is None:
                continue
            if score.score is None:
                scores_written += 1
            else:
                scores_updated += 1
            score.score = score_value
    return {
        "created_students": created_students,
        "scores_written": scores_written,
        "scores_updated": scores_updated,
        "skipped_rows": skipped_rows,
    }


def exam_subjects(session, exam_id: int) -> list[str]:
    """这场考试实际录过分的科目，按名字排序。"""
    rows = (session.query(Score.subject)
            .filter(Score.exam_id == exam_id)
            .distinct().all())
    return sorted({r[0] for r in rows})


def exam_score_rows(session, exam_id: int) -> list[dict]:
    """取一场考试宽表：每个学生一行，含各科分数和总分。"""
    exam = session.get(Exam, exam_id)
    if exam is None:
        return []
    subjects = exam_subjects(session, exam_id)
    scores = session.query(Score).filter(Score.exam_id == exam_id).all()
    student_ids = {s.student_id for s in scores}
    students = {st.id: st for st in session.query(Student)
                .filter(Student.id.in_(student_ids)).all()}
    by_student = defaultdict(dict)
    meta = {}
    for s in scores:
        student = students.get(s.student_id)
        if student is None:
            continue
        by_student[student.id][s.subject] = s.score
        meta[student.id] = {"name": student.name,
                            "class_name": student.class_name,
                            "student_no": student.student_no}
    result = []
    for student_id, subject_scores in by_student.items():
        valid = [v for v in subject_scores.values() if v is not None]
        total = round(sum(valid), 2) if valid else None
        row = {
            "student_id": student_id,
            "name": meta[student_id]["name"],
            "class_name": meta[student_id]["class_name"],
            "student_no": meta[student_id]["student_no"],
        }
        for subject in subjects:
            row[subject] = subject_scores.get(subject)
        row["total"] = total
        result.append(row)
    return result


def _rank_within_class(rows: list[dict], key: str) -> None:
    """给宽表按班内某列降序并列排名，写回 {key}_rank。"""
    groups = defaultdict(list)
    for row in rows:
        groups[row["class_name"]].append(row)
    for class_rows in groups.values():
        ranks = stats_mod.competition_rank([r[key] for r in class_rows])
        for row, rank in zip(class_rows, ranks):
            row[f"{key}_rank"] = rank


def _normalize_thresholds(thresholds: dict | None) -> dict:
    if thresholds is not None:
        from utils.analysis_settings import validate_thresholds
        return validate_thresholds(thresholds)
    from utils.analysis_settings import load_thresholds
    return load_thresholds()


def analyze_exam(session, exam_id: int, class_name: str | None = None,
                 thresholds: dict | None = None) -> dict:
    """组装单次考试总分总览、各科统计、分数段、班内排名和进退步。"""
    exam = session.get(Exam, exam_id)
    if exam is None:
        return {}
    subjects = exam_subjects(session, exam_id)
    full_scores = full_scores_of(exam, subjects)
    rows = exam_score_rows(session, exam_id)
    if class_name:
        rows = [r for r in rows if r["class_name"] == class_name]

    for subject in subjects:
        _rank_within_class(rows, subject)
    _rank_within_class(rows, "total")

    t = _normalize_thresholds(thresholds)
    subject_stats = {}
    for subject in subjects:
        subject_stats[subject] = stats_mod.subject_summary(
            [r[subject] for r in rows], full_scores[subject],
            t["pass_ratio"], t["excellent_ratio"], custom_bands=True)
    total_full = round(sum(full_scores.values()), 2)
    total_summary = stats_mod.subject_summary(
        [r["total"] for r in rows], total_full,
        t["pass_ratio"], t["excellent_ratio"], custom_bands=True)
    deltas = _previous_exam_deltas(session, exam, rows, class_name)
    return {
        "exam": exam,
        "subjects": subjects,
        "full_scores": full_scores,
        "subject_stats": subject_stats,
        "total_full": total_full,
        "total_stats": total_summary,
        "rows": sorted(rows, key=lambda r: (r["class_name"] or "",
                                           r["total_rank"] or 9999)),
        "deltas": deltas,
        "thresholds": t,
    }


def analyze_subject(session, exam_id: int, subject: str,
                    class_name: str | None = None,
                    thresholds: dict | None = None) -> dict:
    """单科视角分析：指标、三段分数分布、班内并列排名、同科进退步。"""
    exam = session.get(Exam, exam_id)
    if exam is None:
        return {}
    subjects = exam_subjects(session, exam_id)
    if subject not in subjects:
        return {}
    full_scores = full_scores_of(exam, subjects)
    full_score = full_scores[subject]
    rows = exam_score_rows(session, exam_id)
    if class_name:
        rows = [r for r in rows if r["class_name"] == class_name]
    rows = [r for r in rows if r.get(subject) is not None]
    if not rows:
        return {}

    _rank_within_class(rows, subject)
    t = _normalize_thresholds(thresholds)
    summary = stats_mod.subject_summary(
        [r[subject] for r in rows], full_score,
        t["pass_ratio"], t["excellent_ratio"], custom_bands=True)
    deltas = _previous_subject_deltas(session, exam, subject, rows, class_name)
    return {
        "exam": exam,
        "subject": subject,
        "full_score": full_score,
        "stats": summary,
        "rows": sorted(rows, key=lambda r: (r["class_name"] or "",
                                            r[f"{subject}_rank"] or 9999)),
        "deltas": deltas,
        "thresholds": t,
    }


def _previous_exam(session, exam: Exam) -> Optional[Exam]:
    return (session.query(Exam)
            .filter(Exam.exam_date < exam.exam_date)
            .order_by(Exam.exam_date.desc(), Exam.id.desc()).first())


def _previous_subject_deltas(session, exam: Exam, subject: str,
                             current_rows: list[dict],
                             class_name: str | None) -> dict:
    prev = _previous_exam(session, exam)
    if prev is None:
        return {}
    prev_rows = exam_score_rows(session, prev.id)
    if class_name:
        prev_rows = [r for r in prev_rows if r["class_name"] == class_name]
    cur_map = {r["student_id"]: r[subject] for r in current_rows}
    prev_map = {r["student_id"]: r.get(subject) for r in prev_rows}
    raw = stats_mod.score_deltas(cur_map, prev_map)
    _rank_within_class(prev_rows, subject)
    prev_rank = {r["student_id"]: r[f"{subject}_rank"] for r in prev_rows}
    cur_rank = {r["student_id"]: r[f"{subject}_rank"] for r in current_rows}
    for sid, info in raw.items():
        if sid in prev_rank and sid in cur_rank and prev_rank[sid] and cur_rank[sid]:
            info["rank_delta"] = prev_rank[sid] - cur_rank[sid]
    return raw


def _previous_exam_deltas(session, exam: Exam, current_rows: list[dict],
                          class_name: str | None) -> dict:
    prev = _previous_exam(session, exam)
    if prev is None:
        return {}
    prev_rows = exam_score_rows(session, prev.id)
    if class_name:
        prev_rows = [r for r in prev_rows if r["class_name"] == class_name]
    raw = stats_mod.score_deltas(
        {r["student_id"]: r["total"] for r in current_rows},
        {r["student_id"]: r["total"] for r in prev_rows})
    _rank_within_class(prev_rows, "total")
    prev_rank = {r["student_id"]: r["total_rank"] for r in prev_rows}
    cur_rank = {r["student_id"]: r["total_rank"] for r in current_rows}
    for sid, info in raw.items():
        if sid in prev_rank and sid in cur_rank and prev_rank[sid] and cur_rank[sid]:
            info["rank_delta"] = prev_rank[sid] - cur_rank[sid]
    return raw


def student_scores_over_time(session, student_id: int,
                             exam_ids: Iterable[int] | None = None) -> list[dict]:
    """某学生历次成绩；传 exam_ids 时只返回指定考试，顺序仍按日期。"""
    exams = list_exams(session)
    if exam_ids is not None:
        wanted = {int(x) for x in exam_ids}
        exams = [e for e in exams if e.id in wanted]
    wanted_ids = {e.id for e in exams}
    scores = (session.query(Score)
              .filter(Score.student_id == student_id)
              .all())
    by_exam = defaultdict(dict)
    subjects_involved = set()
    for s in scores:
        if s.exam_id not in wanted_ids:
            continue
        by_exam[s.exam_id][s.subject] = s.score
        subjects_involved.add(s.subject)
    result = []
    for exam in exams:
        if exam.id not in by_exam:
            continue
        subject_scores = by_exam[exam.id]
        valid = [v for v in subject_scores.values() if v is not None]
        item = {"exam_id": exam.id, "exam_name": exam.name,
                "exam_date": exam.exam_date}
        for subject in subjects_involved:
            item[subject] = subject_scores.get(subject)
        item["total"] = round(sum(valid), 2) if valid else None
        result.append(item)
    return result


def class_trend(session, class_name: str | None = None,
                subject: str | None = None,
                exam_ids: Iterable[int] | None = None) -> dict:
    """
    班级趋势；传 subject 只返回该科，传 exam_ids 只统计指定考试。
    返回考试名、ID、日期、唯一标签、各科/总分均分和满分元信息。
    """
    all_exams = list_exams(session)
    if exam_ids is not None:
        wanted = {int(x) for x in exam_ids}
        all_exams = [e for e in all_exams if e.id in wanted]
    labels = academic_time.unique_exam_labels(all_exams)
    all_subjects = set()
    if subject is not None:
        all_subjects.add(subject)
    points = []
    for exam in all_exams:
        rows = exam_score_rows(session, exam.id)
        if class_name:
            rows = [r for r in rows if r["class_name"] == class_name]
        if not rows:
            continue
        exam_subject_list = exam_subjects(session, exam.id)
        full_scores = full_scores_of(exam, exam_subject_list)
        if subject is not None:
            if subject not in exam_subject_list:
                continue
            vals = [r[subject] for r in rows if r.get(subject) is not None]
            if not vals:
                continue
            points.append({
                "exam_id": exam.id,
                "exam_name": exam.name,
                "exam_date": exam.exam_date,
                "label": labels.get(exam.id, exam.name),
                "values": {subject: round(sum(vals) / len(vals), 2)},
                "full_scores": {subject: full_scores.get(subject, 100.0)},
            })
            continue
        all_subjects.update(exam_subject_list)
        point = {
            "exam_id": exam.id,
            "exam_name": exam.name,
            "exam_date": exam.exam_date,
            "label": labels.get(exam.id, exam.name),
            "values": {},
            "full_scores": full_scores,
            "total_full_score": round(sum(full_scores.values()), 2),
        }
        for item_subject in exam_subject_list:
            vals = [r[item_subject] for r in rows if r[item_subject] is not None]
            point["values"][item_subject] = round(sum(vals) / len(vals), 2) if vals else None
        totals = [r["total"] for r in rows if r["total"] is not None]
        point["values"]["total"] = round(sum(totals) / len(totals), 2) if totals else None
        points.append(point)

    subjects = sorted(all_subjects)
    if subject is not None:
        return {
            "exams": [p["exam_name"] for p in points],
            "labels": [p["label"] for p in points],
            "exam_ids": [p["exam_id"] for p in points],
            "exam_dates": [p["exam_date"] for p in points],
            "exam_meta": points,
            "subjects": [subject],
            "series": {subject: [p["values"].get(subject) for p in points]},
        }
    series = {s: [p["values"].get(s) for p in points] for s in subjects}
    series["total"] = [p["values"].get("total") for p in points]
    return {
        "exams": [p["exam_name"] for p in points],
        "labels": [p["label"] for p in points],
        "exam_ids": [p["exam_id"] for p in points],
        "exam_dates": [p["exam_date"] for p in points],
        "exam_meta": points,
        "subjects": subjects,
        "series": series,
    }


def student_score_history_rows(session, class_name: str | None = None,
                               exam_ids: Iterable[int] | None = None) -> list[dict]:
    """组装跨考试长表：一行=一个学生一场考试一个科目。"""
    exams = list_exams(session)
    if exam_ids is not None:
        wanted = {int(x) for x in exam_ids}
        exams = [e for e in exams if e.id in wanted]
    result = []
    for exam in exams:
        wide = analyze_exam(session, exam.id, class_name=class_name)
        if not wide:
            continue
        for row in wide["rows"]:
            for subject in wide["subjects"]:
                result.append({
                    "考试日期": exam.exam_date,
                    "学期": academic_time.semester_name(exam.exam_date),
                    "学年": academic_time.academic_year_name(exam.exam_date),
                    "考试类型": academic_time.exam_type(exam.name),
                    "考试": exam.name,
                    "班级": row.get("class_name") or "",
                    "学号": row.get("student_no") or "",
                    "姓名": row["name"],
                    "科目": subject,
                    "满分": wide["full_scores"][subject],
                    "分数": row.get(subject),
                    "科目班级排名": row.get(f"{subject}_rank"),
                    "总分": row.get("total"),
                    "总分班级排名": row.get("total_rank"),
                })
    return result

# ---------------------------------------------------------------------------
# 格式化成绩单 Excel
# ---------------------------------------------------------------------------

def build_score_report_xlsx(session, exam_id: int,
                            class_name: str | None = None) -> bytes:
    """生成某场考试的格式化成绩单，返回 xlsx 字节。"""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    analysis = analyze_exam(session, exam_id, class_name=class_name)
    if not analysis:
        raise ValueError("这场考试还没有成绩，无法导出成绩单。")

    exam = analysis["exam"]
    subjects = list(analysis["subjects"])
    rows = [r for r in analysis["rows"] if r.get("total") is not None]
    if not rows:
        raise ValueError("这场考试还没有有效总分，无法导出成绩单。")

    headers = ["姓名", "学号", "班级"] + subjects + ["总分", "班级排名"]
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩单"

    scope = class_name or "全年级"
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(headers))
    title_cell = ws.cell(1, 1, f"{scope} {exam.name} 成绩单")
    title_cell.font = Font(name="微软雅黑", size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    header_fill = PatternFill("solid", fgColor="DEEAF6")
    thin = Side(style="thin", color="9EADCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name="微软雅黑", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(2, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    numeric_cols = set(subjects + ["总分"])
    for row_no, row in enumerate(rows, start=3):
        values = [
            row["name"], row.get("student_no") or "",
            row.get("class_name") or "",
            *[row.get(s) for s in subjects],
            row.get("total"), row.get("total_rank"),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_no, col, value)
            cell.border = border
            cell.alignment = center
            cell.font = Font(name="微软雅黑")
            if headers[col - 1] in numeric_cols:
                cell.number_format = "0.0"

    summary_start = len(rows) + 3
    for offset, label in enumerate(("平均分", "最高分", "最低分")):
        row_no = summary_start + offset
        ws.cell(row_no, 1, label)
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row_no, col)
            cell.border = border
            cell.alignment = center
            cell.font = Font(name="微软雅黑", bold=col == 1)
            if header in numeric_cols:
                vals = [r[header] for r in rows if r.get(header) is not None]
                if vals:
                    value = (round(sum(vals) / len(vals), 1) if label == "平均分"
                             else max(vals) if label == "最高分" else min(vals))
                    cell.value = value
                    cell.number_format = "0.0"

    widths = {"姓名": 12, "学号": 12, "班级": 14, "总分": 10, "班级排名": 10}
    for col, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(header, 10)
    ws.freeze_panes = "A3"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
