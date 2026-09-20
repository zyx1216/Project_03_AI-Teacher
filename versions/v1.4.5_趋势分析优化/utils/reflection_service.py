# -*- coding: utf-8 -*-
"""
教学反思服务层（v1.0）。

职责：
- 聚合“成绩数据 + 作业错题”文本，作为 user 消息注入给主模型；
- 四段式反思（成功之处/不足之处/学生反馈/改进措施）的解析、保存、读取、删除；
- 导出 Word。

所有数据库函数收外部 session，便于用内存库单测；本模块不直接调用 LLM，
逐人/逐次调用由页面负责，保证离线可测。
"""

from __future__ import annotations

import json
from datetime import date, datetime

from models.models import (
    Exam, Homework, HomeworkAnswer, Question, Student, TeachingReflection)
from utils import exam_service, homework_stats as hstat

# 四段标题（AI 输出用 “## 标题” 开头，解析时按这个顺序）
SECTIONS = ["成功之处", "不足之处", "学生反馈", "改进措施"]


# ---------------------------------------------------------------------------
# 数据聚合
# ---------------------------------------------------------------------------

def exams_in_scope(session, scope_type: str, start_exam_id: int,
                   end_exam_id: int | None = None) -> list[Exam]:
    """按范围取考试：exam 只取一场；range 取起止考试日期之间（含两端）的全部考试。"""
    start = exam_service.get_exam(session, start_exam_id)
    if start is None:
        return []
    if scope_type != "range" or not end_exam_id:
        return [start]
    end = exam_service.get_exam(session, end_exam_id)
    if end is None:
        return [start]
    d_from, d_to = sorted([start.exam_date, end.exam_date])
    return [e for e in exam_service.list_exams(session)
            if d_from <= e.exam_date <= d_to]


def _pct(x: float | None) -> str:
    return "—" if x is None else f"{x * 100:.1f}%"


def _exam_metrics_lines(session, exam: Exam, class_name: str | None,
                        thresholds: dict | None = None) -> list[str]:
    """一场考试的指标文本行（复用考试分析口径）。"""
    data = exam_service.analyze_exam(
        session, exam.id, class_name=class_name, thresholds=thresholds)
    if not data or not data["rows"]:
        return [f"- 《{exam.name}》（{exam.exam_date}）：暂无成绩数据。"]
    lines = [f"- 《{exam.name}》（{exam.exam_date}）参考 {len(data['rows'])} 人："]
    for subject in data["subjects"]:
        stt = data["subject_stats"][subject]
        lines.append(
            f"    {subject}（满分{data['full_scores'][subject]:g}）："
            f"均分{stt['mean']}，最高{stt['max']:g}，"
            f"最低{stt['min']:g}，及格率{_pct(stt['pass_rate'])}，"
            f"优秀率{_pct(stt['excellent_rate'])}")
    total = data["total_stats"]
    lines.append(
        f"    总分（满分{data['total_full']:g}）：均分{total['mean']}，"
        f"及格率{_pct(total['pass_rate'])}，优秀率{_pct(total['excellent_rate'])}")
    # 进退步人数（只在有上次考试时才有）
    if data.get("deltas"):
        up = sum(1 for d in data["deltas"].values()
                 if d.get("score_delta") is not None and d["score_delta"] > 0)
        down = sum(1 for d in data["deltas"].values()
                   if d.get("score_delta") is not None and d["score_delta"] < 0)
        lines.append(f"    较上一次考试：进步 {up} 人，退步 {down} 人。")
    return lines


def collect_homework_errors(session, date_from: date | None, date_to: date | None,
                            class_name: str | None) -> dict:
    """
    聚合一段时间内作业的逐题作答，返回：
    {questions:[{content,knowledge_points,judged,wrong,rate}], error_types:{类型:次数}, total_wrong}
    只统计已判对错（is_correct 非空）的作答。
    """
    q = (session.query(HomeworkAnswer, Homework, Question, Student)
         .join(Homework, HomeworkAnswer.homework_id == Homework.id)
         .join(Question, HomeworkAnswer.question_id == Question.id)
         .join(Student, HomeworkAnswer.student_id == Student.id)
         .filter(Homework.is_template.is_(False))
         .filter(HomeworkAnswer.is_correct.isnot(None)))
    if date_from is not None:
        q = q.filter(Homework.created_at >= datetime.combine(date_from, datetime.min.time()))
    if date_to is not None:
        q = q.filter(Homework.created_at <= datetime.combine(date_to, datetime.max.time()))
    if class_name:
        q = q.filter(Student.class_name == class_name)

    bucket: dict[int, dict] = {}
    error_types: dict[str, int] = {}
    total_wrong = 0
    for ans, _hw, question, _stu in q.all():
        g = bucket.setdefault(question.id, {
            "content": question.content, "knowledge_points": _kp_list(question),
            "judged": 0, "wrong": 0})
        g["judged"] += 1
        if not ans.is_correct:
            g["wrong"] += 1
            total_wrong += 1
            if ans.error_type:
                error_types[ans.error_type] = error_types.get(ans.error_type, 0) + 1

    questions = []
    for g in bucket.values():
        g["rate"] = round(g["wrong"] / g["judged"], 4) if g["judged"] else None
        questions.append(g)
    # 高频错题：错误人数多的在前
    questions.sort(key=lambda g: (g["wrong"], g["rate"] or 0), reverse=True)
    return {"questions": questions, "error_types": error_types,
            "total_wrong": total_wrong}


def _kp_list(question: Question) -> list[str]:
    try:
        data = json.loads(question.knowledge_points) if question.knowledge_points else []
        return data if isinstance(data, list) else []
    except (ValueError, TypeError):
        return []


def build_reflection_data_text(session, scope_type: str, start_exam_id: int,
                               end_exam_id: int | None = None,
                               class_name: str | None = None,
                               thresholds: dict | None = None) -> str:
    """组装发给 AI 的数据文本（成绩指标 + 作业错题）。"""
    exams = exams_in_scope(session, scope_type, start_exam_id, end_exam_id)
    if not exams:
        return "没有找到考试数据。"

    scope_label = (f"{exams[0].name} 至 {exams[-1].name}"
                   if scope_type == "range" and len(exams) > 1 else exams[0].name)
    lines = [f"反思范围：{scope_label}",
             f"班级：{class_name or '全部班级'}",
             "",
             "【一、考试成绩数据】"]
    for exam in exams:
        lines.extend(_exam_metrics_lines(session, exam, class_name, thresholds))

    # 作业错题时间窗：单次考上一场之后到本场；时间段就用两端考试日期
    if scope_type == "range" and len(exams) > 1:
        d_from, d_to = exams[0].exam_date, exams[-1].exam_date
    else:
        prev = exam_service._previous_exam(session, exams[0])
        d_from = prev.exam_date if prev else None
        d_to = exams[0].exam_date

    errors = collect_homework_errors(session, d_from, d_to, class_name)
    lines.extend(["", "【二、作业错题数据】"])
    if errors["total_wrong"] == 0:
        lines.append("这段时间没有逐题批改的作业错题数据（可能只录了总分），请主要依据成绩数据分析。")
    else:
        if errors["error_types"]:
            et = "，".join(f"{k}{v} 次" for k, v in
                          sorted(errors["error_types"].items(), key=lambda x: -x[1]))
            lines.append(f"错误类型分布：{et}。")
        lines.append(f"高频错题 TOP5（共 {errors['total_wrong']} 个错误作答）：")
        for i, g in enumerate(errors["questions"][:5], start=1):
            kps = "、".join(g["knowledge_points"]) or "未标注知识点"
            preview = g["content"].replace("\n", " ")[:60]
            lines.append(f"{i}. {preview}……｜知识点：{kps}"
                         f"｜错误 {g['wrong']}/{g['judged']} 人，错误率{_pct(g['rate'])}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 反思内容解析与留档
# ---------------------------------------------------------------------------

def parse_sections(text: str) -> dict:
    """
    把 AI 输出按 “## 标题” 解析成 {四段标题: 正文}。
    解析不出来时整篇归到“成功之处”，保证内容不丢。
    """
    result = {name: "" for name in SECTIONS}
    current = None
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        matched = None
        if line.startswith("#"):
            for name in SECTIONS:
                if name in line:
                    matched = name
                    break
        if matched:
            current = matched
            continue
        if current and line:
            result[current] = (result[current] + "\n" + line).strip()
    if not any(result.values()):
        result[SECTIONS[0]] = (text or "").strip()
    return result


def sections_to_text(sections: dict) -> str:
    """四段字典转成带标题的纯文本（AI 原文风格），供编辑框和导出共用。"""
    return "\n\n".join(f"## {name}\n{(sections.get(name) or '').strip()}"
                       for name in SECTIONS).strip()


def save_reflection(session, title: str, scope_type: str, start_exam_id: int,
                    class_name: str | None, content: str | dict,
                    end_exam_id: int | None = None,
                    reflection_id: int | None = None) -> TeachingReflection:
    """新建或更新一条反思。content 可传四段字典或纯文本（自动解析）。"""
    if isinstance(content, dict):
        sections = {name: str(content.get(name, "")) for name in SECTIONS}
        raw = sections_to_text(sections)
    else:
        raw = (content or "").strip()
        sections = parse_sections(raw)
    stored = json.dumps({"raw": raw, "sections": sections}, ensure_ascii=False)

    if reflection_id is not None:
        row = session.get(TeachingReflection, reflection_id)
        if row is None:
            raise ValueError(f"反思不存在：id={reflection_id}")
        row.title, row.content = title, stored
        row.updated_at = datetime.now()
        return row

    row = TeachingReflection(
        title=title, scope_type=scope_type, start_exam_id=start_exam_id,
        end_exam_id=end_exam_id if scope_type == "range" else None,
        class_name=class_name, content=stored)
    session.add(row)
    session.flush()
    return row


def load_reflection_content(reflection: TeachingReflection) -> dict:
    """读取留档反思，返回 {"raw": 全文, "sections": 四段字典}，兼容旧数据。"""
    try:
        data = json.loads(reflection.content) if reflection.content else {}
        sections = data.get("sections") or parse_sections(data.get("raw", ""))
        return {"raw": data.get("raw", reflection.content or ""),
                "sections": {name: sections.get(name, "") for name in SECTIONS}}
    except (ValueError, TypeError):
        return {"raw": reflection.content or "",
                "sections": parse_sections(reflection.content or "")}


def list_reflections(session, class_name: str | None = None) -> list[TeachingReflection]:
    q = session.query(TeachingReflection)
    if class_name:
        q = q.filter(TeachingReflection.class_name == class_name)
    return q.order_by(TeachingReflection.id.desc()).all()


def delete_reflection(session, reflection_id: int) -> None:
    row = session.get(TeachingReflection, reflection_id)
    if row is None:
        raise ValueError(f"反思不存在：id={reflection_id}")
    session.delete(row)


def export_word(reflection: TeachingReflection) -> bytes:
    """把反思导出成含四个小节标题的 Word。"""
    import io
    from docx import Document

    data = load_reflection_content(reflection)
    doc = Document()
    doc.add_heading(reflection.title, level=0)
    meta = "　".join(x for x in [reflection.class_name,
                                 reflection.created_at.strftime("%Y-%m-%d")
                                 if reflection.created_at else ""] if x)
    if meta:
        doc.add_paragraph(meta)
    for name in SECTIONS:
        doc.add_heading(name, level=1)
        doc.add_paragraph(data["sections"].get(name) or "—")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
