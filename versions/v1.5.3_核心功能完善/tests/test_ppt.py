# -*- coding: utf-8 -*-
"""PPT 生成测试：离线生成后用 python-pptx 重开，断言页数与标题、要点约束。"""

import io

from pptx import Presentation

from utils import ppt_generator as ppt
from utils import lesson_service as ls


def _lesson_plan():
    plan = ls.empty_plan()
    plan["objectives"] = {"knowledge": "理解求根公式", "process": "通过推导探究",
                          "emotion": "体会数学严谨"}
    plan["process"][0]["content"] = "复习旧知。提问引入。激发兴趣。"
    plan["process"][1]["content"] = "讲解公式。例题板演。" * 10  # 内容多，触发分页
    plan["process"][2]["content"] = "课堂练习第一题。课堂练习第二题。"
    return plan


class _FakeLesson:
    title = "一元二次方程求根公式"
    grade = "九年级"
    chapter = "第21章"


def _open(data):
    return Presentation(io.BytesIO(data))


def _slide_texts(prs):
    texts = []
    for slide in prs.slides:
        buf = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                buf.append(shape.text_frame.text)
        texts.append("\n".join(buf))
    return texts


def test_ppt_returns_valid_file_with_core_pages():
    data = ppt.generate_ppt(_FakeLesson(), _lesson_plan())
    prs = _open(data)
    assert len(prs.slides) >= 4  # 封面 + 目标 + 至少导入/新授
    texts = _slide_texts(prs)
    # 封面含课题与年级
    assert "一元二次方程求根公式" in texts[0]
    assert "九年级" in texts[0]
    # 存在学习目标页
    assert any("学习目标" in t for t in texts)


def test_ppt_bullets_limited_and_short():
    data = ppt.generate_ppt(_FakeLesson(), _lesson_plan())
    prs = _open(data)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            bullets = [p.text for p in shape.text_frame.paragraphs
                       if p.text.strip().startswith("•")]
            assert len(bullets) <= 5  # 每页最多 5 条
            for b in bullets:
                # 20 字 + 项目符号和省略号
                assert len(b.lstrip("• ")) <= 21


def test_ppt_title_color_is_deep_blue():
    prs = _open(ppt.generate_ppt(_FakeLesson(), _lesson_plan()))
    # 找一张含“学习目标”标题的页，其标题文字颜色应为深蓝 #1F4E79
    found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and "学习目标" in shape.text_frame.text:
                run = shape.text_frame.paragraphs[0].runs[0]
                rgb = run.font.color.rgb
                assert str(rgb) == "1F4E79"
                found = True
    assert found


def test_split_and_paginate_logic():
    bullets = ppt._split_bullets("第一句话很长很长很长很长很长很长很长很长很长。第二句。")
    assert bullets[0].endswith("…") or len(bullets[0]) <= 21
    pages = ppt._paginate(list(range(12)))
    assert len(pages) == 3 and all(len(p) <= 5 for p in pages)


def test_empty_stage_not_rendered():
    plan = ls.empty_plan()  # 所有环节内容为空
    data = ppt.generate_ppt(_FakeLesson(), plan)
    prs = _open(data)
    # 只有封面 + 学习目标，空环节不出空白页
    assert len(prs.slides) == 2