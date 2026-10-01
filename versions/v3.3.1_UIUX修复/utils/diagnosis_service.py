# -*- coding: utf-8 -*-
"""v2.9.1 AI 教学诊断服务。"""

from __future__ import annotations
import io
import json
from collections import defaultdict
from datetime import date, datetime, timedelta

from models.models import (Exam, Homework, HomeworkAnswer, HomeworkScore,
                           Question, Score, Student, TeachingDiagnosis,
                           TeachingProgress)
from utils import llm_client, question_service as qs, tiered_homework_service


def normalize_period(period):
    period = dict(period or {})
    kind = str(period.get("type") or "month")
    if kind not in ("exam", "month", "semester", "custom"):
        kind = "month"
    return {"type": kind, "exam_id": period.get("exam_id"),
            "start_date": period.get("start_date"),
            "end_date": period.get("end_date")}


def period_dates(session, period):
    period = normalize_period(period)
    today = date.today()
    if period["type"] == "exam" and period.get("exam_id"):
        exam = session.get(Exam, int(period["exam_id"]))
        return (exam.exam_date, exam.exam_date) if exam else (None, None)
    if period["type"] == "month":
        return today - timedelta(days=30), today
    if period["type"] == "semester":
        if today.month >= 9:
            return date(today.year, 9, 1), date(today.year + 1, 1, 31)
        if today.month == 1:
            return date(today.year - 1, 9, 1), date(today.year, 1, 31)
        return date(today.year, 2, 1), date(today.year, 8, 31)
    def parse(value):
        try:
            return date.fromisoformat(str(value)) if value else None
        except ValueError:
            return None
    return parse(period.get("start_date")), parse(period.get("end_date"))


def _in_range(value, start, end):
    if value is None:
        return True
    day = value.date() if isinstance(value, datetime) else value
    return (not start or day >= start) and (not end or day <= end)


def collect_diagnosis_data(session, class_name=None, subject="数学",
                           grade=None, period=None):
    period = normalize_period(period)
    start, end = period_dates(session, period)
    score_rows = (session.query(Score, Exam, Student)
                  .join(Exam, Score.exam_id == Exam.id)
                  .join(Student, Score.student_id == Student.id)
                  .filter(Score.subject == subject, Score.score.isnot(None)).all())
    by_exam, by_student = defaultdict(list), defaultdict(float)
    for score, exam, student in score_rows:
        if class_name and student.class_name != class_name:
            continue
        if grade and exam.grade and exam.grade != grade:
            continue
        if not _in_range(exam.exam_date, start, end):
            continue
        by_exam[exam.id].append(float(score.score))
        by_student[student.id] += float(score.score)
    exams = []
    for exam_id, values in by_exam.items():
        exam = session.get(Exam, exam_id)
        exams.append({"exam_id": exam_id, "name": exam.name if exam else "",
                      "date": str(exam.exam_date if exam else ""),
                      "average": round(sum(values) / len(values), 2),
                      "max": max(values), "min": min(values), "count": len(values)})
    exams.sort(key=lambda item: item["date"])

    hw_rows = (session.query(HomeworkScore, Homework, Student)
               .join(Homework, HomeworkScore.homework_id == Homework.id)
               .join(Student, HomeworkScore.student_id == Student.id)
               .filter(Homework.subject == subject).all())
    hw_total = hw_submitted = hw_timely = 0
    hw_scores = []
    for hs, hw, student in hw_rows:
        if class_name and student.class_name != class_name:
            continue
        if grade and hw.grade and hw.grade != grade:
            continue
        if not _in_range(hw.created_at, start, end):
            continue
        hw_total += 1
        if hs.submitted and hs.total_score is not None:
            hw_submitted += 1
            hw_scores.append(float(hs.total_score))
            if not hw.due_date or not hs.created_at or hs.created_at <= hw.due_date:
                hw_timely += 1

    answer_rows = (session.query(HomeworkAnswer, Homework, Student, Question)
                   .join(Homework, HomeworkAnswer.homework_id == Homework.id)
                   .join(Student, HomeworkAnswer.student_id == Student.id)
                   .join(Question, HomeworkAnswer.question_id == Question.id)
                   .filter(Homework.subject == subject).all())
    kp_values = defaultdict(list)
    for answer, hw, student, question in answer_rows:
        if class_name and student.class_name != class_name:
            continue
        if grade and hw.grade and hw.grade != grade:
            continue
        if not _in_range(hw.created_at, start, end):
            continue
        if answer.is_correct is None:
            continue
        rate = 1.0 if answer.is_correct else 0.0
        for kp in qs.knowledge_points_list(question) or ["未标注知识点"]:
            kp_values[kp].append(rate)
    knowledge = [{"knowledge_point": kp,
                  "mastery": round(sum(values) / len(values), 4),
                  "count": len(values)} for kp, values in kp_values.items()]
    knowledge.sort(key=lambda item: item["mastery"])

    progress_rows = (session.query(TeachingProgress)
                     .filter(TeachingProgress.subject == subject).all())
    planned = sum(1 for row in progress_rows if row.planned_content)
    actual = sum(1 for row in progress_rows if row.actual_content)
    tier_rows = [{"student_id": sid, "name": str(sid), "total": total}
                 for sid, total in by_student.items()]
    tiers = (tiered_homework_service.build_tier_assignment(tier_rows)
             if tier_rows else {"A": [], "B": [], "C": []})
    return {
        "period": period, "start_date": str(start or ""), "end_date": str(end or ""),
        "exams": exams,
        "exam_average": round(sum(x["average"] for x in exams) / len(exams), 2)
                        if exams else None,
        "homework": {"total": hw_total, "submitted": hw_submitted,
                     "completion_rate": round(hw_submitted / hw_total, 4)
                                        if hw_total else None,
                     "timely_rate": round(hw_timely / hw_submitted, 4)
                                    if hw_submitted else None,
                     "average": round(sum(hw_scores) / len(hw_scores), 2)
                                if hw_scores else None},
        "knowledge": knowledge,
        "progress": {"planned": planned, "actual": actual,
                     "rate": round(actual / planned, 4) if planned else None},
        "tiers": {key: len(value) for key, value in tiers.items()},
    }


def _fallback_report(data):
    avg = data.get("exam_average")
    mastery = [item["mastery"] for item in data.get("knowledge", [])]
    knowledge_rate = sum(mastery) / len(mastery) if mastery else None
    homework_rate = data["homework"].get("completion_rate")
    progress_rate = data["progress"].get("rate")
    parts = [(avg / 100 if avg is not None else 0.6, 40),
             (knowledge_rate if knowledge_rate is not None else 0.6, 30),
             (homework_rate if homework_rate is not None else 0.6, 20),
             (progress_rate if progress_rate is not None else 0.6, 10)]
    score = max(1, min(100, round(sum(value * weight for value, weight in parts))))
    rating = "优秀" if score >= 85 else ("良好" if score >= 70 else "需改进")
    weak = [item["knowledge_point"] for item in data.get("knowledge", [])
            if item["mastery"] < 0.6]
    problems = []
    if avg is not None and avg < 60:
        problems.append({"id": "score", "problem": "班级平均成绩偏低，基础题得分不稳定。",
                         "suggestion": "安排基础题专项复习，并按错题类型分组讲评。"})
    if weak:
        problems.append({"id": "knowledge", "problem": "薄弱知识点：" + "、".join(weak[:4]) + "。",
                         "suggestion": "针对薄弱知识点补充例题、变式练习和当堂检测。"})
    if homework_rate is not None and homework_rate < 0.85:
        problems.append({"id": "homework", "problem": "作业完成率不足，部分学生学习任务未落实。",
                         "suggestion": "分层布置作业并建立未交清单，必要时进行个别沟通。"})
    if not problems:
        problems.append({"id": "stable", "problem": "暂未发现明显异常，需继续跟踪趋势。",
                         "suggestion": "保持当前节奏，下一阶段重点观察薄弱知识点变化。"})
    focus = [item["knowledge_point"] for item in data.get("knowledge", [])[:3]] or ["基础概念", "典型题型", "错题复盘"]
    return {"overall_rating": rating, "score": score,
            "pace": "过快" if (progress_rate or 0) > 0.85
                    else ("过慢" if progress_rate is not None and progress_rate < 0.6 else "正常"),
            "knowledge_mastery": {
                "mastered": [item["knowledge_point"] for item in data.get("knowledge", [])
                             if item["mastery"] >= 0.8],
                "weak": weak, "uncovered": []},
            "tier_distribution": data.get("tiers", {}),
            "problems": problems, "next_focus": focus[:5],
            "ai_deepened": False}


def generate_diagnosis(session, class_name=None, subject="数学", grade=None,
                       period=None, chat_func=None):
    data = collect_diagnosis_data(session, class_name, subject, grade, period)
    report = _fallback_report(data)
    try:
        func = chat_func
        if func is None and llm_client.is_content_configured():
            func = llm_client.chat_content
        if func is not None:
            raw = func("你是教学诊断助手，只输出 JSON。",
                       "根据统计生成教学诊断：\n" + json.dumps(data, ensure_ascii=False),
                       temperature=0.3)
            text = str(raw).strip()
            if text.startswith("```"):
                text = text.strip("`").lstrip("json").strip()
            parsed = json.loads(text)
            if isinstance(parsed, dict):
                report.update(parsed)
                report["ai_deepened"] = True
    except Exception:  # noqa: BLE001
        pass
    row = TeachingDiagnosis(
        class_name=class_name, subject=subject, grade=grade,
        period=json.dumps(data.get("period") or normalize_period(period), ensure_ascii=False),
        report_content=json.dumps({"data": data, "report": report}, ensure_ascii=False),
        score=int(report.get("score") or 1))
    session.add(row)
    session.flush()
    return row


def list_diagnoses(session, class_name=None, subject=None, limit=50):
    query = session.query(TeachingDiagnosis)
    if class_name:
        query = query.filter(TeachingDiagnosis.class_name == class_name)
    if subject:
        query = query.filter(TeachingDiagnosis.subject == subject)
    return query.order_by(TeachingDiagnosis.created_at.desc(),
                          TeachingDiagnosis.id.desc()).limit(limit).all()


def _report(row):
    try:
        return json.loads(row.report_content or "{}").get("report") or {}
    except (TypeError, json.JSONDecodeError):
        return {}


def compare_diagnoses(session, first_id, second_id):
    first = session.get(TeachingDiagnosis, int(first_id))
    second = session.get(TeachingDiagnosis, int(second_id))
    if first is None or second is None:
        raise ValueError("诊断记录不存在。")
    a, b = _report(first), _report(second)
    ka = {str(item.get("id") or item.get("problem")): item.get("problem")
          for item in a.get("problems", [])}
    kb = {str(item.get("id") or item.get("problem")): item.get("problem")
          for item in b.get("problems", [])}
    return {"score_delta": int(b.get("score") or 0) - int(a.get("score") or 0),
            "improved": [kb[key] for key in ka.keys() & kb.keys() if ka[key] != kb[key]],
            "persisting": [kb[key] for key in ka.keys() & kb.keys() if ka[key] == kb[key]],
            "new": [kb[key] for key in kb.keys() - ka.keys()]}


def export_diagnosis_word(row):
    from docx import Document
    report = _report(row)
    doc = Document()
    doc.add_heading(f"{row.class_name or '全班'}·{row.subject}·教学诊断报告", level=0)
    doc.add_paragraph(f"诊断评分：{row.score}")
    doc.add_paragraph(f"整体评价：{report.get('overall_rating', '—')}")
    doc.add_paragraph(f"教学进度：{report.get('pace', '—')}")
    for item in report.get("problems", []):
        doc.add_heading(item.get("problem", "问题"), level=2)
        doc.add_paragraph(item.get("suggestion", ""))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_diagnosis_pdf(row):
    """用项目已有的 PyMuPDF 导出 PDF，不新增依赖。"""
    import fitz
    report = _report(row)
    doc = fitz.open()
    page = doc.new_page()
    lines = [f"{row.class_name or '全班'} · {row.subject} · 教学诊断报告",
             f"诊断评分：{row.score}",
             f"整体评价：{report.get('overall_rating', '—')}",
             f"教学进度：{report.get('pace', '—')}"]
    for item in report.get("problems", []):
        lines.extend([f"问题：{item.get('problem', '')}",
                      f"建议：{item.get('suggestion', '')}", ""])
    page.insert_text((50, 60), "\n".join(lines), fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data
