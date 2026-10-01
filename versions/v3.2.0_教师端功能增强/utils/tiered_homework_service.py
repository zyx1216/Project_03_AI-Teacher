# -*- coding: utf-8 -*-
"""v2.9.0 分层作业服务。

分层口径：A=后30%基础、B=中间40%能力、C=前30%拓展。
生成时先保存三份草稿作业，教师选择层级后再布置。
"""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime

from sqlalchemy import func

from models.models import Exam, Homework, HomeworkAnswer, HomeworkScore, Score, Student
from utils import homework_service as hw_svc
from utils import homework_stats as hstat

TIER_RATIOS = {"A": 0.30, "B": 0.40, "C": 0.30}
DIFFICULTY_RATIOS = {
    "A": {1: 70, 2: 30, 3: 0},
    "B": {1: 20, 2: 60, 3: 20},
    "C": {1: 0, 2: 40, 3: 60},
}
TIER_LABELS = {"A": "基础巩固", "B": "能力提升", "C": "拓展挑战"}


def _rows_for_exam(session, exam_id: int, class_name: str | None) -> list[dict]:
    q = (session.query(Score.student_id, Student.name, Student.class_name,
                       func.sum(Score.score))
         .join(Student, Student.id == Score.student_id)
         .filter(Score.exam_id == exam_id, Score.score.isnot(None)))
    if class_name:
        q = q.filter(Student.class_name == class_name)
    return [{"student_id": sid, "name": name, "class_name": cn, "total": float(total)}
            for sid, name, cn, total in q.group_by(
                Score.student_id, Student.name, Student.class_name).all()
            if total is not None]


def select_tier_exam(session, class_name: str, exam_id: int | None = None):
    """选择最近一场有班级总分的考试；没有数据返回 (None, [])。"""
    if exam_id is not None:
        exam = session.get(Exam, int(exam_id))
        return exam, (_rows_for_exam(session, exam.id, class_name) if exam else [])
    exams = (session.query(Exam).order_by(Exam.exam_date.desc(), Exam.id.desc())
             .limit(30).all())
    for exam in exams:
        rows = _rows_for_exam(session, exam.id, class_name)
        if rows:
            return exam, rows
    return None, []


def build_tier_assignment(rows: list[dict], ratios: dict | None = None) -> dict:
    """按总分从低到高划分 A/B/C；并列按 student_id 稳定截断。"""
    ordered = sorted(rows, key=lambda r: (float(r["total"]), int(r["student_id"])))
    n = len(ordered)
    if not n:
        return {"A": [], "B": [], "C": []}
    ratios = ratios or TIER_RATIOS
    a_count = max(1, round(n * float(ratios.get("A", 0.30))))
    c_count = max(1, round(n * float(ratios.get("C", 0.30))))
    if a_count + c_count > n:
        if n == 1:
            a_count, c_count = 1, 0
        else:
            a_count, c_count = 1, 1
    return {
        "A": ordered[:a_count],
        "B": ordered[a_count:n - c_count],
        "C": ordered[n - c_count:] if c_count else [],
    }


def normalize_tier_assignments(rows: list[dict], assignments: dict | None) -> dict:
    """校验教师调整后的名单；每名学生必须且只能出现在一层。"""
    if not assignments:
        return build_tier_assignment(rows)
    by_id = {int(r["student_id"]): r for r in rows}
    result = {"A": [], "B": [], "C": []}
    used = set()
    for tier in ("A", "B", "C"):
        for sid in assignments.get(tier, []):
            sid = int(sid)
            if sid not in by_id:
                raise ValueError(f"分层名单包含不存在的学生：{sid}")
            if sid in used:
                raise ValueError(f"学生重复分到多层：{sid}")
            used.add(sid)
            result[tier].append(by_id[sid])
    missing = set(by_id) - used
    if missing:
        raise ValueError(f"还有 {len(missing)} 名学生未分层。")
    return result


def build_tier_spec(tier: str, count: int = 5) -> dict:
    """按层级难度比例生成自动组卷规格。"""
    count = int(count)
    if count < 1:
        raise ValueError("每层题量必须大于 0。")
    base = count // 3
    remainder = count % 3
    counts = {"choice": base, "fill": base, "solution": base}
    for key in ("choice", "fill", "solution")[:remainder]:
        counts[key] += 1
    return {"counts": counts, "difficulty_ratio": DIFFICULTY_RATIOS[tier]}


def generate_tiered_homework(session, class_name: str, subject: str,
                             weak_points=None, exam_id=None,
                             questions_per_layer: int = 5,
                             tier_assignments: dict | None = None,
                             selected_levels=None) -> dict:
    """生成 A/B/C 三份草稿作业；selected_levels 非空时同步布置。"""
    exam, rows = select_tier_exam(session, class_name, exam_id)
    if not rows:
        raise ValueError("该班暂无可用于分层的考试成绩。")
    tiers = normalize_tier_assignments(rows, tier_assignments)
    grade = exam.grade if exam else None
    batch_time = datetime.now()
    homework_ids, name_lists = {}, {}
    for tier in ("A", "B", "C"):
        members = tiers.get(tier) or []
        if not members:
            continue
        name_lists[tier] = [m["name"] for m in members]
        hw = hw_svc.create_homework(
            session, f"{class_name}·{tier}层作业", homework_type="after_class",
            class_name=class_name, subject=subject, grade=grade,
            level=tier, status="draft", duration=max(10, int(questions_per_layer) * 5))
        hw.created_at = batch_time
        session.flush()
        hw_svc.auto_compose(
            session, hw.id, build_tier_spec(tier, questions_per_layer),
            knowledge_points=weak_points or None)
        homework_ids[tier] = hw.id
    if selected_levels:
        deploy_tiered_homework(session, list(homework_ids.values()),
                               [tier for tier in selected_levels if tier in homework_ids])
    session.commit()
    return {
        "exam_id": exam.id if exam else None,
        "exam_name": exam.name if exam else "",
        "tiers": name_lists,
        "homework_ids": homework_ids,
        "levels": {tier: homework_ids.get(tier) for tier in ("A", "B", "C")},
        "status": "pending" if selected_levels else "draft",
    }


def deploy_tiered_homework(session, homework_ids: list[int], levels: list[str]) -> int:
    """只把选中层级从草稿置为进行中。"""
    wanted = set(level for level in levels if level in ("A", "B", "C"))
    count = 0
    for homework_id in homework_ids:
        hw = session.get(Homework, int(homework_id))
        if hw is not None and hw.level in wanted:
            hw.status = "pending"
            hw.completed_at = None
            count += 1
    session.flush()
    return count


def tier_effect_report(session, homework_ids: list[int]) -> list[dict]:
    """按层汇总完成率和平均分；无成绩时留空，不编造。"""
    by_level = defaultdict(lambda: {"students": 0, "submitted": 0,
                                    "scores": [], "question_count": 0})
    for homework_id in homework_ids:
        hw = session.get(Homework, int(homework_id))
        if hw is None or hw.level not in ("A", "B", "C"):
            continue
        bucket = by_level[hw.level]
        bucket["question_count"] += len(hw.questions)
        rows = (session.query(HomeworkScore)
                .filter(HomeworkScore.homework_id == hw.id).all())
        bucket["students"] += len(rows)
        bucket["submitted"] += sum(1 for r in rows if r.submitted and r.total_score is not None)
        bucket["scores"].extend(r.total_score for r in rows if r.total_score is not None)
    result = []
    for tier in ("A", "B", "C"):
        bucket = by_level.get(tier)
        if not bucket:
            continue
        total = bucket["submitted"]
        mean = round(sum(bucket["scores"]) / len(bucket["scores"]), 2) if bucket["scores"] else None
        result.append({
            "level": tier,
            "label": TIER_LABELS[tier],
            "students": bucket["students"],
            "submitted": total,
            "completion_rate": round(total / bucket["students"], 4) if bucket["students"] else None,
            "average": mean,
            "question_count": bucket["question_count"],
        })
    return result


def export_tier_effect_report(report: list[dict]) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading("分层教学效果报告", level=0)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["层级", "学生数", "已提交", "完成率", "平均分", "题数"]
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for row in report:
        cells = table.add_row().cells
        vals = [row["level"] + "·" + row["label"], row["students"], row["submitted"],
                f"{row['completion_rate'] * 100:.0f}%" if row["completion_rate"] is not None else "—",
                row["average"] if row["average"] is not None else "—",
                row["question_count"]]
        for i, value in enumerate(vals):
            cells[i].text = str(value)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
