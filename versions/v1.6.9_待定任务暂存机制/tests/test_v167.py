# -*- coding: utf-8 -*-
"""v1.6.7 导出优化、勾选导出、上传清空测试。"""

import io
import json

import pytest

from models.models import Question
from utils import question_service as qs


def _question(qid=1, content="题干", answer="答案", qtype="solution",
              difficulty=2, kps=None, analysis="解析", error_points="易错"):
    return Question(
        id=qid, content=content, question_type=qtype, difficulty=difficulty,
        knowledge_points=json.dumps(kps or ["知识点"], ensure_ascii=False),
        answer=answer, analysis=analysis, error_points=error_points,
        status="approved", source="manual", subject="数学")


def _word_text(data: bytes) -> str:
    from docx import Document
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


# ---------------------------------------------------------------------------
# 公式标记清理
# ---------------------------------------------------------------------------

def test_clean_formula_markers():
    assert qs._clean_formula_text("$72 - 9 - 21$") == "72 - 9 - 21"
    text = qs._clean_formula_text("计算 \\(x+1\\) 与 \\[y-2\\]")
    assert "$" not in text
    assert "\\(" not in text and "\\)" not in text
    assert "\\[" not in text and "\\]" not in text
    assert "x+1" in text and "y-2" in text


def test_export_word_strips_formula_markers():
    q = _question(content="$72 - 9 - 21$", answer="42")
    text = _word_text(qs.export_questions_word([q], with_answer=False))
    assert "$" not in text
    assert "72 - 9 - 21" in text


# ---------------------------------------------------------------------------
# 学生卷 / 教师卷
# ---------------------------------------------------------------------------

def test_student_paper_header_and_hidden_meta():
    q = _question(content="学生卷题干", answer="学生卷答案")
    text = _word_text(qs.export_questions_word([q], with_answer=False))
    assert "班级：____________" in text
    assert "姓名：____________" in text
    assert "得分：____________" in text
    assert "解答题" not in text
    assert "中等" not in text and "基础" not in text and "拓展" not in text
    assert "【知识点】" not in text
    assert "学生卷答案" not in text
    assert "1.　学生卷题干" in text


def test_student_paper_continuous_numbering():
    questions = [_question(i, content=f"题干{i}", answer=f"答案{i}") for i in (1, 2, 3)]
    text = _word_text(qs.export_questions_word(questions, with_answer=False))
    for number in ("1.　题干1", "2.　题干2", "3.　题干3"):
        assert number in text


def test_teacher_paper_keeps_meta_and_answer():
    q = _question(content="教师卷题干", answer="教师卷答案", kps=["核心知识点"])
    text = _word_text(qs.export_questions_word([q], with_answer=True))
    assert "1.（解答题，中等）" in text
    assert "【核心知识点】" in text
    assert "【答案】" in text and "教师卷答案" in text
    assert "【解析】" in text
    assert "【易错点】" in text


# ---------------------------------------------------------------------------
# 选择题选项拆分
# ---------------------------------------------------------------------------

CHOICE_CONTENT = "下列计算正确的是 A. 1+1=1 B. 2+2=4 C. 3+3=5 D. 4+4=6"


def test_split_choice_options_four_options():
    parts = qs._split_choice_options(CHOICE_CONTENT)
    assert len(parts) == 5
    assert parts[0] == "下列计算正确的是"
    assert parts[1] == "    A. 1+1=1"
    assert parts[2] == "    B. 2+2=4"
    assert parts[3] == "    C. 3+3=5"
    assert parts[4] == "    D. 4+4=6"


@pytest.mark.parametrize("text", [
    "只有题干没有选项",
    "A. 选项甲 B. 选项乙",  # 缺 C、D
    "B. 乙 A. 甲 C. 丙 D. 丁",  # 顺序错误
    "填空题：1 + 1 = ____",
    "A. 甲 B. 乙 C. 丙 A. 重复",  # 无 D 且 A 重复
])
def test_split_choice_options_no_false_split(text):
    parts = qs._split_choice_options(text)
    assert parts == [text]


def test_choice_options_separate_paragraphs():
    q = _question(content=CHOICE_CONTENT, answer="B", qtype="choice")
    student = Document_paragraphs(qs.export_questions_word([q], with_answer=False))
    assert student[1] == "班级：____________    姓名：____________    得分：____________"
    assert "1.　下列计算正确的是" in student
    assert "    A. 1+1=1" in student
    assert "    D. 4+4=6" in student
    teacher = Document_paragraphs(qs.export_questions_word([q], with_answer=True))
    assert "    A. 1+1=1" in teacher and "    D. 4+4=6" in teacher


def Document_paragraphs(data: bytes) -> list[str]:
    from docx import Document
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


# ---------------------------------------------------------------------------
# 勾选导出
# ---------------------------------------------------------------------------

def test_questions_by_selected_ids_keeps_order():
    questions = [_question(i) for i in (1, 2, 3, 4)]
    result = qs.questions_by_selected_ids(questions, [3, 1, 4])
    assert [q.id for q in result] == [3, 1, 4]


def test_questions_by_selected_ids_skips_missing():
    questions = [_question(i) for i in (1, 2)]
    result = qs.questions_by_selected_ids(questions, [2, 99, 1])
    assert [q.id for q in result] == [2, 1]


def test_export_selected_subset_renumbered():
    questions = [
        _question(1, content="第一题"),
        _question(2, content="第二题"),
        _question(3, content="第三题"),
    ]
    subset = qs.questions_by_selected_ids(questions, [3, 1])
    text = _word_text(qs.export_questions_word(subset, with_answer=False))
    assert "1.　第三题" in text
    assert "2.　第一题" in text
    assert "第二题" not in text


# ---------------------------------------------------------------------------
# AppTest
# ---------------------------------------------------------------------------

import pandas
from streamlit.testing.v1 import AppTest


def _isolated_app_code(patch_dir):
    """复用 smoke 的隔离脚本；假 AgGrid 和导出捕获通过 monkeypatch 注入。"""
    import importlib.util
    from pathlib import Path

    smoke_path = Path(__file__).with_name("test_app_smoke.py")
    spec = importlib.util.spec_from_file_location("test_app_smoke", smoke_path)
    smoke = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(smoke)

    db_file = patch_dir / "v167.db"
    code = smoke._isolated_app_code(db_file, patch_dir / "feature.json")
    return code, db_file


def _patch_aggrid_and_export(monkeypatch, patch_dir):
    """
    从测试侧替换模块属性（monkeypatch 测试结束自动还原，避免污染同进程其他测试）：
    AgGrid 按注入文件返回勾选行；导出函数把 Word 字节写到临时目录供断言。
    """
    import modules.lesson_plan as lesson_plan
    from utils import question_service

    def fake_aggrid(_df, gridOptions=None, key=None, **kwargs):
        rows = patch_dir.get("selected_rows", [])
        return {"data": _df, "selected_rows": pandas.DataFrame(rows)}

    monkeypatch.setattr(lesson_plan, "AgGrid", fake_aggrid)

    orig_export = question_service.export_questions_word

    def captured_export(questions, with_answer, title=None):
        data = orig_export(questions, with_answer, title)
        name = "teacher.docx" if with_answer else "student.docx"
        (patch_dir["dir"] / name).write_bytes(data)
        return data

    monkeypatch.setattr(question_service, "export_questions_word", captured_export)


def _seed_three_questions(db_file):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from models.models import Base

    eng = create_engine(f"sqlite:///{db_file.as_posix()}")
    Base.metadata.create_all(eng)
    session = sessionmaker(bind=eng)()
    for i in (1, 2, 3):
        qs.create_question(
            session,
            {"content": f"题干{i}", "question_type": "solution",
             "difficulty": 2, "knowledge_points": "[]",
             "answer": f"答案{i}", "analysis": "",
             "error_points": "", "verify": None},
            source="manual", status="approved", subject="数学")
    session.commit()
    session.close()
    return eng


def _goto_question_bank(at):
    at.sidebar.radio[0].set_value("📚 备课").run()
    return at.tabs[3]


def test_question_bank_empty_apptest(tmp_path, monkeypatch):
    code, db_file = _isolated_app_code(tmp_path)
    _patch_aggrid_and_export(monkeypatch, {"dir": tmp_path})
    at = AppTest.from_string(code, default_timeout=30)
    at.run()
    _goto_question_bank(at)
    assert not at.exception, [str(e) for e in at.exception]


def test_export_requires_selection_apptest(tmp_path, monkeypatch):
    code, db_file = _isolated_app_code(tmp_path)
    _patch_aggrid_and_export(monkeypatch, {"dir": tmp_path})
    eng = _seed_three_questions(db_file)
    at = AppTest.from_string(code, default_timeout=30)
    at.run()
    _goto_question_bank(at)
    assert not at.exception, [str(e) for e in at.exception]

    student_btn = next(b for b in at.button if b.key == "export_selected_student")
    teacher_btn = next(b for b in at.button if b.key == "export_selected_teacher")
    assert "请先勾选" in student_btn.label
    assert "请先勾选" in teacher_btn.label

    student_btn.click().run()
    assert not at.exception, [str(e) for e in at.exception]
    assert any("请先在表格中勾选要导出的题目" in str(x.value) for x in at.warning)
    assert all(x.key not in ("dl_exercise", "dl_teacher")
               for x in at.download_button)
    eng.dispose()


def test_export_selected_renumber_apptest(tmp_path, monkeypatch):
    code, db_file = _isolated_app_code(tmp_path)
    selected = [{"question_id": 3}, {"question_id": 1}, {"question_id": 2}]
    _patch_aggrid_and_export(
        monkeypatch, {"dir": tmp_path, "selected_rows": selected})
    eng = _seed_three_questions(db_file)

    at = AppTest.from_string(code, default_timeout=30)
    at.run()
    _goto_question_bank(at)
    assert not at.exception, [str(e) for e in at.exception]

    student_btn = next(b for b in at.button if b.key == "export_selected_student")
    teacher_btn = next(b for b in at.button if b.key == "export_selected_teacher")
    assert "已选 3 题" in student_btn.label
    assert "已选 3 题" in teacher_btn.label

    student_btn.click().run()
    teacher_btn.click().run()
    assert not at.exception, [str(e) for e in at.exception]

    dl_keys = {x.key for x in at.download_button}
    assert "dl_exercise" in dl_keys and "dl_teacher" in dl_keys

    student_text = _word_text((tmp_path / "student.docx").read_bytes())
    teacher_text = _word_text((tmp_path / "teacher.docx").read_bytes())
    # 勾选顺序 3、1、2 → 重新编号 1、2、3
    assert "1.　题干3" in student_text
    assert "2.　题干1" in student_text
    assert "3.　题干2" in student_text
    assert "1.（解答题，中等）" in teacher_text
    eng.dispose()


def _make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 有理数", level=1)
    doc.add_paragraph("这是 Word 讲义的正文内容。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_cleared_after_save_apptest(tmp_path, monkeypatch):
    code, db_file = _isolated_app_code(tmp_path)
    _patch_aggrid_and_export(monkeypatch, {"dir": tmp_path})
    at = AppTest.from_string(code, default_timeout=30)
    at.run()
    at.sidebar.radio[0].set_value("📚 备课").run()
    assert not at.exception, [str(e) for e in at.exception]

    file_type = next(x for x in at.selectbox if x.label == "来源类型")
    file_type.set_value("word").run()

    uploader = next(x for x in at.file_uploader if x.key == "material_upload")
    uploader.upload(
        "讲义.docx", _make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document").run()
    next(b for b in at.button if b.label == "提取并确认信息").click().run()
    assert not at.exception, [str(e) for e in at.exception]

    next(b for b in at.button if b.label == "确认").click().run()
    assert not at.exception, [str(e) for e in at.exception]
    next(b for b in at.button if b.key == "save_material").click().run()
    assert not at.exception, [str(e) for e in at.exception]

    # 保存成功并 rerun 后，上传框 widget 与 session_state 都应为空
    assert at.session_state.get("material_upload") is None
    uploader = next(x for x in at.file_uploader if x.key == "material_upload")
    assert not uploader.value
