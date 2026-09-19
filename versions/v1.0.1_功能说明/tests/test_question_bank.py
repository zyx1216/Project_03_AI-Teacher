# -*- coding: utf-8 -*-
"""题库测试：Excel/Word/文本导入 → 筛选 → 审核 → 编辑 → 删除。全部用内存库。"""

import io

import pandas as pd

from utils import question_service as qs
from utils import question_importer as qi


def _valid_record(content="题干", answer="答案", kp="知识点"):
    return {"row_no": 1, "content": content, "question_type": "解答题",
            "answer": answer, "analysis": "解析", "knowledge_points": kp,
            "difficulty": 2, "error_points": "", "problems": []}


def test_import_records_skips_missing_answer(session):
    records = [
        _valid_record("好题1", "答1"),
        {**_valid_record(), "content": "", "answer": "", "problems": ["缺题干", "缺答案"]},
        _valid_record("好题2", "答2"),
    ]
    result = qi.import_records(session, records)
    assert result == {"imported": 2, "rejected": 1}
    assert qs.list_questions(session, status="pending").__len__() == 2


def test_imported_questions_default_pending(session):
    qi.import_records(session, [_valid_record()])
    q = qs.list_questions(session)[0]
    assert q.status == "pending"
    assert q.source == "imported"


def test_excel_detect_columns_and_import(session):
    df = pd.DataFrame({
        "题目": ["1+1=", "2+2=", "无答案题"],
        "题型": ["填空题", "填空题", "填空题"],
        "答案": ["2", "4", None],
        "知识点": ["加法", "加法", "加法"],
    })
    detected = qi.detect_question_columns(df)
    mapping = detected["mapping"]
    assert mapping["content"] == "题目"
    assert mapping["answer"] == "答案"
    records = qi.records_from_excel(df, mapping)
    # 第三题缺答案被标记，但记录仍保留供预览
    bad = [r for r in records if not r["answer"]]
    assert bad and "缺答案" in bad[0]["problems"]
    result = qi.import_records(session, records)
    assert result["imported"] == 2 and result["rejected"] == 1


def test_text_import_numbered(session):
    text = "1. 第一题题干\n答案：甲\n解析：略\n\n2. 第二题题干\n答案：乙"
    records = qi.records_from_text(text)
    assert len(records) == 2
    result = qi.import_records(session, records)
    assert result["imported"] == 2


def test_text_import_without_answer_marked_red(session):
    records = qi.records_from_text("1. 只有题干没有答案标记的题")
    assert records[0]["problems"] == ["缺答案"]
    assert qi.import_records(session, records)["rejected"] == 1


def test_docx_import_roundtrip(session):
    from docx import Document
    doc = Document()
    doc.add_paragraph("1. 第一道题")
    doc.add_paragraph("答案：对")
    doc.add_paragraph("2. 第二道题")
    doc.add_paragraph("答案：错")
    buf = io.BytesIO()
    doc.save(buf)
    records = qi.records_from_docx(buf.getvalue())
    assert len(records) == 2
    assert qi.import_records(session, records)["imported"] == 2


def test_filters(session):
    qi.import_records(session, [
        {**_valid_record("选择", "A"), "question_type": "选择题", "difficulty": 1},
        {**_valid_record("解答", "过程"), "question_type": "解答题", "difficulty": 3},
    ])
    assert len(qs.list_questions(session, question_type="choice")) == 1
    assert len(qs.list_questions(session, difficulty=3)) == 1
    assert len(qs.list_questions(session, keyword="选择")) == 1


def test_approve_edit_delete(session):
    qi.import_records(session, [_valid_record()])
    q = qs.list_questions(session)[0]
    qs.approve_question(session, q.id)
    session.flush()
    assert qs.list_questions(session, status="approved")[0].id == q.id

    qs.update_question(session, q.id, content="改过的题干", answer="新答案")
    session.flush()
    edited = session.get(type(q), q.id)
    assert edited.content == "改过的题干" and edited.answer == "新答案"

    # 答案不允许改成空
    import pytest
    with pytest.raises(ValueError):
        qs.update_question(session, q.id, answer="")

    qs.delete_question(session, q.id)
    session.flush()
    assert qs.list_questions(session) == []


def test_export_word_two_modes(session):
    qi.import_records(session, [_valid_record("导出题干", "导出答案")])
    q = qs.list_questions(session)[0]
    student = qs.export_questions_word([q], with_answer=False)
    teacher = qs.export_questions_word([q], with_answer=True)
    from docx import Document
    s_text = "\n".join(p.text for p in Document(io.BytesIO(student)).paragraphs)
    t_text = "\n".join(p.text for p in Document(io.BytesIO(teacher)).paragraphs)
    assert "导出题干" in s_text and "导出答案" not in s_text
    assert "导出答案" in t_text