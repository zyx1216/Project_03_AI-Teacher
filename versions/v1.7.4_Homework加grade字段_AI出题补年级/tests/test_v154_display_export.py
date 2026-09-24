# -*- coding: utf-8 -*-
"""v1.5.4 展示与导出优化测试。"""

import io
import sys
from datetime import date
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "tests"))

from models.models import Base  # noqa: E402
from utils import (  # noqa: E402
    exam_service,
    homework_score_service,
    homework_service,
    question_service,
    student_service,
)

from test_app_smoke import _isolated_app_code  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402


def _make_exam_scores(session):
    """两班、两学生、一场数学考试。"""
    exam = exam_service.create_exam(
        session, "第一次月考", exam_date=date(2026, 9, 20),
        full_scores={"数学": 100})
    exam_service.import_scores(session, exam.id, [
        {"name": "学生甲", "class_name": "一班", "scores": {"数学": 90}},
        {"name": "学生乙", "class_name": "二班", "scores": {"数学": 80}},
    ])
    session.commit()
    return exam


def test_build_score_report_xlsx_all_classes(session):
    exam = _make_exam_scores(session)
    data = exam_service.build_score_report_xlsx(session, exam.id)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active

    assert ws["A1"].value == "全年级 第一次月考 成绩单"
    assert ws.merged_cells.ranges
    assert [ws.cell(2, col).value for col in range(1, 7)] == [
        "姓名", "学号", "班级", "数学", "总分", "班级排名"]

    student_rows = [
        [ws.cell(3, col).value for col in range(1, 7)],
        [ws.cell(4, col).value for col in range(1, 7)],
    ]
    assert ["学生甲", None, "一班", 90, 90, 1] in student_rows
    assert ["学生乙", None, "二班", 80, 80, 1] in student_rows

    assert ws.cell(5, 4).value == 85
    assert ws.cell(6, 4).value == 90
    assert ws.cell(7, 4).value == 80
    assert ws.freeze_panes == "A3"


def test_build_score_report_xlsx_one_class(session):
    exam = _make_exam_scores(session)
    data = exam_service.build_score_report_xlsx(session, exam.id, "一班")
    ws = load_workbook(io.BytesIO(data)).active

    assert ws["A1"].value == "一班 第一次月考 成绩单"
    assert ws.cell(3, 1).value == "学生甲"
    assert ws.cell(4, 1).value == "平均分"


def test_build_score_report_xlsx_empty_error(session):
    exam = exam_service.create_exam(session, "空考试", exam_date=date(2026, 9, 20))
    session.commit()
    with pytest.raises(ValueError, match="无法导出成绩单"):
        exam_service.build_score_report_xlsx(session, exam.id)


def _wrong_pdf_data(session):
    """同一道双知识点题在两份作业中分别被两名学生做错。"""
    data = question_service.validate_question({
        "content": "分析物体受力并求加速度",
        "question_type": "solution",
        "difficulty": 2,
        "knowledge_points": "力学、运动",
        "answer": "F=ma",
        "analysis": "先画受力图，再列牛顿第二定律。",
    })
    question = question_service.create_question(
        session, data, source="manual", status="approved", subject="数学")
    student1, _ = student_service.get_or_create_student(session, "学生甲", "一班")
    student2, _ = student_service.get_or_create_student(session, "学生乙", "一班")

    homeworks = []
    for student in (student1, student2):
        hw = homework_service.create_homework(
            session, f"作业{student.id}", class_name="一班", subject="数学")
        homework_service.add_questions(session, hw.id, [question.id])
        homework_score_service.save_answers(session, hw.id, [{
            "student_id": student.id,
            "question_id": question.id,
            "order_no": 1,
            "is_correct": False,
            "earned_score": 4,
            "error_type": "概念错误",
        }])
        homeworks.append(hw)
    session.commit()
    return question


def test_export_wrong_pdf_by_student(session):
    question = _wrong_pdf_data(session)
    rows = homework_score_service.list_wrong_answers(session)
    data = homework_score_service.export_wrong_pdf(
        rows, title="数学错题本", group_by="student")

    doc = pymupdf.open(stream=io.BytesIO(data), filetype="pdf")
    text = "".join(page.get_text() for page in doc)
    assert "数学错题本" in text
    assert "学生甲" in text and "学生乙" in text
    assert "F=ma" in text and "错误次数" in text
    assert text.count("分析物体受力并求加速度") == 2
    doc.close()


def test_export_wrong_pdf_by_knowledge(session):
    _wrong_pdf_data(session)
    rows = homework_score_service.list_wrong_answers(session)
    data = homework_score_service.export_wrong_pdf(
        rows, title="数学错题本", group_by="knowledge")

    doc = pymupdf.open(stream=io.BytesIO(data), filetype="pdf")
    text = "".join(page.get_text() for page in doc)
    assert "力学" in text and "运动" in text
    assert text.count("分析物体受力并求加速度") == 2
    assert text.count("错误次数：2") == 2
    doc.close()


def test_export_wrong_pdf_invalid_group(session):
    question = _wrong_pdf_data(session)
    rows = homework_score_service.list_wrong_answers(session)
    with pytest.raises(ValueError, match="student 或 knowledge"):
        homework_score_service.export_wrong_pdf(rows, group_by="class")


def test_score_report_apptest(tmp_path):
    from sqlalchemy import create_engine

    db_file = tmp_path / "score_report.db"
    engine = create_engine(f"sqlite:///{db_file.as_posix()}")
    Base.metadata.create_all(engine)
    session = SessionLocal = None
    from sqlalchemy.orm import sessionmaker
    db_session = sessionmaker(bind=engine)()
    exam = exam_service.create_exam(
        db_session, "第一次月考", exam_date=date(2026, 9, 20),
        full_scores={"数学": 100})
    exam_service.import_scores(db_session, exam.id, [{
        "name": "学生甲", "class_name": "一班", "scores": {"数学": 90}}])
    db_session.commit()
    db_session.close()

    at = AppTest.from_string(
        _isolated_app_code(db_file, tmp_path / "feature.json"),
        default_timeout=30)
    at.run()
    at.sidebar.radio[0].set_value("📊 学情").run()
    next(b for b in at.button if b.key == "make_score_report").click().run()

    assert not at.exception, [str(e) for e in at.exception]
    assert any(b.key == "download_score_report" for b in at.download_button)
    engine.dispose()


def test_wrong_pdf_apptest(tmp_path):
    from test_app_smoke import _seed_homework_data_db

    db_file = tmp_path / "wrong_pdf.db"
    engine = _seed_homework_data_db(db_file)
    feature_file = tmp_path / "feature.json"
    feature_file.write_text('{"wrong_book_subject": "物理"}', encoding="utf-8")

    at = AppTest.from_string(
        _isolated_app_code(db_file, feature_file), default_timeout=30)
    at.run()
    at.sidebar.radio[0].set_value("📝 学业测评").run()
    next(x for x in at.radio if x.key == "wb_pdf_group").set_value("按知识点").run()
    next(b for b in at.button if b.key == "wb_export_pdf").click().run()

    assert not at.exception, [str(e) for e in at.exception]
    assert any(b.key == "wb_export_pdf_dl" for b in at.download_button)
    assert any(b.key == "wb_export" for b in at.button)
    engine.dispose()
