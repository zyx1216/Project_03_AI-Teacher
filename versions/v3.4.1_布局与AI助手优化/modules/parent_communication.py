# -*- coding: utf-8 -*-
"""家校沟通标签页（v3.4.0）。

只做内容生成与导出，不做消息推送：
成绩单 / 学习建议 / 班级报告与批量导出 / 历史留档。
挂在「📊 学情 → 学生画像」下作为第三个标签。
"""
from __future__ import annotations

import streamlit as st

from models.models import Student
from utils import llm_client
from utils import parent_communication_service as pcs
from utils import parent_report_service as prs
from utils.db import SessionLocal
from utils import student_service as ss
from utils import exam_service as es
from utils.app_config import SUBJECT_NAMES


def _chat_func():
    """配置了内容模型时返回 chat 回调，否则 None（走规则兜底）。"""
    if not llm_client.is_content_configured():
        return None
    return lambda system, user: llm_client.chat_content(system, user)


def _student_options(session) -> dict:
    """{显示名: 学生 id}。"""
    rows = ss.list_students(session)
    return {f"{s.name}（{s.class_name or '未分班'}）": s.id for s in rows}


def _exam_options(session) -> dict:
    """{显示名: 考试 id}，含"不指定考试"。"""
    labels = {"不指定考试（用最近成绩）": None}
    for exam in es.list_exams(session):
        labels[f"{exam.name}（{exam.exam_date}）"] = exam.id
    return labels


# ---------------------------------------------------------------------------
# 成绩单
# ---------------------------------------------------------------------------

def _report_card_tab() -> None:
    with SessionLocal() as session:
        students = _student_options(session)
        exams = _exam_options(session)
    if not students:
        st.info("还没有学生，先到「学生管理」导入名单。")
        return
    c1, c2 = st.columns(2)
    student_label = c1.selectbox("学生", list(students), key="family_card_student")
    exam_label = c2.selectbox("考试", list(exams), key="family_card_exam")
    if st.button("📄 生成成绩单", key="family_card_build", type="primary"):
        with SessionLocal() as session:
            card = pcs.build_report_card(session, students[student_label],
                                         exam_id=exams[exam_label])
            st.session_state["family_card_data"] = card
            st.session_state["family_card_student_id"] = students[student_label]
    card = st.session_state.get("family_card_data")
    if not card:
        st.caption("选好学生和考试后点「生成成绩单」。")
        return
    st.text(pcs.report_card_text(card))
    c1, c2, c3 = st.columns(3)
    if c1.button("💾 留档成绩单", key="family_card_save"):
        with SessionLocal() as session:
            pcs.save_record(session, card["student_id"],
                            (card.get("exam") or {}).get("exam_id"), "report",
                            pcs.report_card_text(card))
            session.commit()
        st.toast("成绩单已留档。")
    pdf = None
    with SessionLocal() as session:
        report = prs.generate_parent_report(
            session, card["student_id"],
            exam_id=(card.get("exam") or {}).get("exam_id"), chat_func=None)
        pdf = prs.export_parent_report_pdf(report)
    c2.download_button("⬇️ 导出 PDF", pdf,
                       file_name=f"{card['name']}_家校成绩单.pdf",
                       key="family_card_pdf")
    import pandas as pd
    from utils import export_service
    rows = [{"学科": subject, "得分": item.get("score"),
             "满分": item.get("full"), "分数段": item.get("band"),
             "班级排名": item.get("class_rank"),
             "年级排名": item.get("grade_rank")}
            for subject, item in (card.get("subjects") or {}).items()]
    if rows:
        c3.download_button(
            "⬇️ 导出 Excel", export_service.styled_excel(
                pd.DataFrame(rows), sheet_name="成绩单",
                title=f"{card['name']}（{card['class_name']}）成绩单"),
            file_name=f"{card['name']}_成绩单.xlsx", key="family_card_excel")


# ---------------------------------------------------------------------------
# 学习建议
# ---------------------------------------------------------------------------

def _advice_tab() -> None:
    with SessionLocal() as session:
        students = _student_options(session)
    if not students:
        st.info("还没有学生，先到「学生管理」导入名单。")
        return
    c1, c2 = st.columns(2)
    student_label = c1.selectbox("学生", list(students), key="family_advice_student")
    subject = c2.selectbox("学科", ["全部学科"] + list(SUBJECT_NAMES),
                           key="family_advice_subject")
    if st.button("💡 生成学习建议", key="family_advice_build", type="primary"):
        with SessionLocal() as session:
            advice = pcs.generate_advice(
                session, students[student_label],
                subject=None if subject == "全部学科" else subject,
                chat_func=_chat_func())
            st.session_state["family_advice_data"] = advice
    advice = st.session_state.get("family_advice_data")
    if not advice:
        st.caption("选好学生和学科后点「生成学习建议」。")
        return
    if advice.get("data_notice"):
        st.caption(advice["data_notice"])
    st.text(pcs.advice_text(advice))
    c1, c2 = st.columns(2)
    if c1.button("💾 留档建议", key="family_advice_save"):
        with SessionLocal() as session:
            pcs.save_record(session, advice["student_id"], None, "advice",
                            pcs.advice_text(advice))
            session.commit()
        st.toast("学习建议已留档。")
    if c2.button("🔄 重新生成", key="family_advice_regen"):
        with SessionLocal() as session:
            st.session_state["family_advice_data"] = pcs.generate_advice(
                session, advice["student_id"],
                subject=None if subject == "全部学科" else subject,
                chat_func=_chat_func())
        st.rerun()


# ---------------------------------------------------------------------------
# 班级报告与批量导出
# ---------------------------------------------------------------------------

def _class_tab() -> None:
    with SessionLocal() as session:
        classes = ss.list_classes(session)
        exams = _exam_options(session)
    if not classes:
        st.info("还没有班级，先到「学生管理」导入名单。")
        return
    c1, c2 = st.columns(2)
    class_name = c1.selectbox("班级", classes, key="family_class_pick")
    exam_label = c2.selectbox("考试", list(exams), key="family_class_exam")

    if st.button("🏫 生成班级整体报告", key="family_class_build"):
        with SessionLocal() as session:
            try:
                report = pcs.build_class_report(
                    session, class_name, exam_id=exams[exam_label])
            except ValueError as exc:
                st.warning(str(exc))
                return
            st.session_state["family_class_data"] = report
    report = st.session_state.get("family_class_data")
    if report:
        text = pcs.class_report_text(report)
        st.text(text)
        c1, c2 = st.columns(2)
        c1.download_button("⬇️ 导出 Word", _class_report_docx(report),
                           file_name=f"{class_name}_家长会班级报告.docx",
                           key="family_class_word")
        if c2.button("💾 留档班级报告", key="family_class_save"):
            with SessionLocal() as session:
                rows = ss.list_students(session, class_name=class_name)
                if not rows:
                    st.warning("该班还没有学生，无法留档。")
                else:
                    pcs.save_record(session, rows[0].id, None, "report", text)
                    session.commit()
                    st.toast(f"班级报告已留档（归到 {rows[0].name}）。")

    st.divider()
    st.markdown("**批量导出整班家校成绩单（ZIP）**")
    if st.button("📦 批量导出", key="family_batch_export"):
        with SessionLocal() as session:
            students = ss.list_students(session, class_name=class_name)
            if not students:
                st.warning("该班还没有学生。")
                return
            bar = st.progress(0.0, text="正在生成…")

            def _progress(done, total):
                bar.progress(done / max(1, total),
                             text=f"正在生成 {done}/{total} 份")

            data = pcs.export_family_zip(
                session, [s.id for s in students],
                exam_id=exams[exam_label], progress=_progress)
        st.session_state["family_zip_data"] = data
        st.session_state["family_zip_class"] = class_name
        st.toast("批量导出完成。")
    if st.session_state.get("family_zip_data"):
        st.download_button(
            "⬇️ 下载 ZIP", st.session_state["family_zip_data"],
            file_name=f"{st.session_state.get('family_zip_class', '班级')}"
                      "_家校成绩单.zip", key="family_zip_download")


def _class_report_docx(report: dict) -> bytes:
    """班级整体报告导出 Word（纯文本分段，不依赖模板）。"""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading(f"{report['class_name']} 家长会班级报告", level=1)
    if report.get("exam"):
        doc.add_paragraph(f"考试：{report['exam']['exam_name']}　"
                          f"{report['exam']['date']}")
    doc.add_paragraph(f"学生人数：{report['student_count']}")
    for item in report.get("subject_stats") or []:
        doc.add_paragraph(
            f"{item['subject']}：均分 {item['avg']}／{item['full']}，"
            f"及格率 {item['pass_rate']}%，优秀率 {item['excellent_rate']}%")
    if report.get("weak_points"):
        doc.add_paragraph("班级薄弱知识点：" + "、".join(
            f"{x['knowledge_point']}（{x['avg_rate']}%）"
            for x in report["weak_points"]))
    for tip in report.get("suggestions") or []:
        doc.add_paragraph(f"- {tip}")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 历史留档
# ---------------------------------------------------------------------------

def _records_tab() -> None:
    with SessionLocal() as session:
        records = pcs.list_records(session, limit=50)
    if not records:
        st.caption("还没有留档记录。")
        return
    labels = {r["id"]: f"{r['created_at']:%Y-%m-%d %H:%M}｜{r['student_name']}"
                      f"｜{r['type']}" for r in records}
    picked = st.selectbox("选择记录", list(labels),
                          format_func=lambda x: labels[x],
                          key="family_record_pick")
    row = next(r for r in records if r["id"] == picked)
    st.text(row["content"])


def tab_family() -> None:
    """家校沟通：成绩单 / 学习建议 / 班级报告与批量 / 历史留档。"""
    st.subheader("👨👩👧 家校沟通")
    tabs = st.tabs(["📄 成绩单", "💡 学习建议", "🏫 班级报告与批量",
                    "🗂️ 历史留档"])
    with tabs[0]:
        _report_card_tab()
    with tabs[1]:
        _advice_tab()
    with tabs[2]:
        _class_tab()
    with tabs[3]:
        _records_tab()
