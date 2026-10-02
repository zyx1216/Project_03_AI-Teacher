# -*- coding: utf-8 -*-
"""v2.9.1 单元整体设计服务。"""

from __future__ import annotations
import io
import json
from datetime import datetime

from models.models import Homework, LessonPlan, Textbook, UnitLesson, UnitPlan
from utils import homework_service as hw_svc
from utils import homework_stats as hstat
from utils import lesson_service as lesson_svc
from utils import llm_client


def list_unit_candidates(session, subject=None, grade=None, textbook_id=None):
    query = session.query(Textbook)
    if subject:
        query = query.filter(Textbook.subject == subject)
    if grade:
        query = query.filter(Textbook.grade == grade)
    if textbook_id:
        query = query.filter(Textbook.id == int(textbook_id))
    result = []
    for book in query.order_by(Textbook.created_at.desc(), Textbook.id.desc()).all():
        try:
            chapters = json.loads(book.chapter_info or "[]")
        except (TypeError, json.JSONDecodeError):
            chapters = []
        for chapter in chapters if isinstance(chapters, list) else []:
            title = str(chapter.get("title") or "").strip() if isinstance(chapter, dict) else str(chapter).strip()
            if title:
                result.append({"textbook_id": book.id, "textbook_name": book.name,
                               "chapter": title, "grade": book.grade or ""})
    return result


def _fallback_objectives(name, lesson_count=6):
    return {"knowledge_skills": f"掌握{name}的核心概念和基本方法。",
            "process_methods": "经历观察、操作、归纳和迁移应用的过程。",
            "emotion_values": "培养主动探究和合作交流的学习习惯。",
            "key_points": "单元核心概念与基本技能。",
            "difficult_points": "知识迁移与综合应用。",
            "lesson_hours": int(max(1, lesson_count)),
            "lesson_titles": [f"{name} 第{i}课时" for i in range(1, int(max(1, lesson_count)) + 1)]}


def generate_unit_objectives(session, name, subject, grade="", chapter="",
                             textbook_id=None, lesson_count=6, chat_func=None):
    fallback = _fallback_objectives(name, lesson_count)
    try:
        func = chat_func
        if func is None and llm_client.is_content_configured():
            func = llm_client.chat_content
        if func is not None:
            raw = func("你是单元整体设计教研员，只输出 JSON。",
                       f"学科：{subject}；年级：{grade}；单元：{name}；章节：{chapter}；课时数：{lesson_count}。"
                       "输出 knowledge_skills/process_methods/emotion_values/key_points/"
                       "difficult_points/lesson_hours/lesson_titles。", temperature=0.4)
            text = str(raw).strip().strip("`")
            if text.startswith("json"):
                text = text[4:].strip()
            parsed = json.loads(text)
            if isinstance(parsed, dict):
                fallback.update(parsed)
                fallback["lesson_hours"] = int(fallback.get("lesson_hours") or lesson_count)
    except Exception:  # noqa: BLE001
        fallback["fallback_notice"] = "AI 未完成结构化目标，已使用确定性单元目标。"
    return fallback


def split_unit_lessons(objectives):
    titles = objectives.get("lesson_titles") if isinstance(objectives, dict) else None
    if not isinstance(titles, list) or not titles:
        count = int((objectives or {}).get("lesson_hours") or 6)
        titles = [f"第{i}课时" for i in range(1, max(1, count) + 1)]
    return [{"lesson_index": index, "title": str(title).strip()}
            for index, title in enumerate(titles, start=1) if str(title).strip()]


def create_unit_plan(session, name, subject, grade="", textbook_id=None,
                     chapter="", objectives=None, lessons=None):
    objectives = objectives or _fallback_objectives(name)
    lessons = lessons or split_unit_lessons(objectives)
    unit = UnitPlan(name=str(name).strip(), subject=subject, grade=grade,
                    textbook_id=textbook_id, chapter=chapter,
                    objectives=json.dumps(objectives, ensure_ascii=False),
                    lesson_count=len(lessons), status="active")
    session.add(unit)
    session.flush()
    for item in lessons:
        session.add(UnitLesson(unit_id=unit.id,
                               lesson_index=int(item["lesson_index"]),
                               title=str(item["title"]).strip(),
                               status="pending"))
    session.flush()
    return unit


def _fallback_plan(title, subject, grade, chapter):
    plan = lesson_svc.empty_plan()
    plan["subject"] = subject
    plan["objectives"]["knowledge"] = f"理解{title}的核心知识。"
    plan["objectives"]["process"] = "通过例题和练习形成方法。"
    plan["objectives"]["emotion"] = "保持主动学习。"
    plan["key_points"] = title
    plan["difficult_points"] = "综合应用"
    plan["process"][0]["content"] = f"导入{title}。"
    plan["process"][1]["content"] = f"讲解{title}的关键方法。"
    plan["process"][2]["content"] = "完成课堂练习并反馈。"
    plan["process"][3]["content"] = "总结本课结构。"
    plan["process"][4]["content"] = "布置分层作业。"
    return plan


def generate_lesson_content(session, lesson: UnitLesson, subject, grade="",
                            chapter="", chat_func=None):
    try:
        func = chat_func
        if func is None and llm_client.is_content_configured():
            func = llm_client.chat_content
        if func is not None:
            raw = func(
                "你是中小学教师，输出 JSON 教案。",
                f"学科：{subject}；年级：{grade}；章节：{chapter}；课时：{lesson.title}。",
                temperature=0.5)
            parsed = lesson_svc.parse_lesson_plan(str(raw or ""))
            if parsed:
                plan = parsed
            else:
                plan = _fallback_plan(lesson.title, subject, grade, chapter)
        else:
            plan = _fallback_plan(lesson.title, subject, grade, chapter)
    except Exception:  # noqa: BLE001
        plan = _fallback_plan(lesson.title, subject, grade, chapter)
    saved = lesson_svc.save_plan(session, lesson.title, plan, grade=grade,
                                 chapter=chapter, subject=subject)
    lesson.plan_id = saved.id
    lesson.status = "pending"
    session.flush()
    return saved


def generate_all_lessons(session, unit_id, chat_func=None):
    unit = session.get(UnitPlan, int(unit_id))
    if unit is None:
        raise ValueError("单元不存在。")
    objectives = json.loads(unit.objectives or "{}")
    for lesson in unit.lessons:
        generate_lesson_content(session, lesson, unit.subject, unit.grade or "",
                                unit.chapter or "", chat_func=chat_func)
    session.flush()
    return len(unit.lessons)


def generate_unit_test(session, unit_id, class_name=None, count=10):
    unit = session.get(UnitPlan, int(unit_id))
    if unit is None:
        raise ValueError("单元不存在。")
    counts = {"choice": max(1, round(count * .4)),
              "fill": max(1, round(count * .3)),
              "solution": max(1, count - max(1, round(count * .4)) - max(1, round(count * .3)))}
    spec = {"counts": counts, "difficulty_ratio": {1: 30, 2: 50, 3: 20}}
    hw = hw_svc.create_homework(session, f"{unit.name}·单元测试",
                                homework_type="exam", class_name=class_name,
                                subject=unit.subject, grade=unit.grade,
                                chapter=unit.chapter, status="pending")
    result = hw_svc.auto_compose(session, hw.id, spec, knowledge_points=None)
    session.flush()
    return {"homework_id": hw.id, "name": hw.name, "compose": result}


def update_lesson_status(session, lesson_id, status):
    if status not in ("pending", "in_progress", "completed"):
        raise ValueError("课时状态只能是未开始/进行中/已完成。")
    lesson = session.get(UnitLesson, int(lesson_id))
    if lesson is None:
        raise ValueError("课时不存在。")
    lesson.status = status
    lesson.completed_at = datetime.now() if status == "completed" else None
    unit = lesson.unit
    if unit and unit.lessons:
        done = sum(1 for item in unit.lessons if item.status == "completed")
        unit.status = "completed" if done == len(unit.lessons) else "active"
    session.flush()
    return lesson


def copy_unit_plan(session, unit_id, new_name=None):
    source = session.get(UnitPlan, int(unit_id))
    if source is None:
        raise ValueError("单元不存在。")
    copied = UnitPlan(name=(new_name or f"{source.name}（副本）"), subject=source.subject,
                      grade=source.grade, textbook_id=source.textbook_id,
                      chapter=source.chapter, objectives=source.objectives,
                      lesson_count=source.lesson_count, status="draft")
    session.add(copied)
    session.flush()
    for item in source.lessons:
        session.add(UnitLesson(unit_id=copied.id, lesson_index=item.lesson_index,
                               title=item.title, plan_id=item.plan_id,
                               ppt_path=item.ppt_path, status="pending"))
    session.flush()
    return copied


def export_unit_word(session, unit_id):
    from docx import Document
    unit = session.get(UnitPlan, int(unit_id))
    if unit is None:
        raise ValueError("单元不存在。")
    doc = Document()
    doc.add_heading(f"{unit.name} · 单元教案合集", level=0)
    objectives = json.loads(unit.objectives or "{}")
    doc.add_heading("单元目标", level=1)
    for key, value in objectives.items():
        if isinstance(value, list):
            value = "、".join(str(x) for x in value)
        doc.add_paragraph(f"{key}：{value}")
    for lesson in unit.lessons:
        doc.add_heading(f"{lesson.lesson_index}. {lesson.title}", level=1)
        plan = session.get(LessonPlan, lesson.plan_id) if lesson.plan_id else None
        if plan is None:
            doc.add_paragraph("尚未生成教案。")
            continue
        data = lesson_svc.load_plan(plan)
        doc.add_paragraph("教学目标：" + json.dumps(data.get("objectives", {}), ensure_ascii=False))
        for step in data.get("process", []):
            doc.add_paragraph(f"{step.get('stage', '')}：{step.get('content', '')}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
