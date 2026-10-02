# -*- coding: utf-8 -*-
"""多格式导出服务（v1.9.0）。

- lesson_plan_pdf：用 PyMuPDF Story 直出教案 PDF（不依赖 Word 转 PDF）；
- homework_latex：作业导出为 ctexart LaTeX 源码；
- analysis_report_ppt：成绩分析导出为 PPT（python-pptx）。

原有 Word / Excel / ZIP 导出保留在各自服务里，这里不重复。
"""

from __future__ import annotations

import html
import io
from pathlib import Path


# ---------------------------------------------------------------------------
# 教案 PDF（PyMuPDF Story）
# ---------------------------------------------------------------------------

def _pdf_archive():
    """微软雅黑存在时嵌入；否则用 PyMuPDF 内置中文字体。"""
    import pymupdf
    font = Path(r"C:\Windows\Fonts\msyh.ttc")
    if not font.exists():
        return None, "body { font-family: china-s, sans-serif; }"
    archive = pymupdf.Archive()
    archive.add(font.read_bytes(), "yh.ttc")
    return archive, (
        "@font-face { font-family: yh; src: url(yh.ttc); }"
        "body { font-family: yh, china-s, sans-serif; }")


def _esc(value) -> str:
    return html.escape(str(value or ""), quote=False).replace("\n", "<br>")


def lesson_plan_html(lesson, plan: dict) -> str:
    """教案 PDF 的 HTML；分页交给 Story。"""
    objectives = plan.get("objectives", {})
    parts = [
        "<html><head><meta charset='utf-8'><style>",
        """
        @page { margin: 48px 46px; }
        body { font-size: 10.5pt; line-height: 1.6; color: #222; }
        h1 { color: #1f4e79; font-size: 22pt; margin-bottom: 4px; }
        h2 { color: #1f4e79; font-size: 14pt; margin: 18px 0 8px;
             border-bottom: 1px solid #9eadcc; padding-bottom: 3px;
             page-break-after: avoid; }
        .meta { color: #555; margin-bottom: 16px; }
        .block { page-break-inside: avoid; margin: 6px 0; }
        .stage { page-break-inside: avoid; border: 1px solid #c8d4e6;
                 border-radius: 6px; padding: 8px 10px; margin: 9px 0; }
        .stage-t { color: #1f4e79; font-weight: bold; }
        """,
        "</style></head><body>",
        f"<h1>{_esc(lesson.title)}</h1>",
        f"<div class='meta'>学科：{_esc(plan.get('subject'))}　"
        f"年级：{_esc(lesson.grade or '未指定')}　"
        f"章节：{_esc(lesson.chapter or '未指定')}</div>",
        "<h2>教学目标</h2>",
        f"<div class='block'>知识与技能：{_esc(objectives.get('knowledge'))}</div>",
        f"<div class='block'>过程与方法：{_esc(objectives.get('process'))}</div>",
        f"<div class='block'>情感态度：{_esc(objectives.get('emotion'))}</div>",
        f"<h2>教学重点</h2>{_esc(plan.get('key_points'))}",
        f"<h2>教学难点</h2>{_esc(plan.get('difficult_points'))}",
        "<h2>教学过程</h2>",
    ]
    for step in plan.get("process", []):
        parts.append(
            "<div class='stage'>"
            f"<div class='stage-t'>{_esc(step.get('stage'))}"
            f"（{_esc(step.get('minutes'))}分钟）</div>"
            f"<div>{_esc(step.get('content'))}</div></div>")
    assignment = next(
        (s.get("content") for s in plan.get("process", [])
         if s.get("stage") == "作业布置"), "")
    parts.append(f"<h2>作业布置</h2>{_esc(assignment or '见课堂安排')}")
    if plan.get("board_design"):
        parts.append(f"<h2>板书设计</h2>{_esc(plan.get('board_design'))}")
    parts.append("</body></html>")
    return "".join(parts)


def lesson_plan_pdf(lesson, plan: dict) -> bytes:
    """教案导出为 PDF（PyMuPDF Story 直出，支持中文）。"""
    import pymupdf
    archive, font_css = _pdf_archive()
    page_html = lesson_plan_html(lesson, plan).replace(
        "</style>", font_css + "</style>", 1)
    story = pymupdf.Story(html=page_html, archive=archive)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    mediabox = pymupdf.paper_rect("a4")
    while True:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(mediabox + (46, 48, -46, -48))
        story.draw(dev)
        writer.end_page()
        if not more:
            break
    writer.close()
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 作业 LaTeX
# ---------------------------------------------------------------------------

def _latex_text(value) -> str:
    """普通文本最小转义；保留题干里已有的 $...$ 公式。"""
    text = str(value or "")
    # 按公式段切分，只转义非公式部分。
    parts = []
    cursor = 0
    in_math = False
    i = 0
    while i < len(text):
        if text[i] == "$":
            seg = text[cursor:i]
            parts.append(_latex_escape(seg) if not in_math else seg)
            parts.append("$")
            in_math = not in_math
            cursor = i + 1
        i += 1
    tail = text[cursor:]
    parts.append(_latex_escape(tail) if not in_math else tail)
    return "".join(parts).replace("\n", "\\\\\n")


def _latex_escape(seg: str) -> str:
    for ch in ("\\", "&", "%", "#", "_", "{", "}"):
        seg = seg.replace(ch, "\\" + ch)
    return seg


def homework_latex(homework, pairs) -> str:
    """作业导出为 ctexart LaTeX 源码。pairs 为 [(关联行, 题目)]。"""
    lines = [
        "\\documentclass{ctexart}",
        "\\usepackage[a4paper,margin=2.2cm]{geometry}",
        "\\usepackage{amsmath,amssymb}",
        "\\begin{document}",
        f"\\section*{{{_latex_text(homework.name)}}}",
    ]
    current_type = None
    type_names = {"choice": "一、选择题", "fill": "二、填空题",
                  "judge": "三、判断题", "solution": "四、解答题"}
    for index, (_link, q) in enumerate(pairs, start=1):
        if q.question_type != current_type:
            current_type = q.question_type
            lines.append(f"\\subsection*{{{type_names.get(current_type, '题目')}}}")
        lines.append(f"{index}. {_latex_text(q.content)}")
        if q.answer:
            lines.append("")
    lines.append("\\end{document}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 成绩分析 PPT
# ---------------------------------------------------------------------------

def analysis_report_ppt(session, source_id: int, source_type: str) -> bytes:
    """成绩分析导出 PPT。

    source_type="exam" 走 exam_service.analyze_exam；
    source_type="homework" 走 homework_score_service.analyze_homework。
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title_slide(title, subtitle):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle

    def add_bullets(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0] if bullets else ""
        for text in bullets[1:]:
            frame.add_paragraph().text = text
        return slide

    def add_table(title, headers, rows):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        rows_n = len(rows) + 1
        cols_n = len(headers)
        table_shape = slide.shapes.add_table(
            rows_n, cols_n, Inches(0.6), Inches(1.4),
            Inches(12.1), Inches(min(5.2, 0.45 * rows_n)))
        table = table_shape.table
        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = str(h)
            cell.text_frame.paragraphs[0].font.bold = True
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                table.cell(i, j).text = str(value)

    if source_type == "exam":
        from utils import exam_service
        data = exam_service.analyze_exam(session, source_id)
        exam = data["exam"]
        add_title_slide(
            f"考试分析：{exam.name}",
            f"考试日期：{exam.exam_date}")
        stats = data["total_stats"]
        add_bullets("总体指标", [
            f"参考人数：{stats.get('count')}",
            f"平均分：{stats.get('mean')}",
            f"最高分：{stats.get('max')}",
            f"最低分：{stats.get('min')}",
            f"及格率：{stats.get('pass_rate')}%",
            f"优秀率：{stats.get('excellent_rate')}%",
        ])
        subject_rows = []
        for subject in data["subjects"]:
            s = data["subject_stats"][subject]
            subject_rows.append([
                subject, s.get("count"), s.get("mean"),
                s.get("pass_rate"), s.get("excellent_rate")])
        add_table("各科统计",
                  ["科目", "人数", "平均分", "及格率", "优秀率"],
                  subject_rows)
        add_bullets("结论与建议", [
            "结合各科及格率、优秀率确定后续教学重点。",
            "关注进退步明显的学生，分层布置巩固练习。"])
    else:
        from utils import homework_score_service as hscore
        data = hscore.analyze_homework(session, source_id)
        name = data.get("name", f"作业{source_id}")
        add_title_slide(f"作业分析：{name}", "学业测评")
        stats = data.get("stats", {})
        add_bullets("总体指标", [
            f"提交人数：{stats.get('submitted_count')}",
            f"平均分：{stats.get('mean')}",
            f"最高分：{stats.get('max')}",
            f"最低分：{stats.get('min')}",
        ])
        add_bullets("结论与建议", [
            "依据错题知识点安排专项复习。",
            "对未提交学生及时跟进。"])

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()

# ---------------------------------------------------------------------------
# v3.4.0：Excel / PDF / Word 导出优化
# ---------------------------------------------------------------------------

def _auto_width(ws, headers: list[str], rows: list[list], cap: int = 40) -> None:
    """按内容长度自适应列宽（中文按 2 个字符宽度估）。"""
    from openpyxl.utils import get_column_letter

    for index, header in enumerate(headers, start=1):
        width = sum(2 if ord(ch) > 127 else 1 for ch in str(header or ""))
        for row in rows:
            if index - 1 >= len(row):
                continue
            value = row[index - 1]
            text = "" if value is None else str(value)
            width = max(width, sum(2 if ord(ch) > 127 else 1 for ch in text))
        ws.column_dimensions[get_column_letter(index)].width = min(
            max(10, width + 3), cap)


def styled_excel(df, *, columns=None, header_map=None, sheet_name: str = "数据",
                 title: str | None = None, validations=None) -> bytes:
    """带格式的 Excel：表头加粗、列宽自适应、冻结表头，可选数据验证。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = [str(c) for c in (columns if columns is not None else list(df.columns))]
    header_map = header_map or {}
    display_headers = [str(header_map.get(h, h)) for h in headers]
    data = df[headers].values.tolist() if hasattr(df, "columns") else list(df)

    wb = Workbook()
    ws = wb.active
    ws.title = str(sheet_name)[:31] or "数据"
    start = 1
    if title:
        ws.cell(row=1, column=1, value=str(title)).font = Font(
            name="微软雅黑", size=14, bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(1, len(display_headers)))
        start = 2
    header_font = Font(name="微软雅黑", bold=True)
    fill = PatternFill("solid", fgColor="D9E1F2")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, text in enumerate(display_headers, start=1):
        cell = ws.cell(row=start, column=col, value=text)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = center
    for r, row_values in enumerate(data, start=start + 1):
        for c, value in enumerate(row_values, start=1):
            ws.cell(row=r, column=c, value=value).font = Font(name="微软雅黑")
    _auto_width(ws, display_headers, data)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if data:
        ws.auto_filter.ref = (f"A{start}:"
                              f"{get_column_letter(len(display_headers))}"
                              f"{start + len(data)}")
    for spec in validations or []:
        formula = '"' + ",".join(str(x) for x in spec.get("options") or []) + '"'
        if formula == '""':
            continue
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(str(spec.get("range") or ""))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def multi_sheet_excel(sheets: list[dict]) -> bytes:
    """多 Sheet 导出：每个元素 {name, df, columns?, header_map?, title?}。"""
    from openpyxl import load_workbook

    if not sheets:
        raise ValueError("至少需要一个 Sheet。")
    first = styled_excel(
        sheets[0]["df"], columns=sheets[0].get("columns"),
        header_map=sheets[0].get("header_map"),
        sheet_name=sheets[0].get("name") or "Sheet1",
        title=sheets[0].get("title"))
    wb = load_workbook(io.BytesIO(first))
    for spec in sheets[1:]:
        one = styled_excel(
            spec["df"], columns=spec.get("columns"),
            header_map=spec.get("header_map"),
            sheet_name=spec.get("name") or "Sheet",
            title=spec.get("title"))
        extra = load_workbook(io.BytesIO(one))
        source = extra[extra.sheetnames[0]]
        target = wb.create_sheet(source.title)
        for row in source.iter_rows():
            for cell in row:
                target[cell.coordinate] = cell.value
        for key, dim in source.column_dimensions.items():
            target.column_dimensions[key].width = dim.width
        extra.close()
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def fill_school_template(template_path, mapping: dict) -> bytes:
    """按表头文字把数据填进学校 Excel 模板，返回新文件字节。"""
    from openpyxl import load_workbook

    path = Path(template_path)
    if not path.exists():
        raise ValueError("模板文件不存在。")
    wb = load_workbook(io.BytesIO(path.read_bytes()))
    ws = wb[wb.sheetnames[0]]
    header_row = 0
    headers: dict[int, str] = {}
    for row_index in range(1, min(ws.max_row, 20) + 1):
        for col_index in range(1, ws.max_column + 1):
            value = ws.cell(row=row_index, column=col_index).value
            if value is None or not str(value).strip():
                continue
            text = str(value).strip()
            if text in mapping:
                headers[col_index] = text
        if headers:
            header_row = row_index
            break
    if not headers:
        wb.close()
        raise ValueError("模板里没有找到可识别的表头（需与导出字段名一致）。")
    for col_index, key in headers.items():
        value = mapping[key]
        if isinstance(value, (list, tuple)):
            for offset, item in enumerate(value, start=1):
                ws.cell(row=header_row + offset, column=col_index, value=item)
        else:
            ws.cell(row=header_row + 1, column=col_index, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def pdf_with_header_footer(pdf_bytes: bytes, *, school: str = "",
                           title: str = "", logo: bytes | None = None,
                           footer: str = "", page_size: str = "A4",
                           orientation: str = "portrait",
                           margins: tuple[int, int, int, int] | None = None,
                           toc: list[dict] | None = None) -> bytes:
    """给 PDF 加页眉页脚/页码/书签，并按需要换页面尺寸。

    margins 顺序 (上下左右)，单位 point。
    toc 形如 [{"title": "第一章", "page": 1}]，页码从 1 开始。
    """
    import pymupdf

    margins = margins or (36, 36, 36, 36)
    size_map = {"A4": (595, 842), "A3": (842, 1191), "B5": (499, 709)}
    base_width, base_height = size_map.get(str(page_size).upper(),
                                           size_map["A4"])
    if str(orientation).lower().startswith("land"):
        base_width, base_height = base_height, base_width

    source = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    target = pymupdf.open()
    for index, page in enumerate(source):
        new_page = target.new_page(width=base_width, height=base_height)
        scale = min((base_width - 2) / page.rect.width,
                    (base_height - 2) / page.rect.height, 1.0)
        offset_x = (base_width - page.rect.width * scale) / 2
        offset_y = (base_height - page.rect.height * scale) / 2
        new_page.show_pdf_page(
            pymupdf.Rect(offset_x, offset_y,
                         offset_x + page.rect.width * scale,
                         offset_y + page.rect.height * scale),
            source, index)
        top, bottom, left, right = margins
        if logo:
            try:
                new_page.insert_image(
                    pymupdf.Rect(left, top / 2, left + 48, top / 2 + 48),
                    stream=logo)
            except Exception:  # noqa: BLE001 —— logo 不合法就跳过
                pass
        header_text = "　".join(x for x in (school, title) if x)
        if header_text:
            new_page.insert_text((left + (56 if logo else 0), top - 12),
                                 header_text, fontsize=10,
                                 fontname="china-s")
        footer_text = footer or ""
        page_label = f"第 {index + 1} / {len(source)} 页"
        if footer_text:
            page_label = f"{footer_text}　{page_label}"
        new_page.insert_text((left, base_height - bottom + 14), page_label,
                             fontsize=9, fontname="china-s")
    if toc:
        items = []
        for item in toc:
            page_no = int(item.get("page") or 1)
            level = int(item.get("level") or 1)
            items.append([level, str(item.get("title") or ""),
                          max(1, min(page_no, target.page_count))])
        target.set_toc(items)
    out = io.BytesIO()
    target.save(out)
    target.close()
    source.close()
    return out.getvalue()


def docx_from_school_template(template_path, mapping: dict) -> bytes:
    """把 {{字段}} 占位符替换进学校 Word 模板，返回新文件字节。"""
    from docx import Document

    path = Path(template_path)
    if not path.exists():
        raise ValueError("模板文件不存在。")
    doc = Document(io.BytesIO(path.read_bytes()))
    hit = 0

    def _replace(text: str) -> str:
        nonlocal hit
        result = str(text or "")
        for key, value in mapping.items():
            for token in (f"{{{{{key}}}}}", f"{{{{{key} }}}}"):
                if token in result:
                    result = result.replace(token, str(value))
                    hit += 1
        return result

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "{{" in run.text:
                run.text = _replace(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "{{" in run.text:
                            run.text = _replace(run.text)
    if not hit:
        raise ValueError("模板里没有找到 {{字段}} 占位符。")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def standard_exam_docx(homework, questions, *, school: str = "",
                       with_answer: bool = False,
                       header: str = "", footer: str = "") -> bytes:
    """按标准试卷格式导出 Word：密封线、题号、分值，可选答案与页眉页脚。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)

    sealed = doc.add_table(rows=1, cols=1)
    cell = sealed.cell(0, 0)
    cell.text = (f"学校：{school or '＿＿＿＿'}　班级：＿＿＿＿　"
                 "姓名：＿＿＿＿　学号：＿＿＿＿")
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(getattr(homework, "name", "") or "试卷")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    minutes = getattr(homework, "duration", None)
    meta.add_run(f"满分 {getattr(homework, 'total_score', None) or '—'} 分"
                 + (f"　考试时间 {minutes} 分钟" if minutes else ""))

    for index, (link, question) in enumerate(questions, start=1):
        paragraph = doc.add_paragraph()
        score = getattr(link, "score", None)
        prefix = f"{index}. "
        if score:
            prefix += f"（{score} 分）"
        paragraph.add_run(prefix + str(getattr(question, "content", "")))
        if with_answer:
            answer = doc.add_paragraph()
            answer.add_run(f"答案：{getattr(question, 'answer', '') or '—'}")
            answer.runs[0].font.size = Pt(9)
    if header:
        for section in doc.sections:
            section.header.paragraphs[0].text = header
    if footer:
        for section in doc.sections:
            section.footer.paragraphs[0].text = footer
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
