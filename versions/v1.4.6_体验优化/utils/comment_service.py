# -*- coding: utf-8 -*-
"""
期末评语服务层（v1.0）。

职责：
- 组装单个学生的评语输入文本（复用画像口径：层次/趋势/强弱科/历次成绩，
  再叠加作业错题暴露的薄弱知识点），作为 user 消息注入主模型；
- 评语模板库 CommentTemplate 的增删查；
- 学生评语 StudentComment 按“学生 + 学期” upsert、读取、批量导出 Word。

逐人调用 LLM、进度条与失败重试在页面层做，本层保持离线可测。
"""

from __future__ import annotations

import json
from datetime import datetime

from models.models import (
    CommentTemplate, HomeworkAnswer, Question, Student, StudentComment)
from utils import exam_service, stats as stats_mod

# 三种风格（页面和测试共用）
STYLES = ["鼓励", "中肯", "严格"]
# 评语字数区间
MIN_LEN, MAX_LEN = 100, 150


# ---------------------------------------------------------------------------
# 单个学生数据组装
# ---------------------------------------------------------------------------

def _profile_summary(session, student_id: int) -> dict:
    """复用学生画像的口径，算出层次/趋势/强弱科/历次成绩。"""
    history = exam_service.student_scores_over_time(session, student_id)
    subject_rates_all: dict[str, list[float]] = {}
    total_rates: list[float] = []
    for h in history:
        exam = exam_service.get_exam(session, h["exam_id"])
        subjects = exam_service.exam_subjects(session, h["exam_id"])
        full = exam_service.full_scores_of(exam, subjects)
        for s in subjects:
            if h.get(s) is not None and full.get(s):
                subject_rates_all.setdefault(s, []).append(h[s] / full[s])
        full_total = sum(full.values())
        if h["total"] is not None and full_total:
            total_rates.append(h["total"] / full_total)

    avg_rate = {s: sum(v) / len(v) for s, v in subject_rates_all.items()}
    strengths = stats_mod.subject_strength(avg_rate)
    overall = sum(avg_rate.values()) / len(avg_rate) if avg_rate else None
    return {
        "history": history,
        "subjects": sorted(avg_rate.keys()),
        "avg_rate": avg_rate,
        "level": stats_mod.level_tag(overall),
        "trend": stats_mod.trend_label(total_rates),
        "strong": [s for s, t in strengths.items() if t == "强项"],
        "weak": [s for s, t in strengths.items() if t == "弱项"],
    }


def _weak_knowledge_from_homework(session, student_id: int, top_n: int = 5) -> list[str]:
    """从该生作业错题里统计出现最多的薄弱知识点。"""
    rows = (session.query(HomeworkAnswer, Question)
            .join(Question, HomeworkAnswer.question_id == Question.id)
            .filter(HomeworkAnswer.student_id == student_id,
                    HomeworkAnswer.is_correct.is_(False)).all())
    counter: dict[str, int] = {}
    for ans, question in rows:
        try:
            kps = json.loads(question.knowledge_points) if question.knowledge_points else []
        except (ValueError, TypeError):
            kps = []
        for kp in (kps if isinstance(kps, list) else []):
            counter[str(kp)] = counter.get(str(kp), 0) + 1
    return [k for k, _ in sorted(counter.items(), key=lambda x: -x[1])[:top_n]]


def build_comment_input(session, student: Student, term: str | None = None) -> str:
    """组装发给 AI 的单个学生数据文本。"""
    prof = _profile_summary(session, student.id)
    lines = [f"学生：{student.name}；班级：{student.class_name or '未分班'}"
             + (f"；学期：{term}" if term else "")]
    if not prof["history"]:
        lines.append("该生暂无考试成绩数据，请只写一段以鼓励为主、期待进步的通用评语。")
        return "\n".join(lines)

    rates = "，".join(f"{s}{prof['avg_rate'][s] * 100:.1f}%"
                      for s in prof["subjects"])
    lines.append(f"系统标签：成绩层次「{prof['level']}」，趋势「{prof['trend']}」。")
    lines.append(f"各科平均得分率：{rates}。")
    lines.append("强项科目：" + ("、".join(prof["strong"]) or "暂无明显强项") + "。")
    lines.append("薄弱科目：" + ("、".join(prof["weak"]) or "暂无明显弱项") + "。")
    lines.append(f"共参考 {len(prof['history'])} 次，历次成绩：")
    for h in prof["history"]:
        parts = [f"{s}{h.get(s):g}" for s in prof["subjects"]
                 if h.get(s) is not None]
        lines.append(f"- {h['exam_name']}：" + "，".join(parts)
                     + f"，总分{h['total']:g}")
    weak_kps = _weak_knowledge_from_homework(session, student.id)
    if weak_kps:
        lines.append("作业错题暴露的薄弱知识点：" + "、".join(weak_kps) + "。")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 评语模板库
# ---------------------------------------------------------------------------

def list_templates(session, style: str | None = None) -> list[CommentTemplate]:
    q = session.query(CommentTemplate)
    if style:
        q = q.filter(CommentTemplate.style == style)
    return q.order_by(CommentTemplate.id.desc()).all()


def create_template(session, name: str, content: str, style: str = "中肯") -> CommentTemplate:
    if not name or not name.strip():
        raise ValueError("模板名称不能为空")
    tpl = CommentTemplate(name=name.strip(), content=(content or "").strip(),
                          style=style if style in STYLES else "中肯")
    session.add(tpl)
    session.flush()
    return tpl


def delete_template(session, template_id: int) -> None:
    tpl = session.get(CommentTemplate, template_id)
    if tpl is None:
        raise ValueError(f"模板不存在：id={template_id}")
    session.delete(tpl)


# ---------------------------------------------------------------------------
# 学生评语留档（学生 + 学期 唯一）
# ---------------------------------------------------------------------------

def normalize_style(style: str | None) -> str:
    return style if style in STYLES else "中肯"


def upsert_comment(session, student_id: int, term: str, content: str,
                   style: str = "中肯") -> StudentComment:
    """按“学生+学期”写入评语，已存在则更新（不新增重复行）。"""
    if not term or not term.strip():
        raise ValueError("学期不能为空")
    content = (content or "").strip()
    if not content:
        raise ValueError("评语内容不能为空")
    row = (session.query(StudentComment)
           .filter(StudentComment.student_id == student_id,
                   StudentComment.term == term.strip()).first())
    if row is None:
        row = StudentComment(student_id=student_id, term=term.strip(),
                             content=content, style=normalize_style(style))
        session.add(row)
    else:
        row.content = content
        row.style = normalize_style(style)
        row.updated_at = datetime.now()
    session.flush()
    return row


def get_comment(session, student_id: int, term: str) -> StudentComment | None:
    return (session.query(StudentComment)
            .filter(StudentComment.student_id == student_id,
                    StudentComment.term == term).first())


def list_comments(session, class_name: str | None = None,
                  term: str | None = None) -> list[dict]:
    """读取评语（联表学生，便于按班导出），按班级、学号排序。"""
    q = (session.query(StudentComment, Student)
         .join(Student, StudentComment.student_id == Student.id))
    if class_name:
        q = q.filter(Student.class_name == class_name)
    if term:
        q = q.filter(StudentComment.term == term)
    result = []
    for comment, student in q.order_by(Student.class_name, Student.student_no,
                                       Student.id).all():
        result.append({
            "comment_id": comment.id, "student_id": student.id,
            "student_name": student.name, "class_name": student.class_name,
            "student_no": student.student_no, "term": comment.term,
            "style": comment.style, "content": comment.content,
            "updated_at": comment.updated_at,
        })
    return result


def export_comments_word(rows: list[dict], title: str = "期末评语") -> bytes:
    """按名单排版批量导出 Word：每个学生一段。"""
    import io
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    current_class = None
    for r in rows:
        if r["class_name"] != current_class:
            current_class = r["class_name"]
            doc.add_heading(current_class or "未分班", level=1)
        doc.add_paragraph(f"{r['student_name']}：{r['content']}")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
