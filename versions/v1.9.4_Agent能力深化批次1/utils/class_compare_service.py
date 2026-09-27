# -*- coding: utf-8 -*-
"""班级对比分析服务（v1.9.0）。

支持 2-5 个班级，按一场考试或一段时间对比：
平均分、最高分、最低分、及格率、优秀率和各分数段人数。
对比报告用 python-docx 直出，不新增依赖。
"""

from __future__ import annotations

from models.models import Score, Student
from utils import exam_service, stats as stats_mod


def _load_thresholds(thresholds):
    if thresholds:
        from utils.analysis_settings import validate_thresholds
        return validate_thresholds(thresholds)
    from utils.analysis_settings import load_thresholds
    return load_thresholds()


def _class_score_values(session, exam_id, subject):
    """取一场考试每个学生的分数（按科目或总分），按班级归类。

    返回 {class_name: [float, ...]}
    """
    if subject:
        rows = (session.query(Score, Student)
                .join(Student, Score.student_id == Student.id)
                .filter(Score.exam_id == exam_id, Score.subject == subject).all())
        buckets = {}
        for sc, stu in rows:
            if sc.score is None:
                continue
            buckets.setdefault(stu.class_name or "未分班", []).append(float(sc.score))
        return buckets
    # 总分：用宽表（各科相加）。
    wide = exam_service.exam_score_rows(session, exam_id)
    buckets = {}
    for r in wide:
        if r.get("total") is None:
            continue
        buckets.setdefault(r.get("class_name") or "未分班", []).append(float(r["total"]))
    return buckets



def compare_exam_classes(session, exam_id, class_names, subject=None,
                         thresholds=None) -> dict:
    """按一场考试对比 2-5 个班级。"""
    class_names = list(class_names or [])
    if not 2 <= len(class_names) <= 5:
        raise ValueError("班级对比需要选择 2-5 个班级。")

    exam = exam_service.get_exam(session, exam_id)
    if exam is None:
        raise ValueError(f"考试不存在：id={exam_id}")

    t = _load_thresholds(thresholds)
    subjects = exam_service.exam_subjects(session, exam_id)
    if subject:
        if subject not in subjects:
            raise ValueError(f"这场考试没有“{subject}”成绩。")
        full_scores = exam_service.full_scores_of(exam, [subject])
        full_score = full_scores[subject]
    else:
        full_score = round(sum(
            exam_service.full_scores_of(exam, subjects).values()), 2)

    all_buckets = _class_score_values(session, exam_id, subject)
    result_classes = []
    bands_out = {}
    for cls in class_names:
        values = all_buckets.get(cls, [])
        pass_line = full_score * t["pass_ratio"]
        excellent_line = full_score * t["excellent_ratio"]
        summary = stats_mod.subject_summary(
            values, full_score, t["pass_ratio"], t["excellent_ratio"],
            custom_bands=True)
        result_classes.append({
            "class_name": cls,
            "count": len(values),
            "mean": round(sum(values) / len(values), 2) if values else None,
            "max": max(values) if values else None,
            "min": min(values) if values else None,
            "pass_rate": summary.get("pass_rate"),
            "excellent_rate": summary.get("excellent_rate"),
        })
        bands_out[cls] = summary.get("bands", [])

    return {
        "exam": exam,
        "subject": subject,
        "full_score": full_score,
        "classes": result_classes,
        "bands": bands_out,
        "thresholds": t,
    }


def compare_class_trends(session, class_names, subject=None,
                         exam_ids=None) -> dict:
    """跨考试对比各班均分走势。"""
    class_names = list(class_names or [])
    if not 2 <= len(class_names) <= 5:
        raise ValueError("班级对比需要选择 2-5 个班级。")

    exams = exam_service.list_exams(session)
    if exam_ids is not None:
        wanted = {int(x) for x in exam_ids}
        exams = [e for e in exams if e.id in wanted]

    points = []
    series = {cls: [] for cls in class_names}
    for exam in exams:
        buckets = _class_score_values(session, exam.id, subject)
        if not any(cls in buckets for cls in class_names):
            continue
        row = {"exam_id": exam.id, "exam_name": exam.name,
               "exam_date": exam.exam_date}
        for cls in class_names:
            values = buckets.get(cls, [])
            mean = round(sum(values) / len(values), 2) if values else None
            row[cls] = mean
            series[cls].append(mean)
        points.append(row)

    return {
        "exams": [p["exam_name"] for p in points],
        "exam_ids": [p["exam_id"] for p in points],
        "exam_dates": [p["exam_date"] for p in points],
        "points": points,
        "class_names": class_names,
        "subject": subject,
        "series": series,
    }


def export_compare_report(session, compare_result) -> bytes:
    """把 compare_exam_classes 的结果导出成 Word 报告。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    exam = compare_result["exam"]
    subject_label = compare_result.get("subject") or "总分"
    doc.add_heading(f"班级对比报告：{exam.name}", level=0)
    doc.add_paragraph(f"对比科目：{subject_label}    满分：{compare_result['full_score']}")

    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, name in enumerate(["班级", "人数", "平均分", "最高", "最低",
                              "及格率", "优秀率"]):
        hdr[i].text = name
    for item in compare_result["classes"]:
        cells = table.add_row().cells
        cells[0].text = item["class_name"]
        cells[1].text = str(item["count"])
        cells[2].text = str(item["mean"])
        cells[3].text = str(item["max"])
        cells[4].text = str(item["min"])
        cells[5].text = (f"{item['pass_rate']}%"
                         if item["pass_rate"] is not None else "-")
        cells[6].text = (f"{item['excellent_rate']}%"
                         if item["excellent_rate"] is not None else "-")

    doc.add_heading("各分数段人数", level=1)
    for cls, bands in compare_result["bands"].items():
        doc.add_paragraph(cls, style="Heading 2")
        for band in bands:
            doc.add_paragraph(
                f"{band.get('label')}：{band.get('count')} 人", style="List Bullet")

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
