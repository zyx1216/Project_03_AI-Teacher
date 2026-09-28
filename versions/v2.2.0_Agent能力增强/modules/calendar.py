# -*- coding: utf-8 -*-
"""教学日历页面：AgGrid 月历、当天安排、本周概览、课程安排表。"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
import streamlit as st
from st_aggrid import AgGrid, GridUpdateMode

from models.models import Homework, HomeworkScore
from utils.db import SessionLocal
from utils import calendar_service as cal_svc
from utils import class_service
from utils import schedule_service
from utils import student_service

WEEKDAY_CN = ["一", "二", "三", "四", "五", "六", "日"]
_current_events = {}
MONTH_LABEL_CAP = "教学日历按日期显示考试、作业和备课；作业和备课日期取创建当天。"


def show():
    """渲染教学日历页（月视图 / 学期视图）。"""
    st.title("📅 教学日历")
    st.caption(MONTH_LABEL_CAP)

    today = date.today()
    semester = _semester_settings(today)

    # 学期进度
    current_week, total_weeks, percent = cal_svc.semester_progress(
        today, semester)
    pc1, pc2 = st.columns([3, 2])
    pc1.caption(
        f"本学期第 {current_week} 周，共 {total_weeks} 周，已过 {percent}%")
    pc2.progress(percent / 100.0)


    # 视图切换
    view = st.radio("日历视图", ["📅 月视图", "📆 学期视图"],
                    horizontal=True, key="calendar_view",
                    label_visibility="collapsed")

    if "cal_year" not in st.session_state:
        st.session_state["cal_year"] = today.year
    if "cal_month" not in st.session_state:
        st.session_state["cal_month"] = today.month

    with SessionLocal() as session:
        _render_teaching_plan_panel(session, semester)
        if view == "📅 月视图":
            clicked = _month_view(session, today, semester)
        else:
            clicked = _semester_view(session, semester)

        st.divider()
        _render_day_detail(_current_events, today, clicked)

        if view == "📅 月视图":
            st.divider()
            _render_week_overview(session, today)

        st.divider()
        _render_schedule_panel(session)


def _semester_settings(today) -> dict:
    """学期设置折叠区，返回当前学期配置。"""
    with st.expander("⚙️ 学期设置", expanded=False):
        current = cal_svc.load_semester()
        start = date.fromisoformat(current["start"])
        end = date.fromisoformat(current["end"])
        c1, c2, c3 = st.columns([2, 2, 1])
        new_start = c1.date_input("开学日期", value=start,
                                  key="semester_start_input")
        new_end = c2.date_input("放假日期", value=end,
                                key="semester_end_input")
        if c3.button("保存", key="save_semester"):
            if new_end <= new_start:
                st.error("放假日期必须晚于开学日期。")
            else:
                saved = cal_svc.save_semester(
                    {"start": new_start.isoformat(),
                     "end": new_end.isoformat()})
                st.toast(
                    f"已保存学期：{saved['start']} 至 {saved['end']}")
                st.rerun()
    return cal_svc.load_semester()


def _month_view(session, today, semester):
    """月视图：翻月 + 标周次的月历。"""
    cc1, cc2, cc3 = st.columns([1, 1, 3])
    if cc1.button("⬅️ 上一月", key="cal_prev_month"):
        _shift_month(-1)
    if cc2.button("下一月 ➡️", key="cal_next_month"):
        _shift_month(1)
    cc3.caption(f"当前查看月份：{st.session_state['cal_year']} 年 "
                f"{st.session_state['cal_month']} 月")

    year = st.session_state["cal_year"]
    month = st.session_state["cal_month"]
    start = date(year, month, 1)
    last_week = cal_svc.month_weeks(year, month)[-1]
    end_day = next(d for d in reversed(last_week) if d is not None)

    global _current_events
    _current_events = cal_svc.collect_events(session, start, end_day)
    week_no_map = cal_svc.semester_week_no_map(semester)
    return _render_month(year, month, _current_events, start, end_day,
                         week_no_map=week_no_map)


def _semester_view(session, semester):
    """学期视图：各周表格。"""
    start = date.fromisoformat(semester["start"])
    end = date.fromisoformat(semester["end"])
    global _current_events
    _current_events = cal_svc.collect_events(session, start, end)

    rows = cal_svc.semester_aggrid_rows(semester, _current_events)
    try:
        response = AgGrid(
            pd.DataFrame(rows),
            gridOptions=cal_svc.build_semester_grid_options(),
            key="calendar_semester_grid",
            update_mode=GridUpdateMode.MODEL_CHANGED,
            data_return_mode="AS_INPUT",
            allow_unsafe_jscode=True,
            height=min(420, 40 + len(rows) * 95),
            show_search=False,
            show_toolbar=False,
            show_download_button=False,
            fit_columns_on_grid_load=True,
        )
    except Exception as exc:
        st.error(f"学期表加载失败：{exc}")
        return None
    records = response.data.to_dict("records")
    return cal_svc.clicked_semester_date(records)


def _shift_month(delta: int) -> None:
    year = st.session_state["cal_year"]
    month = st.session_state["cal_month"] + delta
    if month < 1:
        month, year = 12, year - 1
    elif month > 12:
        month, year = 1, year + 1
    st.session_state["cal_year"] = year
    st.session_state["cal_month"] = month


def _render_month(
    year: int,
    month: int,
    events: dict,
    start: date,
    end_day: date,
    week_no_map=None,
) -> date | None:
    """用 AgGrid 渲染月历，并返回本次点击的日期。"""
    rows = cal_svc.month_aggrid_rows(
        year, month, events, week_no_map=week_no_map)
    options = cal_svc.build_month_grid_options(date.today())
    height = 36 + len(rows) * 105 + 20

    try:
        response = AgGrid(
            pd.DataFrame(rows),
            gridOptions=options,
            key="calendar_month_grid",
            update_mode=GridUpdateMode.MODEL_CHANGED,
            data_return_mode="AS_INPUT",
            allow_unsafe_jscode=True,
            height=height,
            show_search=False,
            show_toolbar=False,
            show_download_button=False,
            fit_columns_on_grid_load=True,
        )
    except Exception as exc:
        st.error(f"月历组件加载失败：{exc}")
        return None

    records = response.data.to_dict("records")
    return cal_svc.clicked_date_from_rows(records, start, end_day)


def _render_day_detail(
    events: dict,
    today: date,
    clicked: date | None,
) -> None:
    st.subheader("查看某天")
    busy_dates = sorted(events.keys())
    if busy_dates:
        st.caption("本月有安排的日期：" +
                   "、".join(d.strftime("%m月%d日") for d in busy_dates))

    if clicked is not None:
        # 日期输入框尚未创建，先写它的固定 key；本轮详情就会定位到点击日期。
        st.session_state["cal_picked_day"] = clicked
    chosen = st.date_input("选择日期", value=today, key="cal_picked_day")

    items = events.get(chosen, [])
    if not items:
        st.info("这一天没有安排。")
    else:
        kind_cn = {"exam": "考试", "homework": "作业", "lesson": "备课"}
        for item in items:
            st.write(
                f"{item['icon']} {kind_cn.get(item['kind'], item['kind'])}："
                f"{item['title']}"
            )

    st.write("**快捷前往**")
    b1, b2, b3 = st.columns(3)
    if b1.button("📝 去学业测评", key="cal_go_homework", use_container_width=True):
        _goto_page("📝 学业测评")
    if b2.button("📊 去新建考试", key="cal_go_exam", use_container_width=True):
        _goto_page("📊 学情")
    if b3.button("📚 去备课", key="cal_go_lesson", use_container_width=True):
        _goto_page("📚 备课")


def _goto_page(page: str) -> None:
    """请求跨主导航跳转，由 app.py 在渲染前应用（不自动开弹窗）。"""
    st.session_state["_pending_main_page"] = page
    st.rerun()


def _render_week_overview(session, today) -> None:
    st.subheader("本周概览")
    monday, sunday = cal_svc.week_range(today)
    events = cal_svc.collect_events(session, monday, sunday)
    if not events:
        st.info("本周没有考试、作业或备课安排。")
    else:
        for day in sorted(events):
            line = "、".join(f"{e['icon']}{e['title']}" for e in events[day])
            st.write(f"**{day.strftime('%m月%d日（周')}"
                     f"{WEEKDAY_CN[day.weekday()]}）**：{line}")

    st.write("**待办提醒**")
    upcoming = cal_svc.upcoming_exams(session, today, 7)
    if upcoming:
        for exam in upcoming:
            st.write(f"📊 {exam['date'].strftime('%m月%d日')} 有考试：{exam['title']}")
    else:
        st.caption("未来 7 天没有考试。")

    missing = _homeworks_missing_scores(session, events)
    if missing:
        for title, n in missing:
            st.write(f"📝 作业「{title}」还有 {n} 名学生未录成绩。")


def _homeworks_missing_scores(session, week_events) -> list[tuple]:
    result = []
    for day, items in week_events.items():
        for item in items:
            if item["kind"] != "homework":
                continue
            hw = session.get(Homework, item["id"])
            if not hw or not hw.class_name:
                continue
            recorded = (session.query(HomeworkScore)
                        .filter(HomeworkScore.homework_id == item["id"]).count())
            total = student_service.list_students(session, class_name=hw.class_name)
            n = len(total) - recorded
            if n > 0:
                result.append((item["title"], n))
    return result


def _reset_schedule_editor() -> None:
    st.session_state.pop("class_schedule_editor", None)


def _render_schedule_panel(session) -> None:
    st.subheader("课程安排表")
    classes = class_service.list_class_names(session)
    if not classes:
        st.info("还没有班级。先去“学情 → 学生管理 → 班级管理”创建班级。")
        return

    selected_class = st.selectbox(
        "选择班级", options=classes, key="schedule_class_select",
        on_change=_reset_schedule_editor)

    grid = pd.DataFrame(schedule_service.week_grid(selected_class))
    edited = st.data_editor(
        grid,
        key="class_schedule_editor",
        hide_index=True,
        num_rows="fixed",
        width="stretch",
    )

    if st.button("💾 保存课程表", key="save_class_schedule"):
        try:
            result = schedule_service.sync_week_grid(
                selected_class, edited.to_dict("records"))
        except (ValueError, OSError) as exc:
            st.error(str(exc))
        else:
            st.toast(f"已保存 {result['saved']} 节课，清空 {result['deleted']} 个旧课程格。")

    st.caption("单元格只填写 9 个学科名称；清空单元格并保存，会删除这节课。")
    _render_schedule_excel(session, classes)


def _render_schedule_excel(session, classes: list[str]) -> None:
    st.markdown("#### Excel 导入课程表")
    uploaded = st.file_uploader(
        "上传课程安排表 Excel", type=["xlsx", "xls"],
        key="schedule_excel_file")
    c1, c2 = st.columns(2)
    if c1.button("📥 导入课程表", key="import_schedule_button"):
        if uploaded is None:
            st.warning("请先选择课程安排表 Excel。")
        else:
            try:
                df = pd.read_excel(uploaded)
                result = schedule_service.import_schedule_entries(df, classes)
            except (ValueError, OSError) as exc:
                st.error(str(exc))
            else:
                st.toast(f"已导入 {result['imported']} 节课。")
                st.rerun()

    template = schedule_service.schedule_template_dataframe()
    output = BytesIO()
    template.to_excel(output, index=False)
    c2.download_button(
        "⬇️ 下载 Excel 模板",
        data=output.getvalue(),
        file_name="课程安排表模板.xlsx",
        key="download_schedule_template",
    )


# ---------------------------------------------------------------------------
# v2.2.0：学期/周教学计划与进度跟踪
# ---------------------------------------------------------------------------

def _parse_chapter_lines(text: str) -> list[str]:
    return [line.strip(" -、\t") for line in str(text or "").splitlines()
            if line.strip()]


def _semester_plan_excel(plan: dict) -> bytes:
    output = BytesIO()
    rows = plan.get("weekly_plan", [])
    pd.DataFrame(rows).to_excel(output, index=False, sheet_name="学期计划")
    return output.getvalue()


def _semester_plan_word(plan: dict) -> bytes:
    from docx import Document
    doc = Document()
    meta = plan.get("meta", {})
    doc.add_heading(
        f"{meta.get('subject', '')}{meta.get('grade', '')}学期教学计划", 0)
    for item in plan.get("weekly_plan", []):
        doc.add_heading(f"第{item.get('week_number')}周", level=2)
        doc.add_paragraph(
            f"内容：{item.get('content', '')}；课时：{item.get('hours', '')}")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _render_teaching_plan_panel(session, semester) -> None:
    """教学计划：生成、查看、录入进度、对比和一键调整。"""
    st.divider()
    st.subheader("📆 教学计划")
    c1, c2, c3 = st.columns(3)
    subject = c1.selectbox("计划学科", ["数学", "语文", "英语", "物理", "化学",
                                       "生物", "政治", "历史", "地理"],
                           key="teaching_plan_subject")
    grade = c2.text_input("计划年级", key="teaching_plan_grade",
                          placeholder="如：高一、初二、三年级")
    semester_name = c3.text_input(
        "学期", key="teaching_plan_semester",
        value=f"{date.today().year}年秋季学期")
    c4, c5 = st.columns(2)
    weekly_hours = c4.number_input("每周课时", min_value=1, max_value=12,
                                   value=4, key="teaching_plan_weekly_hours")
    chapters_text = c5.text_area(
        "教材章节（每行一章，可空）", key="teaching_plan_chapters",
        placeholder="第一章 有理数`n第二章 整式的加减", height=90)

    if st.button("🧩 生成学期计划", key="teaching_plan_generate"):
        plan = __import__("utils.teaching_plan_service", fromlist=["x"]).generate_semester_plan(
            subject, grade, semester_name, weekly_hours,
            _parse_chapter_lines(chapters_text))
        st.session_state["current_semester_plan"] = plan
        st.toast("学期计划已生成，可保存到进度记录。")

    plan = st.session_state.get("current_semester_plan")
    if not plan:
        st.info("填写学科、年级、章节后生成学期计划。")
        return

    st.dataframe(pd.DataFrame(plan["weekly_plan"]), hide_index=True,
                 width="stretch")
    b1, b2, b3, b4 = st.columns(4)
    if b1.button("💾 保存计划", key="teaching_plan_save"):
        count = __import__("utils.teaching_plan_service", fromlist=["x"]).save_planned_weeks(
            session, plan)
        st.toast(f"已保存 {count} 周计划。")
    b2.download_button("⬇️ 导出 Excel", _semester_plan_excel(plan),
                       file_name="学期教学计划.xlsx",
                       key="teaching_plan_excel")
    b3.download_button("📄 导出 Word", _semester_plan_word(plan),
                       file_name="学期教学计划.docx",
                       key="teaching_plan_word")
    if b4.button("🔀 查看进度对比", key="teaching_plan_compare"):
        st.session_state["teaching_compare_open"] = True

    st.markdown("#### 周计划与实际进度")
    w1, w2, w3 = st.columns([1, 2, 2])
    week_no = w1.number_input("周次", min_value=1, max_value=20, value=1,
                              key="teaching_progress_week")
    try:
        weekly = __import__("utils.teaching_plan_service", fromlist=["x"]).generate_weekly_plan(
            plan, int(week_no))
        w2.caption(f"计划：{weekly.get('content', '')}")
    except ValueError as exc:
        w2.caption(str(exc))
    actual = w3.text_input("实际完成内容", key="teaching_actual_content")
    note = st.text_input("备注", key="teaching_progress_note")
    if st.button("✅ 录入实际进度", key="teaching_progress_track"):
        __import__("utils.teaching_plan_service", fromlist=["x"]).track_progress(
            session, subject, grade, semester_name, int(week_no),
            actual, note)
        st.toast("已录入本周实际进度。")

    if st.session_state.get("teaching_compare_open"):
        report = __import__("utils.teaching_plan_service", fromlist=["x"]).compare_progress(
            session, subject, grade, semester_name)
        st.dataframe(pd.DataFrame(report["details"]), hide_index=True,
                     width="stretch")
        if report["lag_count"]:
            st.warning(f"落后 {report['lag_count']} 周：{report['lag_weeks']}")
        if report["ahead_count"]:
            st.info(f"超前 {report['ahead_count']} 周：{report['ahead_weeks']}")
        if st.button("🛠️ 一键调整后续计划", key="teaching_plan_adjust"):
            adjusted = __import__("utils.teaching_plan_service", fromlist=["x"]).auto_adjust_plan(
                plan, report["current_week"], report["details"])
            st.session_state["current_semester_plan"] = adjusted
            __import__("utils.teaching_plan_service", fromlist=["x"]).apply_adjusted_plan(
                session, adjusted)
            st.toast("已调整并保存后续计划。")
            st.rerun()
