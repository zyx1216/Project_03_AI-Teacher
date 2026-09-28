# -*- coding: utf-8 -*-
"""
教案服务层。

负责：
- 教案 JSON 的解析与容错修复（LLM 常带代码围栏或多余文字）；
- 教案的保存、读取、列表、删除；
- 教案导出 Word。

教案内容固定结构（存 LessonPlan.content，JSON 字符串）：
{
  "objectives": {"knowledge": 知识技能, "process": 过程方法, "emotion": 情感态度},
  "key_points": 教学重点,
  "difficult_points": 教学难点,
  "process": [{"stage": 环节名, "minutes": 分钟, "content": 具体内容}, ...],
  "subject": 学科,
  "board_design": 板书设计,
  "reflection": 教学反思预设
}
"""

from __future__ import annotations

import difflib
import json
import re
from typing import Optional

from models.models import LessonPlan, LessonPlanVersion
from utils import llm_client
from utils.app_config import DEFAULT_SUBJECT, is_valid_subject

# 教学过程的 5 个固定环节
PROCESS_STAGES = ["导入", "新授", "巩固练习", "课堂小结", "作业布置"]


# AI 试讲脚本五段
TRIAL_SECTIONS = [
    ("import", "课堂导入"),
    ("lecture", "新知讲解"),
    ("practice", "课堂练习"),
    ("summary", "课堂总结"),
    ("board_design", "板书设计"),
]

# 空教案模板（新建/解析失败兜底用）
def empty_plan() -> dict:
    return {
        "objectives": {"knowledge": "", "process": "", "emotion": ""},
        "key_points": "",
        "difficult_points": "",
        "process": [{"stage": s, "minutes": _default_minutes(s), "content": ""}
                    for s in PROCESS_STAGES],
        "subject": DEFAULT_SUBJECT,
        "board_design": "",
        "reflection": "",
    }


def _default_minutes(stage: str) -> int:
    """各环节默认建议时长（45 分钟一节课的常见分配）。"""
    return {"导入": 5, "新授": 20, "巩固练习": 12, "课堂小结": 5,
            "作业布置": 3}.get(stage, 5)


def strip_code_fence(text: str) -> str:
    """去掉 ```json ... ``` 代码围栏。"""
    text = text.strip()
    fence = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fence:
        return fence.group(1).strip()
    return text


def extract_json(text: str):
    """
    从 LLM 输出中尽力提取一个 JSON 对象并解析。
    先去围栏，再找最外层花括号；解析失败返回 None。
    """
    text = strip_code_fence(text)
    try:
        return json.loads(text)
    except (json.JSONDecodeError, TypeError):
        pass
    # 退而求其次：截取第一个 { 到最后一个 }
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        candidate = text[start:end + 1]
        try:
            return json.loads(candidate)
        except json.JSONDecodeError:
            return None
    return None


def normalize_plan(raw: dict, subject: str | None = None) -> dict:
    """
    把 LLM 返回的各种字段命名归一化成固定结构，缺的补空。
    兼容常见别名（如 教学目标/知识与技能/重点/难点）。
    """
    if not isinstance(raw, dict):
        return empty_plan()

    def pick(*keys, default=""):
        for k in keys:
            if k in raw and raw[k] not in (None, ""):
                return raw[k]
        return default

    if is_valid_subject(subject):
        plan_subject = subject
    else:
        raw_subject = raw.get("subject")
        plan_subject = raw_subject if is_valid_subject(raw_subject) else DEFAULT_SUBJECT

    # 目标：可能是 dict，也可能是三个并列字段或一整段
    objectives_raw = pick("objectives", "教学目标", default={})
    objectives = {"knowledge": "", "process": "", "emotion": ""}
    if isinstance(objectives_raw, dict):
        objectives["knowledge"] = str(objectives_raw.get(
            "knowledge") or objectives_raw.get("知识与技能") or objectives_raw.get("知识技能") or "")
        objectives["process"] = str(objectives_raw.get(
            "process") or objectives_raw.get("过程与方法") or "")
        objectives["emotion"] = str(objectives_raw.get(
            "emotion") or objectives_raw.get("情感态度与价值观") or objectives_raw.get("情感态度") or "")
    else:
        objectives["knowledge"] = str(objectives_raw)

    # 教学过程：归一化为 [{stage,minutes,content}]
    process = _normalize_process(pick("process", "教学过程", default=[]))

    return {
        "subject": plan_subject,
        "objectives": objectives,
        "key_points": str(pick("key_points", "keyPoints", "教学重点", "重点")),
        "difficult_points": str(pick("difficult_points", "difficultPoints",
                                     "教学难点", "难点")),
        "process": process,
        "board_design": str(pick("board_design", "boardDesign", "板书设计", "板书")),
        "reflection": str(pick("reflection", "教学反思", "反思预设", "教学反思预设")),
    }


def _normalize_process(raw) -> list[dict]:
    """把教学过程归一化；无法识别时返回 5 个空环节。"""
    result = []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                stage = str(item.get("stage") or item.get("环节") or
                            item.get("name") or "").strip()
                content = str(item.get("content") or item.get("内容") or
                              item.get("detail") or "").strip()
                try:
                    minutes = int(item.get("minutes") or item.get("时长") or
                                  item.get("time") or 0)
                except (TypeError, ValueError):
                    minutes = 0
                if stage:
                    result.append({"stage": stage, "minutes": minutes,
                                   "content": content})
            elif isinstance(item, str) and item.strip():
                result.append({"stage": item.strip()[:10], "minutes": 0, "content": ""})
    elif isinstance(raw, dict):
        # 形如 {"导入": "...", "新授": "..."}
        for stage in PROCESS_STAGES:
            if stage in raw:
                result.append({"stage": stage, "minutes": _default_minutes(stage),
                               "content": str(raw[stage])})
    if not result:
        result = [{"stage": s, "minutes": _default_minutes(s), "content": ""}
                  for s in PROCESS_STAGES]
    return result


def parse_lesson_plan(text: str) -> Optional[dict]:
    """LLM 输出 -> 规范化教案；完全无法解析返回 None。"""
    raw = extract_json(text)
    if raw is None:
        return None
    return normalize_plan(raw)


# ---------------------------------------------------------------------------
# 数据库操作
# ---------------------------------------------------------------------------

def save_plan(session, title: str, plan: dict, grade: str = None,
              chapter: str = None, source: str = None, plan_id: int = None,
              subject: str | None = None) -> LessonPlan:
    """新建或更新教案。plan_id 有值且存在时更新，否则新建。

    学科写在教案 JSON 顶层；旧教案缺学科时按数学兼容。
    """
    if plan_id is not None:
        old = session.get(LessonPlan, plan_id)
        if old is not None:
            old_subject = plan_subject(old)
            plan = normalize_plan(plan, subject or old_subject)
            content = json.dumps(plan, ensure_ascii=False)
            old.title = title
            old.grade = grade
            old.chapter = chapter
            old.content = content
            old.textbook_source = source
            session.flush()
            _create_plan_version(session, old)
            return old
    plan = normalize_plan(plan, subject)
    content = json.dumps(plan, ensure_ascii=False)
    lesson = LessonPlan(title=title, grade=grade, chapter=chapter,
                        content=content, textbook_source=source)
    session.add(lesson)
    session.flush()
    _create_plan_version(session, lesson)
    return lesson


def plan_subject(lesson: LessonPlan) -> str:
    """读取教案自身学科；旧教案、损坏 JSON 都按数学兼容。"""
    try:
        data = json.loads(lesson.content) if lesson.content else {}
        value = data.get("subject") if isinstance(data, dict) else None
        return value if is_valid_subject(value) else DEFAULT_SUBJECT
    except (json.JSONDecodeError, TypeError):
        return DEFAULT_SUBJECT


def load_plan(lesson: LessonPlan) -> dict:
    """读取某条教案的内容 JSON；损坏时退回空模板。"""
    try:
        data = json.loads(lesson.content) if lesson.content else {}
        return normalize_plan(data)
    except (json.JSONDecodeError, TypeError):
        return empty_plan()


def find_plan_by_title(session, title: str, subject: str | None = None):
    """按标题查找一份教案（用于草稿复用）；找不到返回 None。"""
    query = session.query(LessonPlan).filter(LessonPlan.title == title)
    if subject is not None:
        items = query.all()
        return next((item for item in items if plan_subject(item) == subject), None)
    return query.first()


def list_plans(session, subject: str | None = None):
    """所有教案，按更新时间倒序；传 subject 时只列该学科教案。"""
    rows = (session.query(LessonPlan)
            .order_by(LessonPlan.updated_at.desc(), LessonPlan.id.desc()).all())
    if subject:
        rows = [lesson for lesson in rows if plan_subject(lesson) == subject]
    return rows


def delete_plan(session, plan_id: int) -> None:
    lesson = session.get(LessonPlan, plan_id)
    if lesson is not None:
        session.delete(lesson)


# ---------------------------------------------------------------------------
# Word 导出
# ---------------------------------------------------------------------------

def _set_doc_style_font(doc) -> None:
    """统一 Word 默认字体和标题字体，中文走 eastAsia。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for style_name, size in (("Normal", 11), ("Title", 20),
                             ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style.font.size = Pt(size)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "微软雅黑")
        rfonts.set(qn("w:ascii"), "微软雅黑")
        rfonts.set(qn("w:hAnsi"), "微软雅黑")


def _text_or_dash(value) -> str:
    text = str(value or "").strip()
    return text or "—"


def _homework_assignment(plan: dict) -> str:
    """从教学过程的“作业布置”环节提取内容。"""
    for step in plan.get("process", []):
        if str(step.get("stage") or "").strip() == "作业布置":
            return _text_or_dash(step.get("content"))
    return "—"


def export_word(lesson: LessonPlan, plan: dict) -> bytes:
    """把教案渲染成 Word，返回字节内容供页面下载。"""
    import io
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    _set_doc_style_font(doc)

    title = doc.add_heading(lesson.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subject = plan.get("subject") or DEFAULT_SUBJECT
    meta = "　".join(x for x in [subject, lesson.grade, lesson.chapter,
                                 lesson.textbook_source] if x)
    meta_paragraph = doc.add_paragraph(meta)
    meta_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in meta_paragraph.runs:
        run.font.size = Pt(10)

    obj = plan["objectives"]
    doc.add_heading("一、教学目标", level=1)
    doc.add_paragraph(f"知识与技能：{_text_or_dash(obj['knowledge'])}")
    doc.add_paragraph(f"过程与方法：{_text_or_dash(obj['process'])}")
    doc.add_paragraph(f"情感态度与价值观：{_text_or_dash(obj['emotion'])}")

    doc.add_heading("二、教学重点", level=1)
    doc.add_paragraph(_text_or_dash(plan.get("key_points")))
    doc.add_heading("三、教学难点", level=1)
    doc.add_paragraph(_text_or_dash(plan.get("difficult_points")))

    doc.add_heading("四、教学过程", level=1)
    for i, step in enumerate(plan["process"], start=1):
        stage = step.get("stage", f"环节{i}")
        minutes = step.get("minutes")
        heading = f"{stage}（{minutes} 分钟）" if minutes else str(stage)
        doc.add_heading(heading, level=2)
        doc.add_paragraph(_text_or_dash(step.get("content")))

    doc.add_heading("五、板书设计", level=1)
    doc.add_paragraph(_text_or_dash(plan.get("board_design")))

    doc.add_heading("六、作业布置", level=1)
    doc.add_paragraph(_homework_assignment(plan))

    doc.add_heading("七、教学反思预设", level=1)
    doc.add_paragraph(_text_or_dash(plan.get("reflection")))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


VERSION_LIMIT = 5


def _create_plan_version(session, lesson: LessonPlan) -> LessonPlanVersion:
    """给当前教案内容追加一版，并只保留最近 5 版。"""
    version = LessonPlanVersion(
        plan_id=lesson.id, title=lesson.title, content_json=lesson.content or "")
    session.add(version)
    session.flush()
    old_rows = (session.query(LessonPlanVersion)
                .filter(LessonPlanVersion.plan_id == lesson.id)
                .order_by(LessonPlanVersion.created_at.desc(),
                          LessonPlanVersion.id.desc()).all())
    for old in old_rows[VERSION_LIMIT:]:
        session.delete(old)
    return version


def list_plan_versions(session, plan_id: int) -> list[LessonPlanVersion]:
    """按时间倒序列出某教案的历史版本。"""
    return (session.query(LessonPlanVersion)
            .filter(LessonPlanVersion.plan_id == plan_id)
            .order_by(LessonPlanVersion.created_at.desc(),
                      LessonPlanVersion.id.desc()).all())


def rollback_plan_version(session, plan_id: int,
                          version_id: int) -> LessonPlan:
    """用指定历史版本内容再保存一次，形成带新版本的回退。"""
    lesson = session.get(LessonPlan, plan_id)
    version = session.get(LessonPlanVersion, version_id)
    if lesson is None or version is None or version.plan_id != plan_id:
        raise ValueError("历史版本不存在。")
    content = json.loads(version.content_json)
    subject = plan_subject(lesson)
    return save_plan(
        session, lesson.title, content, grade=lesson.grade,
        chapter=lesson.chapter, source=lesson.textbook_source,
        plan_id=lesson.id, subject=subject)


# ---------------------------------------------------------------------------
# v2.0.0：教案版本 diff
# ---------------------------------------------------------------------------

def _flat_plan_fields(content: dict) -> dict[str, str]:
    """把教案 JSON 摊平成“字段名 -> 文本”，供逐字段对比。"""
    fields = {
        "教学目标·知识": content.get("objectives", {}).get("knowledge", ""),
        "教学目标·方法": content.get("objectives", {}).get("process", ""),
        "教学目标·情感": content.get("objectives", {}).get("emotion", ""),
        "教学重点": content.get("key_points", ""),
        "教学难点": content.get("difficult_points", ""),
        "板书设计": content.get("board_design", ""),
        "教学反思": content.get("reflection", ""),
    }
    for index, step in enumerate(content.get("process", []), start=1):
        stage = step.get("stage") or f"环节{index}"
        fields[f"教学过程·{index}-{stage}"] = (
            f"时长：{step.get('minutes', '')}\n{step.get('content', '')}")
    return fields


def compare_plan_versions(session, version_id1: int,
                          version_id2: int) -> str:
    """对比两个教案版本，返回中文差异文本。

    逐字段列出增删改；两版完全一致时返回“无差异”。
    """
    v1 = session.get(LessonPlanVersion, version_id1)
    v2 = session.get(LessonPlanVersion, version_id2)
    if v1 is None or v2 is None:
        raise ValueError("历史版本不存在。")
    try:
        c1 = json.loads(v1.content_json)
        c2 = json.loads(v2.content_json)
    except (TypeError, ValueError):
        raise ValueError("版本内容不是合法 JSON。")

    f1 = _flat_plan_fields(c1)
    f2 = _flat_plan_fields(c2)
    all_keys = list(dict.fromkeys(list(f1.keys()) + list(f2.keys())))

    lines = []
    for key in all_keys:
        a = str(f1.get(key, "")).strip()
        b = str(f2.get(key, "")).strip()
        if a == b:
            continue
        if not a:
            lines.append(f"【新增】{key}")
        elif not b:
            lines.append(f"【删除】{key}")
        else:
            lines.append(f"【修改】{key}")
            # 用 difflib 给出逐行差异，便于看具体改动
            diff = difflib.unified_diff(
                a.splitlines(), b.splitlines(), lineterm="")
            detail = [d for d in list(diff)
                      if d.startswith(("+", "-"))
                      and not d.startswith(("+++", "---"))]
            lines.extend(f"    {d}" for d in detail[:8])
    if not lines:
        return "两个版本内容无差异。"
    header = (f"版本对比：#{version_id1} → #{version_id2}\n"
              f"（{v1.title}）")
    return header + "\n" + "\n".join(lines)


# ---------------------------------------------------------------------------
# v1.9.1：AI 试讲
# ---------------------------------------------------------------------------


def _lesson_objective_text(plan: dict) -> str:
    """汇总教学目标。"""
    obj = plan.get("objectives", {})
    return "\n".join(str(x) for x in (
        obj.get("knowledge", ""), obj.get("process", ""),
        obj.get("emotion", "")) if str(x).strip())


def _lesson_process_text(plan: dict) -> str:
    """汇总教案教学过程。"""
    lines = []
    for index, step in enumerate(plan.get("process", []), start=1):
        stage = step.get("stage", f"环节{index}")
        minutes = step.get("minutes", "")
        lines.append(
            f"{stage}（{minutes} 分钟）\n{step.get('content', '')}")
    return "\n\n".join(lines)


def trial_lecture(lesson: LessonPlan, plan: dict, settings: dict) -> dict:
    """根据教案和课堂设置生成五段式试讲脚本。"""
    system_prompt = (
        "你是一名经验丰富的教师培训师，擅长模拟课堂教学。"
        "请根据教案生成可直接照着演练的试讲脚本。"
        "只输出 JSON 对象，包含 import、lecture、practice、summary、"
        "board_design 五个字段。脚本中的关键提问以【关键提问】开头，"
        "学生可能回答以【学生回答】开头，应对策略以【应对策略】开头。"
    )
    user_text = (
        f"教案标题：{lesson.title}\n教学目标：{_lesson_objective_text(plan)}\n"
        f"教学重点：{plan.get('key_points', '')}\n"
        f"教学难点：{plan.get('difficult_points', '')}\n"
        f"教学过程：{_lesson_process_text(plan)}\n\n"
        f"班级层次：{settings.get('class_level', '平行班')}\n"
        f"试讲时长：{settings.get('duration', 40)} 分钟\n"
        f"学生活跃度：{settings.get('activity', '中')}"
    )
    raw = llm_client.chat_content(system_prompt, user_text, temperature=0.8)
    data = extract_json(raw)
    if data is None:
        raise ValueError("AI 返回的试讲脚本不是有效 JSON。")
    return normalize_trial_script(data)


def normalize_trial_script(raw: dict) -> dict:
    """把试讲脚本归一为固定五段；缺失字段补空字符串。"""
    if not isinstance(raw, dict):
        raise ValueError("试讲脚本必须是 JSON 对象。")
    return {
        key: str(raw.get(key, "") or "").strip()
        for key, _label in TRIAL_SECTIONS
    }


def export_trial_word(lesson: LessonPlan, plan: dict,
                      script: dict, settings: dict) -> bytes:
    """导出 AI 试讲脚本 Word。"""
    import io
    from docx import Document

    doc = Document()
    _set_doc_style_font(doc)
    title = doc.add_heading(f"AI 试讲脚本：{lesson.title}", level=0)
    title.alignment = 1
    doc.add_paragraph(
        f"班级层次：{settings.get('class_level', '平行班')}　"
        f"时长：{settings.get('duration', 40)} 分钟　"
        f"学生活跃度：{settings.get('activity', '中')}")
    for index, (key, label) in enumerate(TRIAL_SECTIONS, start=1):
        doc.add_heading(f"{index}. {label}", level=1)
        for line in str(script.get(key, "")).splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# ---------------------------------------------------------------------------
# v1.9.4：自然语言修正教案
# ---------------------------------------------------------------------------


_MODIFY_PLAN_SYSTEM = """你是教案修改器。只按老师的修改意见调整指定部分，未提到的内容必须原样保留。
输入包含原教案 JSON 和修改意见。输出完整教案 JSON，继续使用原字段结构。
不要输出解释，不要使用代码围栏。"""


def modify_lesson_plan(plan: dict, instruction: str) -> dict:
    """根据自然语言修改教案，返回完整教案字典。"""
    old_plan = normalize_plan(plan)
    raw = llm_client.chat_content(
        _MODIFY_PLAN_SYSTEM,
        f"修改意见：{instruction}\n原教案 JSON：{json.dumps(old_plan, ensure_ascii=False)}",
        temperature=0.4,
    )
    modified = extract_json(raw)
    if not isinstance(modified, dict):
        raise ValueError("教案修正失败，请重新描述修改意见。")
    return normalize_plan(modified, old_plan.get("subject"))
