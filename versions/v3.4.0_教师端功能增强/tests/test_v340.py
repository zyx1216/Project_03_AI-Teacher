# -*- coding: utf-8 -*-
"""v3.4.0 教师端功能增强（第三批）测试。"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from streamlit.testing.v1 import AppTest

from models.models import (
    Base, Exam, Homework, HomeworkQuestion, Question, RecycleBin, Score,
    Student,
)
from tests.test_app_smoke import _isolated_app_code, goto_sub
from utils import export_service as es
from utils import parent_communication_service as pcs
from utils import recycle_service as rs
from utils import version_service as vs


@pytest.fixture()
def session(tmp_path):
    eng = create_engine(
        f"sqlite:///{(tmp_path / 'v340.db').as_posix()}",
        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    s = sessionmaker(bind=eng)()
    yield s
    s.close()
    eng.dispose()


def _seed_student(session, name="张三", class_name="一班", score=85.0):
    """建一名学生 + 一场数学考试 + 成绩，返回 (student, exam)。"""
    student = Student(name=name, class_name=class_name)
    exam = Exam(name="期中考试", exam_date=date.today(),
                full_scores='{"数学": 100}')
    session.add_all([student, exam])
    session.flush()
    session.add(Score(exam_id=exam.id, student_id=student.id, subject="数学",
                      score=score, class_rank=5, grade_rank=20))
    session.commit()
    return student, exam


def _seed_exam_homework(session, name="单元测试卷"):
    """建一道题 + 一份试卷，返回 (homework, question, link)。"""
    question = Question(content="1+1=?", answer="2", subject="数学",
                        question_type="fill")
    session.add(question)
    session.flush()
    hw = Homework(name=name, homework_type="exam", subject="数学",
                  total_score=100)
    session.add(hw)
    session.flush()
    link = HomeworkQuestion(homework_id=hw.id, question_id=question.id,
                            order=1, score=10)
    session.add(link)
    session.commit()
    return hw, question, link


# ---------------------------------------------------------------------------
# 数据层
# ---------------------------------------------------------------------------

def test_v340_tables_and_migration():
    from migrations.v3_4_0_migration import migrate
    assert len(Base.metadata.tables) == 38
    assert {"parent_communication", "version_history", "recycle_bin",
            "export_templates"} <= set(Base.metadata.tables)
    eng = create_engine("sqlite:///:memory:")
    out = migrate(eng)
    assert set(out["tables"]) == {"parent_communication", "version_history",
                                  "recycle_bin", "export_templates"}
    assert migrate(eng)["tables"] == out["tables"]  # 幂等
    eng.dispose()


# ---------------------------------------------------------------------------
# 功能8：版本历史
# ---------------------------------------------------------------------------

def test_version_snapshot_compare_rollback(session, tmp_path, monkeypatch):
    monkeypatch.setattr(vs, "VERSION_SETTINGS_PATH",
                        tmp_path / "version_settings.json")
    hw, question, _link = _seed_exam_homework(session)
    v1 = vs.snapshot_homework(session, hw.id, "初始版本")
    session.commit()
    hw.name = "单元测试卷（改）"
    second = Question(content="2+2=?", answer="4", subject="数学")
    session.add(second)
    session.flush()
    session.add(HomeworkQuestion(homework_id=hw.id,
                                 question_id=second.id, order=2, score=10))
    session.commit()
    v2 = vs.snapshot_homework(session, hw.id, "加了一道题")
    session.commit()

    versions = vs.list_versions(session, "homework", hw.id)
    assert [row.version for row in versions] == [2, 1]
    diff = vs.compare_versions(session, v1, v2)
    assert "名称" in diff and "第2题" in diff

    # 回滚：内容回到 v1，并自动把回滚前的内容再存一版。
    out = vs.rollback_version(session, v1)
    session.commit()
    assert out["restored_version"] == 1
    assert session.get(Homework, hw.id).name == "单元测试卷"
    links = (session.query(HomeworkQuestion)
             .filter(HomeworkQuestion.homework_id == hw.id).all())
    assert len(links) == 1
    assert len(vs.list_versions(session, "homework", hw.id)) >= 3


def test_version_limit_configurable(session, tmp_path, monkeypatch):
    monkeypatch.setattr(vs, "VERSION_SETTINGS_PATH",
                        tmp_path / "version_settings.json")
    assert vs.get_version_limit() == 10
    assert vs.set_version_limit(2) == 2
    assert vs.get_version_limit() == 2
    hw, _q, _l = _seed_exam_homework(session)
    for index in range(4):
        vs.snapshot(session, "homework", hw.id, {"name": f"v{index}"},
                    f"第{index}次")
    session.commit()
    rows = vs.list_versions(session, "homework", hw.id)
    assert len(rows) == 2
    assert rows[0].version == 4


# ---------------------------------------------------------------------------
# 功能8：回收站
# ---------------------------------------------------------------------------

def test_recycle_question_roundtrip(session):
    from utils import question_service as qs
    question = Question(content="要被删的题", answer="a", subject="数学")
    session.add(question)
    session.commit()
    qid = question.id
    qs.delete_question(session, qid)
    session.commit()
    assert session.get(Question, qid) is None
    rows = rs.list_items(session)
    assert len(rows) == 1 and rows[0]["item_type"] == "question"
    assert rs.restore(session, rows[0]["id"])["item_id"] == qid
    session.commit()
    assert session.get(Question, qid) is not None
    assert rs.list_items(session) == []


def test_recycle_homework_and_purge(session):
    from utils import homework_service as hws
    hw, _q, _l = _seed_exam_homework(session)
    hws.delete_homework(session, hw.id)
    session.commit()
    assert session.get(Homework, hw.id) is None
    rows = rs.list_items(session, "homework")
    assert len(rows) == 1
    assert rs.purge(session, rows[0]["id"]) is True
    session.commit()
    assert rs.list_items(session, "homework") == []


def test_recycle_empty_and_expired(session):
    hw, _q, _l = _seed_exam_homework(session)
    rs.move_to_bin(session, "homework", hw.id)
    session.commit()
    assert rs.empty(session, "homework") == 1
    session.commit()

    old = RecycleBin(item_type="lesson", item_id=1, title="过期条目",
                     content_json="{}",
                     deleted_at=datetime.now() - timedelta(days=40),
                     expire_at=datetime.now() - timedelta(days=10))
    session.add(old)
    session.commit()
    assert rs.purge_expired(30, session=session) == 1
    assert rs.list_items(session) == []


# ---------------------------------------------------------------------------
# 功能7：家校沟通
# ---------------------------------------------------------------------------

def test_report_card_and_advice(session):
    student, exam = _seed_student(session)
    card = pcs.build_report_card(session, student.id, exam_id=exam.id)
    assert card["name"] == "张三"
    assert card["subjects"]["数学"]["score"] == 85
    assert card["subjects"]["数学"]["class_rank"] == 5
    assert card["subjects"]["数学"]["band"] == "良好"
    text = pcs.report_card_text(card)
    assert "张三" in text and "数学" in text

    advice = pcs.generate_advice(session, student.id, subject="数学",
                                 chat_func=None)
    assert advice["family"] and advice["practice"] and advice["cautions"]
    body = pcs.advice_text(advice)
    assert "在家辅导" in body and "推荐练习" in body


def test_class_report(session):
    first, exam = _seed_student(session, name="张三", score=85)
    second, _ = _seed_student(session, name="李四", score=55)
    session.add(Score(exam_id=exam.id, student_id=second.id, subject="数学",
                      score=55))
    session.commit()
    report = pcs.build_class_report(session, "一班", exam_id=exam.id)
    assert report["student_count"] == 2
    stats = report["subject_stats"][0]
    assert stats["subject"] == "数学"
    assert stats["count"] == 2
    assert 0 < stats["pass_rate"] <= 100
    assert "一班" in pcs.class_report_text(report)


def test_family_records_and_zip(session):
    student, exam = _seed_student(session)
    row = pcs.save_record(session, student.id, exam.id, "report", "成绩单内容")
    session.commit()
    records = pcs.list_records(session, student_id=student.id)
    assert records and records[0]["content"] == "成绩单内容"
    assert records[0]["id"] == row.id

    steps = []
    data = pcs.export_family_zip(session, [student.id], exam_id=exam.id,
                                 progress=lambda done, total:
                                 steps.append((done, total)))
    assert data[:2] == b"PK"
    assert steps == [(1, 1)]


# ---------------------------------------------------------------------------
# 功能9：导出优化
# ---------------------------------------------------------------------------

def test_export_template_service(session, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from utils import export_template_service as ets
    monkeypatch.setattr(ets, "EXPORT_TEMPLATE_DIR", tmp_path / "exports")

    buffer = io.BytesIO()
    wb = Workbook()
    wb.active["A1"] = "姓名"
    wb.save(buffer)
    row = ets.save_export_template(session, "成绩单模板", "excel",
                                   buffer.getvalue())
    session.commit()
    assert ets.list_export_templates(session)[0].id == row.id
    ets.set_default(session, row.id)
    session.commit()
    assert ets.get_default(session, "excel").id == row.id
    preview = ets.preview(session, row.id)
    assert preview["kind"] == "excel" and preview["rows"][0][0] == "姓名"
    assert ets.delete_export_template(session, row.id) is True
    session.commit()
    with pytest.raises(ValueError):
        ets.save_export_template(session, "坏模板", "excel", b"not-excel")


def test_styled_excel_and_multi_sheet():
    import pandas as pd
    from openpyxl import load_workbook
    df = pd.DataFrame([{"姓名": "张三", "分数": 85}])
    data = es.styled_excel(df, sheet_name="成绩", title="成绩单")
    assert data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(data))
    ws = wb["成绩"]
    assert ws["A1"].value == "成绩单"
    assert ws["A2"].value == "姓名"
    assert ws["A2"].font.bold is True
    wb.close()

    multi = es.multi_sheet_excel([
        {"name": "明细", "df": df},
        {"name": "统计", "df": pd.DataFrame([{"项目": "均分", "值": 85}])},
    ])
    wb2 = load_workbook(io.BytesIO(multi))
    assert wb2.sheetnames == ["明细", "统计"]
    wb2.close()


def test_fill_school_template(tmp_path):
    from openpyxl import Workbook, load_workbook
    path = tmp_path / "school.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "分数"])
    ws.append([None, None])
    wb.save(path)

    data = es.fill_school_template(path, {"姓名": ["张三", "李四"],
                                          "分数": [85, 90]})
    out = load_workbook(io.BytesIO(data))
    sheet = out[out.sheetnames[0]]
    assert sheet["A2"].value == "张三"
    assert sheet["B3"].value == 90
    out.close()
    with pytest.raises(ValueError):
        es.fill_school_template(path, {"没有这个表头": 1})


def test_pdf_header_footer_and_bookmarks():
    import fitz
    doc = fitz.open()
    doc.new_page()
    doc.new_page()
    raw = doc.tobytes()
    doc.close()
    out = es.pdf_with_header_footer(
        raw, school="测试小学", title="家校成绩单", footer="教务处",
        page_size="A4", toc=[{"title": "第一章", "page": 1}])
    with fitz.open(stream=out, filetype="pdf") as check:
        assert check.page_count == 2
        assert check.get_toc()
        text = check[0].get_text()
        assert "测试小学" in text and "第 1 / 2 页" in text


def test_docx_and_exam_exports(session, tmp_path):
    from docx import Document
    path = tmp_path / "school.docx"
    doc = Document()
    doc.add_paragraph("学生：{{姓名}}　班级：{{班级}}")
    doc.save(path)
    data = es.docx_from_school_template(path, {"姓名": "张三", "班级": "一班"})
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "张三" in text and "一班" in text
    with pytest.raises(ValueError):
        es.docx_from_school_template(path, {"不存在": 1})

    hw, question, link = _seed_exam_homework(session)
    exam_doc = es.standard_exam_docx(hw, [(link, question)],
                                     school="测试小学", with_answer=True,
                                     header="单元测试卷")
    parsed = Document(io.BytesIO(exam_doc))
    body = "\n".join(p.text for p in parsed.paragraphs)
    assert "单元测试卷" in body
    assert "1+1=?" in body and "答案：2" in body
    assert len(parsed.tables) >= 1  # 密封线表格


# ---------------------------------------------------------------------------
# 页面冒烟
# ---------------------------------------------------------------------------

def test_v340_pages_apptest(tmp_path):
    at = AppTest.from_string(
        _isolated_app_code(tmp_path / "v340_app.db",
                           tmp_path / "feature.json"),
        default_timeout=30)
    at.run()
    assert not at.exception, [str(e) for e in at.exception]

    goto_sub(at, "analysis_tab", "学生画像", "📊 学情")
    assert not at.exception, [str(e) for e in at.exception]
    assert any("家校沟通" in str(tab.label) for tab in at.tabs)

    at.session_state["app_top_page"] = "📁 我的模板"
    at.run()
    assert not at.exception, [str(e) for e in at.exception]
    assert any("导出模板" in str(tab.label) for tab in at.tabs)

    at.session_state["app_top_page"] = "⚙️ 设置"
    at.run()
    assert not at.exception, [str(e) for e in at.exception]
    assert any("回收站" in str(item.value) for item in at.subheader)