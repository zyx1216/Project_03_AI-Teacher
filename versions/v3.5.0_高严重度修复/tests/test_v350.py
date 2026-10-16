# -*- coding: utf-8 -*-
"""v3.5.0 高严重度问题修复测试。"""
from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from streamlit.testing.v1 import AppTest

from models.models import (
    Base, Homework, HomeworkQuestion, Question, Textbook,
)
from tests.test_app_smoke import _isolated_app_code, goto_sub
from utils import excel_handler as eh
from utils import export_service, homework_service, user_config


@pytest.fixture()
def session(tmp_path):
    eng = create_engine(f"sqlite:///{(tmp_path / 'v350.db').as_posix()}",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    s = sessionmaker(bind=eng)()
    yield s
    s.close()
    eng.dispose()


def _seed_homework(session, homework_type="exam", name="期中考试卷"):
    """两道题 + 一份作业/试卷（分值 4、6）。"""
    questions = [
        Question(content="下列函数定义域为 R 的是（  ）\nA. y=1/x\n"
                         "B. y=√x\nC. y=x²\nD. y=lnx",
                 question_type="choice", answer="C", subject="数学",
                 grade="十年级", difficulty=1, status="approved",
                 knowledge_points='["函数"]'),
        Question(content="已知 f(x)=2x+1，则 f(3)=____。",
                 question_type="fill", answer="7", subject="数学",
                 grade="十年级", difficulty=2, status="approved",
                 knowledge_points='["函数"]'),
    ]
    session.add_all(questions)
    session.flush()
    homework = Homework(name=name, homework_type=homework_type,
                        class_name="高一1班", subject="数学", grade="十年级",
                        total_score=100, duration=90)
    session.add(homework)
    session.flush()
    for index, question in enumerate(questions, start=1):
        session.add(HomeworkQuestion(homework_id=homework.id,
                                     question_id=question.id, order=index,
                                     score=[4, 6][index - 1]))
    session.commit()
    return homework, questions


def _school(tmp_path, monkeypatch, name="测试中学"):
    monkeypatch.setattr(user_config, "SCHOOL_PROFILE_FILE",
                        tmp_path / "school_profile.json")
    user_config.save_school_profile(school_name=name)


def _docx_text(data: bytes) -> str:
    """读 docx 文本；去掉换行以兼容逐字竖排的密封线。"""
    doc = Document(io.BytesIO(data))
    body = "".join(p.text for p in doc.paragraphs)
    tables = "".join(c.text for t in doc.tables for r in t.rows
                     for c in r.cells)
    return (body + tables).replace("\n", "").replace("\r", "")

# ---------------------------------------------------------------------------
# 修复1：试卷导出标准化
# ---------------------------------------------------------------------------

def test_exam_word_has_seal_notice_score(session, tmp_path, monkeypatch):
    """标准试卷 Word：学校名/密封线/注意事项/逐题分值/得分栏齐全。"""
    _school(tmp_path, monkeypatch)
    homework, _questions = _seed_homework(session)
    data = homework_service.export_word(session, homework.id,
                                        with_answer=False)
    text = _docx_text(data)
    assert "测试中学" in text
    assert "密封线内不要答题" in text
    assert "注意事项" in text and "答题前请将姓名" in text
    assert "4分" in text and "6分" in text      # 逐题分值
    assert "得分" in text                        # 得分栏
    assert "班级：" in text and "学号：" in text
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) >= 2                  # 密封线区块 + 得分栏


def test_normal_homework_word_has_no_seal(session, tmp_path, monkeypatch):
    """普通作业不加密封线/注意事项（保持原有习题卷样式）。"""
    _school(tmp_path, monkeypatch)
    homework, _questions = _seed_homework(session, homework_type="after_class",
                                          name="课后练习")
    data = homework_service.export_word(session, homework.id,
                                        with_answer=False)
    text = _docx_text(data)
    assert "密封线" not in text and "注意事项" not in text


def test_exam_pdf_has_seal_notice_score(session, tmp_path, monkeypatch):
    """PDF 试卷：能打开且包含与 Word 一致的标准要素。"""
    _school(tmp_path, monkeypatch)
    homework, _questions = _seed_homework(session)
    pairs = homework_service.homework_questions(session, homework.id)
    pdf = export_service.exam_pdf(homework, pairs)
    assert pdf[:4] == b"%PDF"
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        assert doc.page_count >= 1
        text = doc[0].get_text()
    assert "测试中学" in text
    assert "密封线" in text and "注意事项" in text
    assert "4分" in text and "得分" in text


def test_standard_exam_docx_uses_same_shell(session, tmp_path, monkeypatch):
    """standard_exam_docx 与作业/试卷导出共用同一外壳。"""
    _school(tmp_path, monkeypatch)
    homework, questions = _seed_homework(session)
    pairs = homework_service.homework_questions(session, homework.id)
    text = _docx_text(export_service.standard_exam_docx(
        homework, pairs, school="另一所学校"))
    assert "另一所学校" in text
    assert "密封线内不要答题" in text
    assert "注意事项" in text and "得分" in text
    text2 = _docx_text(export_service.standard_exam_docx(homework, questions))
    assert "密封线内不要答题" in text2


def test_homework_latex_has_exam_elements(session, tmp_path, monkeypatch):
    """LaTeX 试卷补齐密封线/注意事项/逐题分值/得分栏。"""
    _school(tmp_path, monkeypatch)
    homework, _questions = _seed_homework(session)
    tex = export_service.homework_latex(
        homework, homework_service.homework_questions(session, homework.id))
    assert "密封线" in tex and "注意事项" in tex
    assert "4分" in tex and "得分栏" in tex
    plain, _q2 = _seed_homework(session, homework_type="after_class",
                                name="普通作业")
    tex2 = export_service.homework_latex(
        plain, homework_service.homework_questions(session, plain.id))
    assert "密封线" not in tex2

# ---------------------------------------------------------------------------
# 修复2：合并单元格向下填充
# ---------------------------------------------------------------------------

def _merged_student_xlsx() -> io.BytesIO:
    buffer = io.BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", "班级", "性别"])
    sheet.append(["甲", "高一2班", "男"])
    sheet.append(["乙", None, "女"])
    sheet.append(["丙", None, "男"])
    sheet.merge_cells("B2:B4")
    book.save(buffer)
    buffer.seek(0)
    return buffer


def test_merged_cells_student_class_filled():
    """合并单元格的学生名单：向下填充后所有行都有班级。"""
    df = eh.load_dataframe(_merged_student_xlsx())
    filled_df, filled_cols = eh.fill_merged_columns(df, ["班级"])
    assert filled_cols == ["班级"]
    detected = eh.detect_student_columns(filled_df)
    records = eh.extract_students(filled_df, detected["mapping"])
    assert [r["class_name"] for r in records] == ["高一2班"] * 3
    # 不先调用 fill_merged_columns 时，extract_students 内部也会兜底填充
    records2 = eh.extract_students(df, detected["mapping"])
    assert all(r["class_name"] for r in records2)


def test_merged_cells_score_class_filled():
    """成绩表合并的班级列：extract_scores 兜底填充。"""
    buffer = io.BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", "班级", "数学"])
    sheet.append(["甲", "高一2班", 90])
    sheet.append(["乙", None, 80])
    sheet.merge_cells("B2:B3")
    book.save(buffer)
    buffer.seek(0)
    df = eh.load_dataframe(buffer)
    detected = eh.detect_score_columns(df)
    records = eh.extract_scores(df, detected)
    assert [r["class_name"] for r in records] == ["高一2班", "高一2班"]


# ---------------------------------------------------------------------------
# 修复3：选中值失效不再崩页
# ---------------------------------------------------------------------------

def _seed_textbook(db_file, *, name, subject, grade, chapter_title):
    """往隔离库里塞一份带章节目录的资料，返回资料 id。"""
    import json

    eng = create_engine(f"sqlite:///{db_file.as_posix()}")
    Base.metadata.create_all(eng)
    session = sessionmaker(bind=eng)()
    book = Textbook(name=name, file_type="pdf", subject=subject, grade=grade,
                    chapter_info=json.dumps(
                        [{"title": chapter_title}], ensure_ascii=False))
    session.add(book)
    session.commit()
    book_id = book.id
    session.close()
    eng.dispose()
    return book_id


def test_stale_material_selection_resets(tmp_path):
    """资料被年级筛选掉后，旧选中值自动重置，页面不崩。"""
    db_file = tmp_path / "stale.db"
    text_dir = tmp_path / "text"
    text_dir.mkdir()
    book_id = _seed_textbook(db_file, name="八年级数学", subject="数学",
                             grade="八年级", chapter_title="第一单元")
    (text_dir / f"{book_id}.txt").write_text("第一单元\n正文内容。",
                                             encoding="utf-8")

    at = AppTest.from_string(
        _isolated_app_code(db_file, tmp_path / "feature.json",
                           text_dir=text_dir),
        default_timeout=30)
    at.run()
    goto_sub(at, "lesson_plan_tab", "AI 备课", "📚 备课")
    next(x for x in at.selectbox
         if x.key == "lesson_grade_select").set_value("初二").run()
    next(x for x in at.selectbox
         if x.key == "lesson_material_select").set_value(book_id).run()
    assert not at.exception, [str(e) for e in at.exception]

    next(x for x in at.selectbox
         if x.key == "lesson_grade_select").set_value("初一").run()
    assert not at.exception, [str(e) for e in at.exception]
    assert at.session_state.get("lesson_material_select") is None


def test_unit_design_candidate_switch_no_crash(tmp_path):
    """单元设计：切换年级导致候选变化时，旧选中项不再触发 index() 崩溃。"""
    db_file = tmp_path / "unit.db"
    _seed_textbook(db_file, name="八年级数学", subject="数学", grade="八年级",
                   chapter_title="第八单元")
    _seed_textbook(db_file, name="高一数学", subject="数学", grade="十年级",
                   chapter_title="第一单元")

    at = AppTest.from_string(
        _isolated_app_code(db_file, tmp_path / "feature.json"),
        default_timeout=30)
    at.run()
    goto_sub(at, "lesson_plan_tab", "📚 单元整体设计", "📚 备课")
    # 学科默认是语文，先切到数学，否则候选为空、不会渲染 unit_candidate
    next(x for x in at.selectbox
         if x.key == "unit_subject").set_value("数学").run()
    next(x for x in at.text_input if x.key == "unit_grade").set_value(
        "八年级").run()
    assert not at.exception, [str(e) for e in at.exception]
    next(x for x in at.selectbox
         if x.key == "unit_candidate").set_value(
             "第八单元（八年级数学）").run()
    assert not at.exception, [str(e) for e in at.exception]
    next(x for x in at.text_input if x.key == "unit_grade").set_value(
        "十年级").run()
    assert not at.exception, [str(e) for e in at.exception]