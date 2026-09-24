# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r'C:\Users\zyx\Downloads\数学习题卷.docx')
out = []
out.append(f'=== 段落总数: {len(doc.paragraphs)} ===')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        style = p.style.name if p.style else ''
        out.append(f'[{i}][{style}] {p.text}')
out.append(f'\n=== 表格数: {len(doc.tables)} ===')
for ti, table in enumerate(doc.tables):
    out.append(f'表格{ti}: {len(table.rows)}行 x {len(table.columns)}列')
with open(r'D:\Codex\Project_03_AI数学教师工作台\docx_content.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('done')
