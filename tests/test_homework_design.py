# -*- coding: utf-8 -*-
"""作业设计测试：建作业、加题、调序、分值、模板复制、级联删除、两卷导出。"""

import io

from utils import homework_service as hw
from utils import question_service as qs
from utils import question_importer as qi
from models.models import Homework, HomeworkQuestion, HomeworkScore, HomeworkAnswer


def _make_question(session, qid_marker="题", qtype="solution", diff=2,
                   answer="答", status="approved", source="manual"):
    data = qs.validate_question({
        "content": f"{qid_marker}{qtype}{diff}", "question_type": qtype,
        "difficulty": diff, "answer": answer, "analysis": "解析",
        "knowledge_points": "方程"})
    return qs.create_question(session, data, source=source, status=status)


def test_create_homework_defaults(session):
    h = hw.create_homework(session, "预习1", homework_type="preview")
    session.flush()
    assert h.duration == hw.DEFAULT_PARAMS["preview"]["duration"]
    assert h.total_score == hw.DEFAULT_PARAMS["preview"]["total_score"]
    assert h.is_template is False


def test_add_bank_questions_and_dedup(session):
    h = hw.create_homework(session, "作业")
    q1 = _make_question(session, "甲")
    q2 = _make_question(session, "乙")
    added = hw.add_questions(session, h.id, [q1.id, q2.id, q1.id])
    assert added == 2  # 重复加入被忽略
    assert hw.question_ids_in_homework(session, h.id) == {q1.id, q2.id}


def test_remove_reorders(session):
    h = hw.create_homework(session, "作业")
    qs3 = [_make_question(session, f"题{i}") for i in range(3)]
    hw.add_questions(session, h.id, [q.id for q in qs3])
    hw.remove_question(session, h.id, qs3[0].id)
    pairs = hw.homework_questions(session, h.id)
    assert [q.id for _l, q in pairs] == [qs3[1].id, qs3[2].id]
    assert [link.order for link, _q in pairs] == [1, 2]


def test_move_question(session):
    h = hw.create_homework(session, "作业")
    qs3 = [_make_question(session, f"题{i}") for i in range(3)]
    hw.add_questions(session, h.id, [q.id for q in qs3])
    hw.move_question(session, h.id, qs3[0].id, 1)  # 第一题下移
    pairs = hw.homework_questions(session, h.id)
    assert [q.id for _l, q in pairs] == [qs3[1].id, qs3[0].id, qs3[2].id]
    # 边界不动
    hw.move_question(session, h.id, qs3[1].id, -5)
    assert hw.homework_questions(session, h.id)[0][1].id == qs3[1].id


def test_score_sum(session):
    h = hw.create_homework(session, "作业")
    qs3 = [_make_question(session, f"题{i}") for i in range(3)]
    hw.add_questions(session, h.id, [q.id for q in qs3])
    hw.set_question_score(session, h.id, qs3[0].id, 10)
    hw.set_question_score(session, h.id, qs3[1].id, 15)
    assert hw.summed_score(session, h.id) == 25


def test_imported_and_manual_sources(session):
    h = hw.create_homework(session, "作业")
    # 外部导入题
    rec = {"row_no": 1, "content": "导入题", "question_type": "填空题",
           "answer": "答案", "analysis": "", "knowledge_points": "",
           "difficulty": 1, "error_points": "", "problems": []}
    result = qi.import_records(session, [rec])
    assert result["imported"] == 1
    imported = qs.list_questions(session, source="imported")
    hw.add_questions(session, h.id, [q.id for q in imported])
    # 手动题
    manual = _make_question(session, "手动", source="manual")
    hw.add_questions(session, h.id, [manual.id])
    assert len(hw.homework_questions(session, h.id)) == 2


def test_save_as_template_and_copy(session):
    h = hw.create_homework(session, "原作业")
    q1 = _make_question(session, "模板题")
    hw.add_questions(session, h.id, [q1.id], default_score=8)
    tpl = hw.save_as_template(session, h.id)
    session.flush()
    assert tpl.is_template is True
    # 模板出现在模板列表、不在普通作业列表
    assert tpl in hw.list_homeworks(session, templates=True)
    assert tpl not in hw.list_homeworks(session, templates=False)

    # 从模板新建作业，题目（含分值）被复制
    h2 = hw.create_homework(session, "新作业", template_id=tpl.id)
    pairs = hw.homework_questions(session, h2.id)
    assert len(pairs) == 1 and pairs[0][1].id == q1.id
    assert pairs[0][0].score == 8
    # 模板与作业的关联行是各自独立的
    assert pairs[0][0].id != hw.homework_questions(session, h.id)[0][0].id


def test_delete_homework_cascades(session):
    from utils import student_service as ss
    h = hw.create_homework(session, "待删作业")
    q1 = _make_question(session, "题")
    hw.add_questions(session, h.id, [q1.id])
    stu, _ = ss.get_or_create_student(session, "学生01", "一班")
    session.add(HomeworkScore(homework_id=h.id, student_id=stu.id, total_score=80))
    session.add(HomeworkAnswer(homework_id=h.id, student_id=stu.id,
                               question_id=q1.id, is_correct=False))
    session.flush()
    hw.delete_homework(session, h.id)
    session.flush()
    # 关联的题目链接、成绩、作答都被级联删除；题库题目本身保留
    assert session.query(HomeworkQuestion).count() == 0
    assert session.query(HomeworkScore).count() == 0
    assert session.query(HomeworkAnswer).count() == 0
    assert session.get(type(q1), q1.id) is not None


def test_export_word_two_modes(session):
    h = hw.create_homework(session, "导出作业")
    q1 = _make_question(session, "导题干", answer="导出答案")
    hw.add_questions(session, h.id, [q1.id])
    from docx import Document
    stu_bytes = hw.export_word(session, h.id, with_answer=False)
    tea_bytes = hw.export_word(session, h.id, with_answer=True)
    stu_text = "\n".join(p.text for p in Document(io.BytesIO(stu_bytes)).paragraphs)
    tea_text = "\n".join(p.text for p in Document(io.BytesIO(tea_bytes)).paragraphs)
    assert "导题干" in stu_text and "导出答案" not in stu_text
    assert "导出答案" in tea_text and "教师卷" in tea_text