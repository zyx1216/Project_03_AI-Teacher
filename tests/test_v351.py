# -*- coding: utf-8 -*-
"""v3.5.1 中严重度问题修复测试。"""
from __future__ import annotations

import io
import json
from datetime import date

import pandas as pd
import pytest
from docx import Document
from docx.shared import Pt
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from models.models import Base, Exam, LessonPlan, Score, Student
from utils import (exam_service, excel_handler as eh, export_service,
                   homework_service, lesson_service, material_service,
                   parent_report_service, question_service,
                   student_report_service)


@pytest.fixture()
def session(tmp_path):
    eng = create_engine(f"sqlite:///{(tmp_path / 'v351.db').as_posix()}",
                        connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    s = sessionmaker(bind=eng)()
    yield s
    s.close()
    eng.dispose()


# ---------------------------------------------------------------------------
# 1) 章节提取三连
# ---------------------------------------------------------------------------

def test_split_chapters_filters_toc_lines():
    text = ("目录\n第一章 有理数 …………………… 1\n"
            "1.1 正数和负数 ………………… 2\n第二章 整式 ……………………… 10\n"
            "\n第一章 有理数\n1.1 正数和负数\n内容甲。")
    chapters = material_service.split_chapters(text)
    titles = [c["title"] for c in chapters]
    assert all("………" not in t for t in titles)      # 目录行不当章节
    assert all(c["content"] for c in chapters)        # 不产出空章节
    assert titles.count("1.1 正数和负数") == 1        # 无重复章节
    assert material_service.is_toc_line("第一章 有理数 …… 3")
    assert material_service.is_toc_line("1.1 正数和负数 ... 2")
    assert not material_service.is_toc_line("第一章 有理数")


def test_split_chapters_levels_and_display_title():
    text = "第一章 有理数\n1.1 正数和负数\n内容甲。\n1.1.1 相反数\n内容乙。"
    chapters = material_service.split_chapters(text)
    by_title = {c["title"]: c for c in chapters}
    assert by_title["1.1 正数和负数"]["level"] == 2
    assert by_title["1.1.1 相反数"]["level"] == 3
    assert by_title["1.1 正数和负数"]["display_title"] == \
        "第一章 有理数 > 1.1 正数和负数"
    # 空的一级标题不单列，只作子节的 parent
    assert "第一章 有理数" not in by_title
    assert material_service.chapter_titles(chapters, hierarchical=True) == [
        "第一章 有理数 > 1.1 正数和负数",
        "第一章 有理数 > 1.1.1 相反数"]


def test_split_chapters_review_attached():
    text = "第一章 有理数\n章节内容。\n整理与复习\n复习题。"
    chapters = material_service.split_chapters(text)
    review = next(c for c in chapters if c["title"] == "整理与复习")
    assert review["is_review"] is True
    assert review["level"] == 2
    assert review["parent"] == "第一章 有理数"
    assert review["display_title"] == "第一章 有理数 > 整理与复习"

# ---------------------------------------------------------------------------
# 2) 成绩导入满分上限校验
# ---------------------------------------------------------------------------

def test_extract_scores_over_full_score_blanked():
    df = pd.DataFrame({"姓名": ["张三", "李四"], "数学": [99, 999]})
    det = eh.detect_score_columns(df)
    records = eh.extract_scores(df, det, full_scores={"数学": 120})
    good = next(r for r in records if r["name"] == "张三")
    bad = next(r for r in records if r["name"] == "李四")
    assert good["scores"]["数学"] == 99 and good["problems"] == []
    assert bad["scores"]["数学"] is None
    assert any("超出满分" in p for p in bad["problems"])


def test_extract_scores_negative_blanked():
    df = pd.DataFrame({"姓名": ["张三", "李四"], "数学": [99, -5]})
    det = eh.detect_score_columns(df)
    records = eh.extract_scores(df, det, full_scores={"数学": 120})
    bad = next(r for r in records if r["name"] == "李四")
    assert bad["scores"]["数学"] is None
    assert any("分数为负" in p for p in bad["problems"])


# ---------------------------------------------------------------------------
# 3) 排名（报告时现算）
# ---------------------------------------------------------------------------

def _seed_exam(session):
    exam = Exam(name="七年级期中", exam_date=date(2026, 5, 1), grade="七年级",
                full_scores=json.dumps({"数学": 120}, ensure_ascii=False))
    session.add(exam)
    session.flush()
    students = []
    for name, class_name, score in [("张三", "一班", 110), ("李四", "一班", 100),
                                    ("王五", "二班", 120), ("赵六", "二班", 90)]:
        stu = Student(name=name, class_name=class_name)
        session.add(stu)
        session.flush()
        session.add(Score(exam_id=exam.id, student_id=stu.id,
                          subject="数学", score=score))
        students.append(stu)
    session.commit()
    return exam, students


def test_student_ranks_class_and_grade(session):
    exam, students = _seed_exam(session)
    zhang = students[0]
    ranks = exam_service.student_ranks(session, exam.id, zhang.id, subject="数学")
    assert ranks["class_rank"] == 1 and ranks["class_size"] == 2
    assert ranks["grade_rank"] == 2 and ranks["grade_size"] == 4


# ---------------------------------------------------------------------------
# 4) 家长 / 学生报告补参照信息
# ---------------------------------------------------------------------------

def test_parent_report_includes_reference(session):
    exam, students = _seed_exam(session)
    zhang = students[0]
    report = parent_report_service.generate_parent_report(
        session, zhang.id, exam_id=exam.id, chat_func=lambda s, u: "")
    html = parent_report_service.render_html(report)
    assert "满分" in html and "得分率" in html
    assert "班级排名" in html and "年级排名" in html
    assert "110分" in report["subjects"]
    assert "满分120" in report["subjects"]
    assert "得分率91.7%" in report["subjects"]
    assert "班级1/2" in report["subjects"]
    assert "年级2/4" in report["subjects"]
    assert "优秀" in report["subjects"]      # 110/120 = 91.7%
    assert "None" not in html


def test_student_report_table_and_advice(session):
    exam, students = _seed_exam(session)
    zhang = students[0]
    report = student_report_service.generate_student_report(
        session, zhang.id, chat_func=lambda s, u: "")
    html = student_report_service.render_student_report_html(report)
    assert "满分" in html and "得分率" in html
    assert "班级排名" in html and "年级排名" in html
    assert "110.0" not in html               # 整数分不带 .0
    assert "table-layout:fixed" in html

    class _Stu:
        name = "张三"
    ok = student_report_service._build_advice(
        _Stu(), [], [], chat_func=lambda s, u: "")
    weak = student_report_service._build_advice(
        _Stu(), [], [{"knowledge_point": "一元一次方程", "rate": 0.3}],
        chat_func=lambda s, u: "")
    assert "各科表现均衡" in ok
    assert "建议重点复习" in weak and "一元一次方程" in weak

# ---------------------------------------------------------------------------
# 5) Word 字体字号规范（教案 + 试卷）
# ---------------------------------------------------------------------------

def _lesson_plan_dict():
    return {
        "subject": "数学",
        "objectives": {"knowledge": "认识有理数", "process": "探究过程",
                       "emotion": "培养兴趣"},
        "key_points": "有理数的概念", "difficult_points": "运算",
        "process": [{"stage": "导入", "minutes": 5, "content": "举例"}],
        "board_design": "板书内容", "reflection": "反思内容",
    }


def test_lesson_word_fonts_follow_spec():
    lesson = LessonPlan(title="有理数第一课时", grade="七年级",
                        chapter="有理数", content="{}")
    data = lesson_service.export_word(lesson, _lesson_plan_dict())
    doc = Document(io.BytesIO(data))
    normal = doc.styles["Normal"]
    assert normal.font.name == "宋体" and normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5
    title = doc.styles["Title"]
    assert title.font.name == "黑体" and title.font.size == Pt(16)
    assert title.font.bold is True
    h1 = doc.styles["Heading 1"]
    assert h1.font.name == "黑体" and h1.font.size == Pt(14)
    assert h1.font.bold is True
    h2 = doc.styles["Heading 2"]
    assert h2.font.name == "宋体" and h2.font.size == Pt(12)
    from docx.oxml.ns import qn
    rfonts = normal.element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == "宋体"


def _seed_exam_homework(session):
    q = question_service.create_question(
        session, {"content": "1+1=?____。", "question_type": "fill",
                  "answer": "2", "difficulty": 1},
        status="approved", subject="数学", grade="七年级")
    hw = homework_service.create_homework(
        session, name="七年级期中试卷", homework_type="exam", subject="数学",
        grade="七年级", class_name="一班", total_score=120, duration=90)
    homework_service.add_questions(session, hw.id, [q.id], default_score=120)
    session.commit()
    return hw


def test_exam_seal_font_size(session):
    hw = _seed_exam_homework(session)
    data = homework_service.export_word(session, hw.id, with_answer=False)
    doc = Document(io.BytesIO(data))
    seal_sizes = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "密封线内不要答题" in cell.text.replace("\n", ""):
                    for par in cell.paragraphs:
                        for run in par.runs:
                            if run.font.size is not None:
                                seal_sizes.append(round(run.font.size.pt, 1))
    assert 10.5 in seal_sizes


def test_school_template_docx_keeps_font(tmp_path):
    tpl = tmp_path / "template.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("{{标题}}")
    run.font.name = "楷体"
    run.font.size = Pt(18)
    doc.save(str(tpl))
    data = export_service.docx_from_school_template(tpl, {"标题": "新标题"})
    out = Document(io.BytesIO(data))
    runs = [r for p in out.paragraphs for r in p.runs]
    filled = next(r for r in runs if r.text == "新标题")
    assert filled.font.name == "楷体" and filled.font.size == Pt(18)
