# -*- coding: utf-8 -*-
"""审计脚本5：模板齐全度 + Word 字体（样式级/eastAsia）检查。"""
from __future__ import annotations
import io
import sys
from pathlib import Path

ROOT = Path(r"D:\Codex\Project_03_AI数学教师工作台")
OUT = ROOT / ".tmp_audit"
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn
from utils import template_service

print("#" * 70)
print("# 1. 教案模板清单与模块齐全度")
print("#" * 70)
templates = template_service.list_lesson_templates()
print(f"共 {len(templates)} 个教案模板")
MODULES = ["教学目标", "重点", "难点", "教学过程", "导入", "新授", "练习",
           "小结", "作业", "板书", "反思"]
for t in templates:
    name = t.get("name")
    content = template_service.lesson_template_content(name)
    missing = [m for m in MODULES if m not in content]
    print("=" * 70)
    print(f"## 模板：{name}（描述：{t.get('description') or '—'}）")
    print(f"   提示词长度={len(content)} 字")
    print(f"   缺失模块: {missing or '无'}")
    print(f"   片段: {content[:160].replace(chr(10), ' / ')!r}")

print()
print("#" * 70)
print("# 2. 生成的 Word 字体（样式级 + eastAsia + 字号）")
print("#" * 70)


def font_report(path: Path, label: str) -> None:
    doc = Document(str(path))
    normal = doc.styles["Normal"].font
    east = None
    rpr = doc.styles["Normal"].element.find(qn("w:rPr"))
    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            east = rf.get(qn("w:eastAsia"))
    # 统计正文运行级字体/字号
    run_fonts, run_sizes = set(), set()
    for p in doc.paragraphs:
        for r in p.runs:
            rpr2 = r._element.find(qn("w:rPr"))
            if rpr2 is not None:
                rf2 = rpr2.find(qn("w:rFonts"))
                if rf2 is not None and rf2.get(qn("w:eastAsia")):
                    run_fonts.add(rf2.get(qn("w:eastAsia")))
            if r.font.name:
                run_fonts.add(r.font.name)
            if r.font.size:
                run_sizes.add(round(r.font.size.pt, 1))
    print("=" * 70)
    print(f"### {label}")
    print(f"   Normal 样式字体: {normal.name!r} | eastAsia: {east!r}"
          f" | 字号: {normal.size.pt if normal.size else None}")
    print(f"   运行级字体: {sorted(run_fonts) or '（未在 run 上设置）'}")
    print(f"   运行级字号: {sorted(run_sizes) or '（未在 run 上设置）'}")
    styles_used = sorted({p.style.name for p in doc.paragraphs if p.style})
    print(f"   段落样式: {styles_used}")


for fname, label in (("lesson.docx", "教案 Word"),
                     ("exam_pathA.docx", "试卷 Word 路径A"),
                     ("exam_pathB.docx", "试卷 Word 路径B")):
    p = OUT / fname
    if p.exists():
        font_report(p, label)

print()
print("#" * 70)
print("# 3. 试卷 Word 正文（检查题号/分值/密封线等要素）")
print("#" * 70)
for fname in ("exam_pathA.docx", "exam_pathB.docx"):
    doc = Document(str(OUT / fname))
    body = "\n".join(p.text for p in doc.paragraphs)
    tables = [" | ".join(c.text for c in row.cells)
              for t in doc.tables for row in t.rows]
    print("=" * 70)
    print(f"### {fname}")
    print("   段落文本前 300 字:")
    print("   " + body[:300].replace("\n", "\n   "))
    print("   表格内容:", tables[:4])